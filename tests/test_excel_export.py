import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from src.reports.excel_report import ExcelReport


class DummyConfig:
    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir


def _sample_frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Cuenta": ["1000", "2000"],
            "Importe": [1250.5, 10.0],
            "Fecha": pd.to_datetime(["2026-01-15", "2026-02-01"]),
        }
    )


def _tmp_names(directory: Path) -> list[Path]:
    return list(directory.glob(".MATRIZ_FBL1N_*.tmp.xlsx"))


class ExcelExportTestCase(unittest.TestCase):
    def test_export_creates_workbook_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))

            dataframe = pd.DataFrame(
                {
                    "Cuenta": ["1000"],
                    "Importe": [1250.5],
                }
            )

            output_path = exporter.export(dataframe)

            self.assertTrue(output_path.exists())
            workbook = load_workbook(output_path)
            self.assertEqual(workbook.active.title, "MATRIZ_FBL1N")
            self.assertEqual(workbook.active.max_row, 2)
            self.assertEqual(workbook.active.max_column, 2)
            workbook.close()

    def test_sheet_order_styles_formats_and_success_cleanup(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            additional = {
                "TD_CLP_NO_COM": pd.DataFrame({"Concepto": ["A"], "Total": [1.0]}),
                "RESUMEN_CONCEPTOS": pd.DataFrame(
                    {"Concepto": ["A"], "Cantidad": [2]}
                ),
            }

            output_path = exporter.export(_sample_frame(), additional_sheets=additional)

            self.assertEqual(output_path, output_dir / output_path.name)
            self.assertTrue(output_path.name.startswith("MATRIZ_FBL1N_"))
            self.assertTrue(output_path.suffix == ".xlsx")
            self.assertTrue(output_path.exists())
            self.assertGreater(output_path.stat().st_size, 0)
            self.assertEqual(_tmp_names(output_dir), [])
            self.assertEqual(len(list(output_dir.glob("MATRIZ_FBL1N_*.xlsx"))), 1)

            workbook = load_workbook(output_path)
            self.assertEqual(
                workbook.sheetnames,
                ["MATRIZ_FBL1N", "TD_CLP_NO_COM", "RESUMEN_CONCEPTOS"],
            )
            sheet = workbook["MATRIZ_FBL1N"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["Cuenta", "Importe", "Fecha"],
            )
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertIsNotNone(sheet.auto_filter.ref)
            header = sheet["A1"]
            self.assertEqual(str(header.fill.fgColor.rgb)[-6:].upper(), "1F4E78")
            self.assertEqual(str(header.fill.patternType), "solid")
            self.assertIsNotNone(header.border.left.style)
            self.assertEqual(sheet["B2"].number_format, "#,##0.00")
            self.assertEqual(sheet["C2"].number_format, "yyyy-mm-dd")
            workbook.close()

    def test_save_failure_leaves_no_final_and_removes_temp(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            with patch.object(Workbook, "save", side_effect=OSError("save failed")):
                with self.assertRaises(OSError):
                    exporter.export(_sample_frame())
            self.assertEqual(list(output_dir.glob("MATRIZ_FBL1N_*.xlsx")), [])
            self.assertEqual(_tmp_names(output_dir), [])

    def test_ooxml_validation_failure_leaves_no_final_and_removes_temp(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            with patch.object(
                ExcelReport,
                "_validate_exported_workbook",
                side_effect=ValueError("bad ooxml"),
            ):
                with self.assertRaises(ValueError):
                    exporter.export(_sample_frame())
            self.assertEqual(list(output_dir.glob("MATRIZ_FBL1N_*.xlsx")), [])
            self.assertEqual(_tmp_names(output_dir), [])

    def test_replace_failure_leaves_no_final_and_removes_temp(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            with patch(
                "src.reports.excel_report.os.replace",
                side_effect=OSError("replace failed"),
            ):
                with self.assertRaises(OSError):
                    exporter.export(_sample_frame())
            self.assertEqual(list(output_dir.glob("MATRIZ_FBL1N_*.xlsx")), [])
            self.assertEqual(_tmp_names(output_dir), [])

    def test_existing_destination_aborts_without_modification(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            dest = output_dir / "MATRIZ_FBL1N_20990101_000000.xlsx"
            dest.write_bytes(b"PREEXISTING")
            with patch.object(exporter, "_build_output_path", return_value=dest):
                with self.assertRaises(FileExistsError):
                    exporter.export(_sample_frame())
            self.assertEqual(dest.read_bytes(), b"PREEXISTING")
            self.assertEqual(_tmp_names(output_dir), [])

    def test_progress_logs_without_large_dataset(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            records: list[str] = []

            def _capture(message: str, *args: object) -> None:
                records.append(message % args if args else message)

            exporter._emit = _capture  # type: ignore[method-assign]
            frame = pd.DataFrame({"Cuenta": ["a", "b", "c", "d"], "Importe": [1, 2, 3, 4]})
            with patch("src.reports.excel_report._MATRIX_PROGRESS_EVERY", 2):
                output_path = exporter.export(frame)

            joined = "\n".join(records)
            self.assertIn("export pid=", joined)
            self.assertIn("fase inicio construccion workbook", joined)
            self.assertIn("hoja=MATRIZ_FBL1N filas=4 columnas=2", joined)
            self.assertIn("MATRIZ_FBL1N progreso filas=2/4", joined)
            self.assertIn("MATRIZ_FBL1N progreso filas=4/4", joined)
            self.assertIn("fase inicio append MATRIZ_FBL1N", joined)
            self.assertIn("fase inicio estilos MATRIZ_FBL1N", joined)
            self.assertIn("fase inicio anchos MATRIZ_FBL1N", joined)
            self.assertIn("fase inicio tipos MATRIZ_FBL1N", joined)
            self.assertIn("fase inicio workbook.save", joined)
            self.assertIn("fase fin workbook.save", joined)
            self.assertIn("fase inicio validacion OOXML", joined)
            self.assertIn("fase inicio os.replace", joined)
            self.assertTrue(output_path.exists())

    def test_rss_helper_failure_does_not_break_export(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_dir = Path(tmp_dir)
            exporter = ExcelReport(config_obj=DummyConfig(output_dir))
            with patch("src.reports.excel_report._rss_mb", return_value=None):
                output_path = exporter.export(_sample_frame())
            self.assertTrue(output_path.exists())


if __name__ == "__main__":
    unittest.main()
