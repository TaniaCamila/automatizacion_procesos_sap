from __future__ import annotations

import logging
import os
import time
import tracemalloc
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook

from ..config.config import VERSION, Config
from ..config.logger import LoggerManager
from ..modules.margen.module import InformeMargenModule
from ..services.output_artifact_service import OutputArtifactService


logger = LoggerManager.get_logger(__name__)


@dataclass
class _MatrixLoadResult:
    matrix: pd.DataFrame
    source_file: Path
    total_rows: int
    classified_rows: int
    usd_rows: int
    no_usd_rows: int


class PivotReportService:
    """
    Genera un Excel con modelo listo para Power BI (BUSINESS-05/06/09).

    BUSINESS-15/18A: dinámicas Power BI. Hojas exportadas:
    PIVOT, TD_CLP_NO_COM, TD_CLP_COM_2026, CLP_USD_NO_COM, CLP_USD_COM,
    KPIS, METADATA. Filtrado por columnas persistidas en la MATRIZ;
    no reconsulta concept_classifier.
    """

    SHEET_NAME = "MATRIZ_FBL1N"
    TARGET_YEAR = 2026
    MODEL_VERSION = "BUSINESS-18B"
    MONTH_LABELS: tuple[str, ...] = (
        "enero",
        "febrero",
        "marzo",
        "abril",
        "mayo",
        "junio",
        "julio",
        "agosto",
        "septiembre",
        "octubre",
        "noviembre",
        "diciembre",
    )

    # Dinámicas por Tipo (BUSINESS-13/18A).
    TD_SHEET_NO_COM = InformeMargenModule._TD_SHEET_NO_COM
    TD_SHEET_COM = InformeMargenModule._TD_SHEET_COM
    TD_SHEET_CLP_USD_NO_COM = InformeMargenModule._TD_SHEET_CLP_USD_NO_COM
    TD_SHEET_CLP_USD_COM = InformeMargenModule._TD_SHEET_CLP_USD_COM
    TD_SHEET_ORDER: tuple[str, ...] = (
        TD_SHEET_NO_COM,
        TD_SHEET_COM,
        TD_SHEET_CLP_USD_NO_COM,
        TD_SHEET_CLP_USD_COM,
    )
    TD_MONTH_LABELS: tuple[str, ...] = InformeMargenModule._TD_MONTH_LABELS

    def __init__(
        self,
        config_obj: Config | None = None,
        output_service: OutputArtifactService | None = None,
        logger_obj: logging.Logger | None = None,
    ) -> None:
        self.config = config_obj or Config()
        self.output_service = output_service or OutputArtifactService(
            config_obj=self.config
        )
        self.logger = logger_obj or logger
        self.last_metrics: dict[str, float] = {}
        self._last_load: _MatrixLoadResult | None = None

    @staticmethod
    def _resolve_header_index(
        header: list[object],
        *candidates: str,
    ) -> int:
        index_by_name = {
            str(name).strip(): index
            for index, name in enumerate(header)
            if name is not None and str(name).strip()
        }
        for candidate in candidates:
            if candidate in index_by_name:
                return index_by_name[candidate]
        raise KeyError(
            "No se encontró ninguna de las columnas: "
            + ", ".join(candidates)
        )

    @staticmethod
    def _try_header_index(
        header: list[object],
        *candidates: str,
    ) -> int | None:
        try:
            return PivotReportService._resolve_header_index(header, *candidates)
        except KeyError:
            return None

    @staticmethod
    def _year_month(value: object) -> tuple[int, int] | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.year, value.month
        if isinstance(value, date):
            return value.year, value.month
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.isna(parsed):
            return None
        return int(parsed.year), int(parsed.month)

    @staticmethod
    def _to_float(value: object) -> float:
        if value is None:
            return 0.0
        if isinstance(value, bool):
            return 0.0
        if isinstance(value, (int, float)):
            if pd.isna(value):
                return 0.0
            return float(value)
        number = pd.to_numeric(value, errors="coerce")
        if pd.isna(number):
            return 0.0
        return float(number)

    @staticmethod
    def _current_user() -> str:
        return (
            os.getenv("USERNAME")
            or os.getenv("USER")
            or os.getenv("LOGNAME")
            or "N/D"
        )

    def load_matrix(self, path: Path | None = None) -> pd.DataFrame:
        """
        Leer streaming columnas necesarias + métricas para KPIS.

        Filtra año TARGET_YEAR para el pivot.
        """

        artifact = path or self.output_service.latest_or_raise()
        self.logger.info("Leyendo hoja %s de %s", self.SHEET_NAME, artifact.name)

        workbook = CalamineWorkbook.from_path(str(artifact))
        sheet = workbook.get_sheet_by_name(self.SHEET_NAME)
        rows = sheet.iter_rows()
        header = next(rows, None)
        empty = pd.DataFrame(
            columns=[
                "concepto_detectado",
                "Fecha compensación",
                "Importe en moneda doc.",
            ]
        )
        if header is None:
            self._last_load = _MatrixLoadResult(
                matrix=empty,
                source_file=artifact,
                total_rows=0,
                classified_rows=0,
                usd_rows=0,
                no_usd_rows=0,
            )
            return empty

        header_list = list(header)
        concept_idx = self._resolve_header_index(
            header_list,
            "concepto_detectado",
            "concepto",
        )
        date_idx = self._resolve_header_index(
            header_list,
            "Fecha compensación",
            "fecha_compensacion",
        )
        amount_idx = self._resolve_header_index(
            header_list,
            "Importe en moneda doc.",
            "importe_en_moneda_doc",
        )
        currency_idx = self._try_header_index(header_list, "currency_group")
        tipo_idx = self._try_header_index(header_list, "tipo")
        moneda_idx = self._try_header_index(
            header_list,
            "moneda_del_documento",
            "Moneda del documento",
        )
        sociedad_idx = self._try_header_index(header_list, "sociedad", "Sociedad")
        acreedor_idx = self._try_header_index(header_list, "acreedor", "Acreedor")
        texto_idx = self._try_header_index(
            header_list,
            "texto_cabdocumento",
            "Texto cab.documento",
        )

        concepts: list[str] = []
        dates: list[object] = []
        amounts: list[float] = []
        tipos: list[str] = []
        monedas: list[str] = []
        sociedades: list[str] = []
        acreedores: list[str] = []
        textos: list[str] = []
        total_rows = 0
        classified_rows = 0
        usd_rows = 0
        no_usd_rows = 0

        for row in rows:
            total_rows += 1
            raw_concept = row[concept_idx] if concept_idx < len(row) else None
            if raw_concept is None or (
                isinstance(raw_concept, float) and pd.isna(raw_concept)
            ):
                concept = ""
            else:
                concept = str(raw_concept).strip()

            if concept:
                classified_rows += 1

            currency_group = ""
            if currency_idx is not None and currency_idx < len(row):
                raw_currency = row[currency_idx]
                if raw_currency is not None and not (
                    isinstance(raw_currency, float) and pd.isna(raw_currency)
                ):
                    currency_group = str(raw_currency).strip().upper()

            year_month = self._year_month(
                row[date_idx] if date_idx < len(row) else None
            )
            if year_month is None or year_month[0] != self.TARGET_YEAR:
                continue
            if not concept:
                continue

            if currency_group == "USD":
                usd_rows += 1
            else:
                # Incluye NO_USD y vacío: alineado a currency_group del clasificador.
                no_usd_rows += 1

            concepts.append(concept)
            dates.append(row[date_idx])
            amounts.append(
                self._to_float(row[amount_idx] if amount_idx < len(row) else None)
            )

            raw_tipo = (
                row[tipo_idx]
                if tipo_idx is not None and tipo_idx < len(row)
                else None
            )
            tipos.append(
                ""
                if raw_tipo is None
                or (isinstance(raw_tipo, float) and pd.isna(raw_tipo))
                else str(raw_tipo).strip().upper()
            )

            raw_moneda = (
                row[moneda_idx]
                if moneda_idx is not None and moneda_idx < len(row)
                else None
            )
            monedas.append(
                ""
                if raw_moneda is None
                or (isinstance(raw_moneda, float) and pd.isna(raw_moneda))
                else str(raw_moneda).strip().upper()
            )

            raw_sociedad = (
                row[sociedad_idx]
                if sociedad_idx is not None and sociedad_idx < len(row)
                else None
            )
            sociedades.append(
                ""
                if raw_sociedad is None
                or (isinstance(raw_sociedad, float) and pd.isna(raw_sociedad))
                else str(raw_sociedad).strip().upper()
            )

            raw_acreedor = (
                row[acreedor_idx]
                if acreedor_idx is not None and acreedor_idx < len(row)
                else None
            )
            acreedores.append(
                InformeMargenModule._normalize_acreedor_id(raw_acreedor)
            )

            raw_texto = (
                row[texto_idx]
                if texto_idx is not None and texto_idx < len(row)
                else None
            )
            textos.append(
                ""
                if raw_texto is None
                or (isinstance(raw_texto, float) and pd.isna(raw_texto))
                else str(raw_texto)
            )

        if not concepts:
            matrix = empty
        else:
            matrix = pd.DataFrame(
                {
                    "concepto_detectado": pd.Series(concepts, dtype="category"),
                    "Fecha compensación": dates,
                    "Importe en moneda doc.": pd.Series(amounts, dtype="float64"),
                    "tipo": pd.Series(tipos, dtype="category"),
                    "moneda_del_documento": pd.Series(monedas, dtype="category"),
                    "sociedad": pd.Series(sociedades, dtype="category"),
                    "acreedor": pd.Series(acreedores, dtype="category"),
                    "texto_cabdocumento": textos,
                }
            )

        self._last_load = _MatrixLoadResult(
            matrix=matrix,
            source_file=artifact,
            total_rows=total_rows,
            classified_rows=classified_rows,
            usd_rows=usd_rows,
            no_usd_rows=no_usd_rows,
        )
        return matrix

    def build_pivot(self, matrix: pd.DataFrame) -> pd.DataFrame:
        """
        Construir pivot:

            Filas    → concepto (concepto_detectado)
            Columnas → meses enero–diciembre
            Valores  → suma de Importe en moneda doc.
            + columna Total
        """

        empty = pd.DataFrame(
            columns=["Concepto de Pago", *self.MONTH_LABELS, "Total"]
        )
        if matrix.empty:
            return empty

        concept_col = (
            "concepto_detectado"
            if "concepto_detectado" in matrix.columns
            else "concepto"
        )
        date_col = (
            "Fecha compensación"
            if "Fecha compensación" in matrix.columns
            else "fecha_compensacion"
        )
        amount_col = (
            "Importe en moneda doc."
            if "Importe en moneda doc." in matrix.columns
            else "importe_en_moneda_doc"
        )

        year_month = matrix[date_col].map(self._year_month)
        months = year_month.map(lambda item: item[1] if item else None)
        valid = months.notna()
        if not bool(valid.any()):
            return empty

        lean = pd.DataFrame(
            {
                "Concepto de Pago": matrix.loc[valid, concept_col]
                .fillna("")
                .astype(str),
                "mes": months.loc[valid].astype("int8"),
                "importe": matrix.loc[valid, amount_col].to_numpy(
                    dtype="float64",
                    copy=False,
                ),
            }
        )

        pivot = lean.pivot_table(
            index="Concepto de Pago",
            columns="mes",
            values="importe",
            aggfunc="sum",
            fill_value=0.0,
            dropna=False,
        )

        for month_number in range(1, 13):
            if month_number not in pivot.columns:
                pivot[month_number] = 0.0
        pivot = pivot.reindex(columns=list(range(1, 13)), fill_value=0.0)
        pivot.columns = list(self.MONTH_LABELS)
        pivot["Total"] = pivot.loc[:, list(self.MONTH_LABELS)].sum(axis=1)
        pivot.index = pivot.index.astype(str)
        pivot = pivot.sort_index(kind="mergesort")
        pivot = pivot.reset_index()
        pivot.columns = ["Concepto de Pago", *self.MONTH_LABELS, "Total"]

        self.logger.info(
            "Pivot margen %d: %d conceptos, importe total=%.2f",
            self.TARGET_YEAR,
            len(pivot),
            float(pivot["Total"].sum()),
        )
        return pivot

    def build_td_dynamics(
        self,
        matrix: pd.DataFrame,
    ) -> dict[str, pd.DataFrame]:
        """
        Construir dinámicas CLP y CLP_USD usando columnas persistidas
        en la MATRIZ (BUSINESS-13A/15A/18A). No reconsulta clasificador.

        Reglas:
            COM / no-COM según `tipo`
            Segmento USD (sociedad + texto USD + acreedor) → CLP_USD_*
        """

        empty = pd.DataFrame(
            columns=["Concepto de Pago", *self.TD_MONTH_LABELS, "Total"]
        )
        dynamics = {name: empty.copy() for name in self.TD_SHEET_ORDER}

        if matrix.empty:
            return dynamics

        work = matrix
        if "moneda_del_documento" in work.columns:
            moneda = (
                work["moneda_del_documento"].astype(str).str.strip().str.upper()
            )
            work = work.loc[moneda == "CLP"]

        if "tipo" in work.columns:
            tipo_series = (
                work["tipo"].fillna("").astype(str).str.strip().str.upper()
            )
            com_mask = tipo_series == "COM"
        else:
            self.logger.warning(
                "Columna 'tipo' ausente en la MATRIZ; "
                "COM vacía y NO_COM con todos los registros."
            )
            com_mask = pd.Series(False, index=work.index)

        # BUSINESS-18A: reutilizar la máscara del módulo margen.
        helper = InformeMargenModule.__new__(InformeMargenModule)
        helper.logger = self.logger
        usd_com = helper.clp_usd_segment_mask(
            work,
            InformeMargenModule._CLP_USD_ACREEDORES_COM,
        )
        usd_no_com = helper.clp_usd_segment_mask(
            work,
            InformeMargenModule._CLP_USD_ACREEDORES_NO_COM,
        )

        rename_months = dict(
            zip(self.MONTH_LABELS, self.TD_MONTH_LABELS, strict=True)
        )

        segments = (
            (self.TD_SHEET_NO_COM, ~com_mask & ~usd_no_com),
            (self.TD_SHEET_COM, com_mask & ~usd_com),
            (self.TD_SHEET_CLP_USD_NO_COM, ~com_mask & usd_no_com),
            (self.TD_SHEET_CLP_USD_COM, com_mask & usd_com),
        )
        for sheet_name, mask in segments:
            pivot = self.build_pivot(work.loc[mask])
            if pivot.empty:
                continue
            dynamics[sheet_name] = pivot.rename(columns=rename_months)

        return dynamics

    def build_kpis(
        self,
        pivot: pd.DataFrame,
        duration_s: float,
        generated_at: datetime,
    ) -> pd.DataFrame:
        """Tabla Indicador|Valor lista para Power BI."""

        load = self._last_load
        total_rows = load.total_rows if load else 0
        classified = load.classified_rows if load else 0
        usd_rows = load.usd_rows if load else 0
        no_usd_rows = load.no_usd_rows if load else 0
        coverage = (
            round((classified / total_rows) * 100, 1) if total_rows else 0.0
        )

        rows = [
            ("Fecha ejecución", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Duración proceso", round(duration_s, 3)),
            ("Registros MATRIZ", total_rows),
            ("Conceptos detectados", usd_rows + no_usd_rows),
            ("Conceptos distintos", int(len(pivot))),
            ("USD", usd_rows),
            ("NO_USD", no_usd_rows),
            ("Cobertura clasificación", coverage),
        ]
        return pd.DataFrame(rows, columns=["Indicador", "Valor"])

    def build_metadata(
        self,
        output_path: Path,
        duration_s: float,
        generated_at: datetime,
    ) -> pd.DataFrame:
        """Tabla de metadatos lista para Power BI."""

        source = self._last_load.source_file.name if self._last_load else "N/D"
        rows = [
            ("Archivo origen", source),
            ("Fecha generación", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Versión", self.MODEL_VERSION),
            ("Tiempo generación", round(duration_s, 3)),
            ("Usuario", self._current_user()),
            ("Pipeline version", VERSION),
            ("Archivo generado", output_path.name),
        ]
        return pd.DataFrame(rows, columns=["Campo", "Valor"])

    def _build_output_path(self) -> Path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = self.config.output_dir
        output_dir.mkdir(parents=True, exist_ok=True)
        return output_dir / f"PIVOT_MARGEN_{timestamp}.xlsx"

    def _style_sheet(
        self,
        worksheet,
        dataframe: pd.DataFrame,
        *,
        numeric_value_col: int | None = None,
        numeric_from_col: int | None = None,
        freeze: str = "B2",
    ) -> None:
        header_font = Font(bold=True)
        thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        center = Alignment(horizontal="center", vertical="center")
        right = Alignment(horizontal="right", vertical="center")
        vertical = Alignment(vertical="center")

        for cell in worksheet[1]:
            cell.font = header_font
            cell.alignment = center
            cell.border = thin

        max_row = worksheet.max_row
        for row in worksheet.iter_rows(min_row=2, max_row=max_row):
            for cell in row:
                cell.border = thin
                is_numeric_col = False
                if numeric_from_col is not None and cell.column >= numeric_from_col:
                    is_numeric_col = True
                if numeric_value_col is not None and cell.column == numeric_value_col:
                    is_numeric_col = isinstance(cell.value, (int, float)) and not isinstance(
                        cell.value, bool
                    )
                if is_numeric_col and isinstance(cell.value, (int, float)) and not isinstance(
                    cell.value, bool
                ):
                    cell.number_format = "#,##0.00"
                    cell.alignment = right
                else:
                    cell.alignment = vertical

        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.freeze_panes = freeze

        headers = list(dataframe.columns)
        for col_idx, name in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            max_length = len(str(name))
            for value in dataframe.iloc[:, col_idx - 1].head(50):
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[letter].width = min(max_length + 2, 40)

    def _write_dataframe(
        self,
        workbook: Workbook,
        title: str,
        dataframe: pd.DataFrame,
        *,
        numeric_from_col: int | None = None,
        numeric_value_col: int | None = None,
        replace_active: bool = False,
    ) -> None:
        if replace_active:
            worksheet = workbook.active
            worksheet.title = title
        else:
            worksheet = workbook.create_sheet(title=title)

        worksheet.append(list(dataframe.columns))
        for row_values in dataframe.itertuples(index=False, name=None):
            worksheet.append(list(row_values))

        self._style_sheet(
            worksheet,
            dataframe,
            numeric_from_col=numeric_from_col,
            numeric_value_col=numeric_value_col,
        )

    def export(
        self,
        pivot: pd.DataFrame,
        path: Path | None = None,
        kpis: pd.DataFrame | None = None,
        metadata: pd.DataFrame | None = None,
        td_dynamics: dict[str, pd.DataFrame] | None = None,
    ) -> Path:
        """
        Exportar modelo Power BI (BUSINESS-15/18A):

            PIVOT | TD_CLP_NO_COM | TD_CLP_COM_2026 |
            CLP_USD_NO_COM | CLP_USD_COM | KPIS | METADATA
        """

        output_path = path or self._build_output_path()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if kpis is None:
            kpis = pd.DataFrame(columns=["Indicador", "Valor"])
        if metadata is None:
            metadata = pd.DataFrame(columns=["Campo", "Valor"])
        if td_dynamics is None:
            td_dynamics = {}

        workbook = Workbook()
        self._write_dataframe(
            workbook,
            "PIVOT",
            pivot,
            numeric_from_col=2,
            replace_active=True,
        )
        for sheet_name in self.TD_SHEET_ORDER:
            frame = td_dynamics.get(sheet_name)
            if frame is None:
                frame = pd.DataFrame(
                    columns=["Concepto de Pago", *self.TD_MONTH_LABELS, "Total"]
                )
            self._write_dataframe(
                workbook,
                sheet_name,
                frame,
                numeric_from_col=2,
            )
        self._write_dataframe(
            workbook,
            "KPIS",
            kpis,
            numeric_value_col=2,
        )
        self._write_dataframe(
            workbook,
            "METADATA",
            metadata,
            numeric_value_col=2,
        )

        workbook.save(output_path)
        workbook.close()
        self.logger.info(
            "Modelo Power BI exportado: %s "
            "(hojas=PIVOT,%s,KPIS,METADATA)",
            output_path.name,
            ",".join(self.TD_SHEET_ORDER),
        )
        return output_path

    def generate(self, path: Path | None = None) -> tuple[pd.DataFrame, Path]:
        """Leer matriz → pivot → KPIS/METADATA → exportar Excel."""

        tracemalloc.start()
        total_started = time.perf_counter()
        generated_at = datetime.now()

        read_started = time.perf_counter()
        matrix = self.load_matrix(path)
        read_s = time.perf_counter() - read_started

        pivot_started = time.perf_counter()
        pivot = self.build_pivot(matrix)
        td_dynamics = self.build_td_dynamics(matrix)
        pivot_s = time.perf_counter() - pivot_started

        # Duración hasta antes de export (visible en KPIS/METADATA).
        pre_export_s = time.perf_counter() - total_started
        output_path = self._build_output_path()
        kpis = self.build_kpis(pivot, duration_s=pre_export_s, generated_at=generated_at)
        metadata = self.build_metadata(
            output_path,
            duration_s=pre_export_s,
            generated_at=generated_at,
        )

        export_started = time.perf_counter()
        output = self.export(
            pivot,
            path=output_path,
            kpis=kpis,
            metadata=metadata,
            td_dynamics=td_dynamics,
        )
        export_s = time.perf_counter() - export_started

        total_s = time.perf_counter() - total_started
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()

        self.last_metrics = {
            "read_s": read_s,
            "pivot_s": pivot_s,
            "export_s": export_s,
            "total_s": total_s,
            "peak_mb": peak / (1024 * 1024),
            "sheets": 7,
        }
        self.logger.info(
            "Métricas pivot: read=%.2fs pivot=%.2fs export=%.2fs total=%.2fs peak=%.1fMB sheets=%d",
            read_s,
            pivot_s,
            export_s,
            total_s,
            self.last_metrics["peak_mb"],
            7,
        )
        return pivot, output


__all__ = ("PivotReportService",)
