from __future__ import annotations

import errno
import hashlib
import json
import os
import uuid
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd

from ..config.config import CONCEPTOS_ADMIN_PATH, CONCEPTOS_PATH, Config
from ..config.logger import LoggerManager


"""
BUSINESS-12A — Sincronización automática CONCEPTOS_ADMIN → CONCEPTOS.

Negocio administra solo CONCEPTOS_ADMIN.xlsx.
Antes del proceso se actualiza CONCEPTOS.xlsx sin borrar registros legacy.
La escritura del destino es atómica: temporal validado + os.replace.
Si el contenido semántico no cambia, no se reescribe el archivo.
"""


logger = LoggerManager.get_logger(__name__)

ADMIN_REQUIRED_COLUMNS: tuple[str, ...] = (
    "Concepto encontrado",
    "Concepto estándar",
    "Grupo",
    "Tipo",
    "Activo",
)

CONCEPTOS_COLUMN = "Conceptos"
STANDARD_COLUMN = "Concepto estándar"
GROUP_COLUMN = "Grupo"
CONCEPTOS_SHEET_NAME = "CONCEPTOS"

OUTPUT_COLUMNS: tuple[str, ...] = (
    CONCEPTOS_COLUMN,
    STANDARD_COLUMN,
    GROUP_COLUMN,
)


class ConceptosConcurrentModificationError(RuntimeError):
    """El destino cambió, apareció o desapareció; el reemplazo fue cancelado."""


def _is_active(value: object) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return False
    text = str(value).strip().upper()
    return text in {"1", "SI", "SÍ", "TRUE", "YES", "Y", "ACTIVO"}


@dataclass(frozen=True)
class ConceptosSyncResult:
    """Resumen de la sincronización ADMIN → CONCEPTOS."""

    added: int
    updated: int
    unchanged: int
    legacy_preserved: int
    total: int
    admin_path: Path
    conceptos_path: Path
    changed: bool
    written: bool = False


def _empty_conceptos_frame() -> pd.DataFrame:
    return pd.DataFrame(columns=list(OUTPUT_COLUMNS))


def _ensure_output_columns(frame: pd.DataFrame) -> pd.DataFrame:
    """Garantizar columnas de salida sin perder columnas legacy extra."""

    work = frame.copy()
    if CONCEPTOS_COLUMN not in work.columns:
        # Compatibilidad: si solo hay otra columna de texto, no inventar datos.
        raise ValueError(
            f"CONCEPTOS.xlsx debe contener la columna '{CONCEPTOS_COLUMN}'"
        )

    for column in OUTPUT_COLUMNS:
        if column not in work.columns:
            work[column] = ""

    # Ordenar columnas canónicas primero; conservar extras al final.
    extras = [col for col in work.columns if col not in OUTPUT_COLUMNS]
    return work[list(OUTPUT_COLUMNS) + extras]


def _normalize_text_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Nulos → '', strip en todas las columnas; OUTPUT_COLUMNS como str."""

    work = _ensure_output_columns(frame)
    for column in work.columns:
        work[column] = work[column].fillna("").astype(str).str.strip()
    return work


def _semantic_records(frame: pd.DataFrame) -> list[tuple[str, ...]]:
    """Colección canónica ordenada: conserva multiplicidad; ignora solo el orden físico."""

    work = _normalize_text_frame(frame)
    columns = list(work.columns)
    rows = [
        tuple(str(work.iat[idx, col_idx]) for col_idx in range(len(columns)))
        for idx in range(len(work))
    ]
    return sorted(rows)


def _semantic_fingerprint(frame: pd.DataFrame) -> str:
    payload = json.dumps(
        _semantic_records(frame),
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _file_sha256(path: Path) -> str | None:
    if not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _read_destination_bytes(path: Path) -> bytes | None:
    """Leer el destino una sola vez.

    None = no existía al abrir. FileNotFoundError tras abrir = concurrente.
    """

    try:
        handle = path.open("rb")
    except FileNotFoundError:
        return None
    except OSError as exc:
        if exc.errno == errno.ENOENT:
            return None
        raise
    try:
        try:
            return handle.read()
        except FileNotFoundError as exc:
            raise ConceptosConcurrentModificationError(
                "CONCEPTOS.xlsx desapareció; reemplazo cancelado"
            ) from exc
        except OSError as exc:
            if exc.errno == errno.ENOENT:
                raise ConceptosConcurrentModificationError(
                    "CONCEPTOS.xlsx desapareció; reemplazo cancelado"
                ) from exc
            raise
    finally:
        handle.close()


@dataclass(frozen=True)
class _DestinationSnapshot:
    existed: bool
    sha256: str | None
    frame: pd.DataFrame
    schema_complete: bool


def _snapshot_conceptos_destination(path: Path) -> _DestinationSnapshot:
    """DataFrame y SHA256 salen de los mismos bytes."""

    payload = _read_destination_bytes(path)
    if payload is None:
        return _DestinationSnapshot(
            existed=False,
            sha256=None,
            frame=_empty_conceptos_frame(),
            schema_complete=False,
        )
    digest = _sha256_bytes(payload)
    raw = pd.read_excel(BytesIO(payload), sheet_name=0)
    schema_complete = all(column in raw.columns for column in OUTPUT_COLUMNS)
    return _DestinationSnapshot(
        existed=True,
        sha256=digest,
        frame=_ensure_output_columns(raw),
        schema_complete=schema_complete,
    )


def _assert_destination_unchanged(
    target: Path,
    expected_sha256: str | None,
) -> None:
    """Guarda inmediata previa a os.replace. No es un CAS nativo portable."""

    if expected_sha256 is None:
        if target.is_file():
            raise ConceptosConcurrentModificationError(
                "CONCEPTOS.xlsx apareció; reemplazo cancelado"
            )
        return
    if not target.is_file():
        raise ConceptosConcurrentModificationError(
            "CONCEPTOS.xlsx desapareció; reemplazo cancelado"
        )
    try:
        current_sha256 = _file_sha256(target)
    except FileNotFoundError as exc:
        raise ConceptosConcurrentModificationError(
            "CONCEPTOS.xlsx desapareció; reemplazo cancelado"
        ) from exc
    if current_sha256 is None:
        raise ConceptosConcurrentModificationError(
            "CONCEPTOS.xlsx desapareció; reemplazo cancelado"
        )
    if current_sha256 != expected_sha256:
        raise ConceptosConcurrentModificationError(
            "CONCEPTOS.xlsx cambió; reemplazo cancelado"
        )


def _temp_xlsx_path(target: Path) -> Path:
    token = uuid.uuid4().hex
    return target.parent / f"{target.stem}.{token}.tmp.xlsx"


def _unlink_known_temp(path: Path) -> None:
    """Eliminar únicamente el temporal conocido. Nunca el destino."""

    try:
        path.unlink(missing_ok=True)
    except OSError:
        logger.warning("No se pudo eliminar temporal CONCEPTOS %s", path.name)


def _validate_temp_workbook(temp_path: Path, expected: pd.DataFrame) -> None:
    if not zipfile.is_zipfile(temp_path):
        raise ValueError("temporal CONCEPTOS no es OOXML válido")

    workbook = openpyxl.load_workbook(temp_path, read_only=True, data_only=True)
    try:
        if CONCEPTOS_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"temporal CONCEPTOS sin hoja '{CONCEPTOS_SHEET_NAME}'"
            )
        sheet = workbook[CONCEPTOS_SHEET_NAME]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("temporal CONCEPTOS sin encabezados")
        headers = ["" if value is None else str(value) for value in rows[0]]
        expected_headers = [str(column) for column in expected.columns]
        if headers != expected_headers:
            raise ValueError("temporal CONCEPTOS con encabezados inesperados")
        data_rows = rows[1:]
        if len(data_rows) != len(expected):
            raise ValueError("temporal CONCEPTOS con número de filas inesperado")
    finally:
        workbook.close()

    loaded = pd.read_excel(temp_path, sheet_name=CONCEPTOS_SHEET_NAME)
    if _semantic_fingerprint(loaded) != _semantic_fingerprint(expected):
        raise ValueError("temporal CONCEPTOS no coincide con el fingerprint esperado")


def _fsync_closed_file(path: Path) -> None:
    """Flush a disco si la plataforma lo permite. En Windows O_RDONLY no admite fsync."""

    flags = os.O_RDWR
    if getattr(os, "O_BINARY", 0):
        flags |= os.O_BINARY
    try:
        fd = os.open(path, flags)
    except OSError:
        logger.debug("fsync omitido (no se pudo abrir temporal CONCEPTOS)")
        return
    try:
        os.fsync(fd)
    except OSError:
        logger.debug("fsync no soportado para temporal CONCEPTOS")
    finally:
        os.close(fd)


def _atomic_write_conceptos(
    frame: pd.DataFrame,
    target: Path,
    expected_sha256: str | None,
) -> None:
    """Escribir CONCEPTOS.xlsx vía temporal validado + os.replace.

    Existe una ventana mínima entre la guarda SHA256 y os.replace que
    Python no convierte en un CAS nativo portable. Todos los callers
    internos pasan por esta función.
    """

    target.parent.mkdir(parents=True, exist_ok=True)
    temp_path = _temp_xlsx_path(target)
    replaced = False
    try:
        frame.to_excel(temp_path, sheet_name=CONCEPTOS_SHEET_NAME, index=False)
        _validate_temp_workbook(temp_path, frame)
        _fsync_closed_file(temp_path)
        _assert_destination_unchanged(target, expected_sha256)
        os.replace(temp_path, target)
        replaced = True
    finally:
        if not replaced:
            _unlink_known_temp(temp_path)


def _load_admin_active_rows(admin_path: Path) -> pd.DataFrame:
    """Leer filas activas de CONCEPTOS_ADMIN.xlsx."""

    if not admin_path.exists():
        raise FileNotFoundError(
            f"CONCEPTOS_ADMIN.xlsx no encontrado: {admin_path}"
        )

    frame = pd.read_excel(admin_path, sheet_name=0)
    missing = [col for col in ADMIN_REQUIRED_COLUMNS if col not in frame.columns]
    if missing:
        raise ValueError(
            "CONCEPTOS_ADMIN.xlsx incompleto. Faltan columnas: "
            + ", ".join(missing)
        )

    rows: list[dict[str, str]] = []
    for _, item in frame.iterrows():
        if not _is_active(item.get("Activo")):
            continue
        encontrado = str(item.get("Concepto encontrado", "") or "").strip()
        estandar = str(item.get(STANDARD_COLUMN, "") or "").strip()
        grupo = str(item.get(GROUP_COLUMN, "") or "").strip()
        if not encontrado or not estandar:
            continue
        rows.append(
            {
                CONCEPTOS_COLUMN: encontrado,
                STANDARD_COLUMN: estandar,
                GROUP_COLUMN: grupo,
            }
        )

    return pd.DataFrame(rows, columns=list(OUTPUT_COLUMNS))


def _load_conceptos(conceptos_path: Path) -> pd.DataFrame:
    """Cargar CONCEPTOS.xlsx o crear estructura vacía."""

    if not conceptos_path.exists():
        logger.warning(
            "CONCEPTOS.xlsx no existe en %s — se creará desde ADMIN",
            conceptos_path,
        )
        return _empty_conceptos_frame()

    frame = pd.read_excel(conceptos_path, sheet_name=0)
    return _ensure_output_columns(frame)


def sync_conceptos_from_admin(
    admin_path: Path | None = None,
    conceptos_path: Path | None = None,
) -> ConceptosSyncResult:
    """
    Sincronizar CONCEPTOS.xlsx desde CONCEPTOS_ADMIN.xlsx.

    Reglas:
        - agregar conceptos nuevos (Concepto encontrado)
        - actualizar Concepto estándar y Grupo si ya existen
        - no eliminar registros legacy ausentes en ADMIN
        - no reescribir el destino si el contenido semántico no cambia
        - si hay que escribir, temporal validado + guarda concurrente + os.replace
    """

    config = Config()
    source_admin = Path(admin_path or config.conceptos_admin_path or CONCEPTOS_ADMIN_PATH)
    target = Path(conceptos_path or config.conceptos_path or CONCEPTOS_PATH)

    admin_rows = _load_admin_active_rows(source_admin)
    snapshot = _snapshot_conceptos_destination(target)
    destination_exists = snapshot.existed
    current = snapshot.frame
    previous_sha256 = snapshot.sha256
    if not destination_exists:
        logger.warning(
            "CONCEPTOS.xlsx no existe en %s — se creará desde ADMIN",
            target.name,
        )
    existing_fingerprint = (
        _semantic_fingerprint(current) if destination_exists else None
    )

    # Índice por Conceptos (concepto encontrado / legacy).
    index_by_concept: dict[str, int] = {}
    for idx, value in enumerate(current[CONCEPTOS_COLUMN].tolist()):
        key = "" if pd.isna(value) else str(value).strip()
        if key and key not in index_by_concept:
            index_by_concept[key] = idx

    added = 0
    updated = 0
    unchanged = 0
    new_rows: list[dict[str, object]] = []

    for _, admin in admin_rows.iterrows():
        key = str(admin[CONCEPTOS_COLUMN]).strip()
        standard = str(admin[STANDARD_COLUMN]).strip()
        group = str(admin[GROUP_COLUMN]).strip()

        if key in index_by_concept:
            row_idx = index_by_concept[key]
            prev_standard = str(current.at[row_idx, STANDARD_COLUMN] or "").strip()
            prev_group = str(current.at[row_idx, GROUP_COLUMN] or "").strip()
            if prev_standard != standard or prev_group != group:
                current.at[row_idx, STANDARD_COLUMN] = standard
                current.at[row_idx, GROUP_COLUMN] = group
                updated += 1
            else:
                unchanged += 1
            continue

        new_rows.append(
            {
                CONCEPTOS_COLUMN: key,
                STANDARD_COLUMN: standard,
                GROUP_COLUMN: group,
            }
        )
        added += 1

    if new_rows:
        current = pd.concat(
            [current, pd.DataFrame(new_rows)],
            ignore_index=True,
            sort=False,
        )

    current = _ensure_output_columns(current)

    # Limpiar NaN en columnas de texto.
    for column in OUTPUT_COLUMNS:
        current[column] = current[column].fillna("").astype(str)

    admin_keys = set(admin_rows[CONCEPTOS_COLUMN].astype(str).str.strip())
    legacy_preserved = 0
    for value in current[CONCEPTOS_COLUMN].tolist():
        key = str(value).strip()
        if key and key not in admin_keys:
            legacy_preserved += 1

    changed = added > 0 or updated > 0
    final_fingerprint = _semantic_fingerprint(current)
    schema_complete = snapshot.schema_complete
    skip_write = (
        destination_exists
        and schema_complete
        and existing_fingerprint == final_fingerprint
    )

    written = False
    if skip_write:
        logger.info(
            "Sincronización CONCEPTOS: escritura omitida "
            "(sin cambio semántico) added=%d updated=%d unchanged=%d "
            "legacy=%d total=%d fp=%s sha256=%s → %s",
            added,
            updated,
            unchanged,
            legacy_preserved,
            len(current),
            final_fingerprint,
            previous_sha256,
            target.name,
        )
    else:
        _atomic_write_conceptos(current, target, previous_sha256)
        written = True
        after_sha256 = _file_sha256(target)
        logger.info(
            "Sincronización CONCEPTOS: escritura atómica "
            "added=%d updated=%d unchanged=%d legacy=%d total=%d "
            "fp_antes=%s fp_despues=%s sha256_antes=%s sha256_despues=%s "
            "validacion=ok → %s",
            added,
            updated,
            unchanged,
            legacy_preserved,
            len(current),
            existing_fingerprint,
            final_fingerprint,
            previous_sha256,
            after_sha256,
            target.name,
        )

    return ConceptosSyncResult(
        added=added,
        updated=updated,
        unchanged=unchanged,
        legacy_preserved=legacy_preserved,
        total=len(current),
        admin_path=source_admin,
        conceptos_path=target,
        changed=changed,
        written=written,
    )


__all__ = (
    "ConceptosConcurrentModificationError",
    "ConceptosSyncResult",
    "sync_conceptos_from_admin",
)
