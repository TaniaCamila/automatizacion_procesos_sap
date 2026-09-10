from __future__ import annotations

import os
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from ..config.config import Config
from ..config.logger import LoggerManager

MATRIX_SHEET_NAME = "MATRIZ_FBL1N"
_MATRIX_PROGRESS_EVERY = 10_000


def _rss_mb() -> float | None:
    """RSS del proceso en MB. Best-effort; None si no se puede medir."""

    try:
        import ctypes
        from ctypes import wintypes

        class PROCESS_MEMORY_COUNTERS(ctypes.Structure):
            _fields_ = [
                ("cb", wintypes.DWORD),
                ("PageFaultCount", wintypes.DWORD),
                ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t),
                ("QuotaPeakPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPagedPoolUsage", ctypes.c_size_t),
                ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t),
                ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t),
            ]

        counters = PROCESS_MEMORY_COUNTERS()
        counters.cb = ctypes.sizeof(PROCESS_MEMORY_COUNTERS)
        handle = ctypes.windll.kernel32.GetCurrentProcess()
        get_info = ctypes.windll.psapi.GetProcessMemoryInfo
        get_info.argtypes = [
            wintypes.HANDLE,
            ctypes.POINTER(PROCESS_MEMORY_COUNTERS),
            wintypes.DWORD,
        ]
        get_info.restype = wintypes.BOOL
        ok = get_info(handle, ctypes.byref(counters), counters.cb)
        if not ok:
            return None
        return round(counters.WorkingSetSize / (1024 * 1024), 1)
    except Exception:
        return None


class ExcelReport:
    """Exportador profesional de DataFrames a Excel usando openpyxl."""

    def __init__(self, config_obj: Config | None = None) -> None:
        self.config = config_obj or Config()
        self.logger = LoggerManager.get_logger(__name__)

    def _emit(self, message: str, *args: object) -> None:
        """Registrar y forzar flush para que el progreso sea observable."""

        self.logger.info(message, *args)
        for handler in self.logger.handlers:
            try:
                handler.flush()
            except Exception:
                pass

    def _phase_start(self, name: str) -> float:
        self._emit(
            "fase inicio %s pid=%s rss_mb=%s",
            name,
            os.getpid(),
            _rss_mb(),
        )
        return time.perf_counter()

    def _phase_end(self, name: str, started: float) -> None:
        self._emit(
            "fase fin %s dur=%.3fs pid=%s rss_mb=%s",
            name,
            time.perf_counter() - started,
            os.getpid(),
            _rss_mb(),
        )

    def _build_output_path(self) -> Path:
        """Crear la ruta del archivo Excel con timestamp."""

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return self.config.output_dir / f"MATRIZ_FBL1N_{timestamp}.xlsx"

    def _apply_styles(self, worksheet) -> None:
        """Aplicar formato visual profesional a la hoja."""

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    def _apply_column_widths(self, worksheet) -> None:
        """Calcular anchos de columna a partir de los valores escritos."""

        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def _apply_data_types(self, worksheet, dataframe: pd.DataFrame) -> None:
        """Aplicar formato numérico y de fecha sin modificar valores."""

        for column in dataframe.columns:
            if dataframe[column].dtype.kind in "iufc":
                for row in worksheet.iter_rows(
                    min_row=2,
                    min_col=worksheet.column_dimensions.get(column, None) and 1,
                    max_col=worksheet.column_dimensions.get(column, None) and 1,
                    max_row=worksheet.max_row,
                ):
                    pass

        for index, column_name in enumerate(dataframe.columns, start=1):
            series = dataframe[column_name]
            if pd.api.types.is_numeric_dtype(series):
                for row in worksheet.iter_rows(
                    min_row=2,
                    max_row=worksheet.max_row,
                    min_col=index,
                    max_col=index,
                ):
                    for cell in row:
                        cell.number_format = "#,##0.00"
            elif pd.api.types.is_datetime64_any_dtype(series):
                for row in worksheet.iter_rows(
                    min_row=2,
                    max_row=worksheet.max_row,
                    min_col=index,
                    max_col=index,
                ):
                    for cell in row:
                        cell.number_format = "yyyy-mm-dd"

    def _append_dataframe(
        self,
        worksheet,
        dataframe: pd.DataFrame,
        *,
        sheet_name: str,
        progress_every: int | None = None,
    ) -> None:
        total = int(len(dataframe))
        cols = int(len(dataframe.columns))
        self._emit(
            "hoja=%s filas=%d columnas=%d",
            sheet_name,
            total,
            cols,
        )
        started = self._phase_start(f"append {sheet_name}")
        worksheet.append(list(dataframe.columns))
        for index, row in enumerate(dataframe.itertuples(index=False, name=None), start=1):
            worksheet.append(list(row))
            if progress_every and index % progress_every == 0:
                self._emit(
                    "%s progreso filas=%d/%d",
                    sheet_name,
                    index,
                    total,
                )
        self._phase_end(f"append {sheet_name}", started)

        started = self._phase_start(f"estilos {sheet_name}")
        self._apply_styles(worksheet)
        self._phase_end(f"estilos {sheet_name}", started)

        started = self._phase_start(f"anchos {sheet_name}")
        self._apply_column_widths(worksheet)
        self._phase_end(f"anchos {sheet_name}", started)

        started = self._phase_start(f"tipos {sheet_name}")
        self._apply_data_types(worksheet, dataframe)
        self._phase_end(f"tipos {sheet_name}", started)

    def _expected_sheet_names(
        self,
        additional_sheets: dict[str, pd.DataFrame] | None,
    ) -> list[str]:
        names = [MATRIX_SHEET_NAME]
        if additional_sheets:
            for sheet_name, sheet_df in additional_sheets.items():
                if isinstance(sheet_df, pd.DataFrame):
                    names.append(sheet_name)
        return names

    def _unlink_known_temp(self, path: Path | None) -> None:
        if path is None or not path.exists():
            return
        try:
            path.unlink()
        except OSError:
            self._emit("No se pudo eliminar temporal MATRIZ %s", path.name)

    def _validate_exported_workbook(
        self,
        path: Path,
        expected_names: list[str],
    ) -> None:
        if not path.is_file() or path.stat().st_size <= 0:
            raise ValueError(f"workbook exportado inexistente o vacio: {path}")
        if not zipfile.is_zipfile(path):
            raise ValueError(f"workbook exportado no es OOXML valido: {path}")

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            actual = list(workbook.sheetnames)
            if actual != expected_names:
                raise ValueError(
                    f"hojas inesperadas: {actual} != {expected_names}"
                )
            if MATRIX_SHEET_NAME not in actual:
                raise ValueError(f"falta hoja {MATRIX_SHEET_NAME}")
        finally:
            workbook.close()

    def export(
        self,
        dataframe: pd.DataFrame,
        additional_sheets: dict[str, pd.DataFrame] | None = None,
    ) -> Path:
        """
        Exportar DataFrames a un archivo Excel profesional.

        Args:
            dataframe: DataFrame principal (será la hoja MATRIZ_FBL1N).
            additional_sheets: Diccionario opcional con {nombre_hoja: DataFrame}.
                              Si None, solo exporta la matriz.

        Returns:
            Ruta del archivo Excel generado.

        Raises:
            ValueError: Si dataframe es None.
            TypeError: Si dataframe no es un DataFrame.
        """

        if dataframe is None:
            raise ValueError("El DataFrame proporcionado es None.")

        if not isinstance(dataframe, pd.DataFrame):
            raise TypeError("El objeto proporcionado no es un DataFrame.")

        workbook: Workbook | None = None
        temp_path: Path | None = None
        phase = "inicio"
        self._emit(
            "export pid=%s rss_mb=%s",
            os.getpid(),
            _rss_mb(),
        )

        try:
            phase = "construccion"
            started = self._phase_start("construccion workbook")
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = MATRIX_SHEET_NAME
            self._append_dataframe(
                worksheet,
                dataframe,
                sheet_name=MATRIX_SHEET_NAME,
                progress_every=_MATRIX_PROGRESS_EVERY,
            )

            if additional_sheets:
                for sheet_name, sheet_df in additional_sheets.items():
                    if not isinstance(sheet_df, pd.DataFrame):
                        continue
                    new_worksheet = workbook.create_sheet(title=sheet_name)
                    self._append_dataframe(
                        new_worksheet,
                        sheet_df,
                        sheet_name=sheet_name,
                    )
            self._phase_end("construccion workbook", started)

            expected_names = self._expected_sheet_names(additional_sheets)
            output_path = self._build_output_path()
            if output_path.exists():
                raise FileExistsError(
                    f"destino MATRIZ ya existe; no se modifica: {output_path}"
                )

            timestamp = output_path.stem.removeprefix("MATRIZ_FBL1N_")
            temp_path = (
                output_path.parent
                / f".MATRIZ_FBL1N_{timestamp}.{uuid.uuid4().hex}.tmp.xlsx"
            )

            phase = "save"
            started = self._phase_start("workbook.save")
            workbook.save(temp_path)
            self._phase_end("workbook.save", started)

            phase = "cierre"
            workbook.close()
            workbook = None

            phase = "validacion"
            started = self._phase_start("validacion OOXML")
            self._validate_exported_workbook(temp_path, expected_names)
            self._phase_end("validacion OOXML", started)

            phase = "replace"
            started = self._phase_start("os.replace")
            os.replace(temp_path, output_path)
            temp_path = None
            self._phase_end("os.replace", started)

            phase = "confirmacion"
            self._validate_exported_workbook(output_path, expected_names)
            self._emit("archivo generado: %s", output_path)
            return output_path
        except Exception:
            self._emit("export abortado en fase=%s pid=%s", phase, os.getpid())
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            self._unlink_known_temp(temp_path)
            raise
