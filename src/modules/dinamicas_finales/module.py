"""
Módulo EF-11.0 / EF-12.2 / EF-12.3 / EF-12.5 / EF-12.6 / EF-13.2 / EF-13.7
/ EF-14.0 / EF-14.1 / EF-14.2 / EF-14.4 / EF-14.7 / EF-14.8 / EF-14.9
— Dinámicas finales.

Construye el entregable final con las hojas:
    COMB_CLP, NO_COMB_CLP, COMB_USD, NO_COMB_USD, CLP_USD,
    COMB_GNL CHILE, CONFIRMING_USD, PENDIENTES_USD, NO_ECM

Fuente única: MATRIZ_FBL1N (artefacto existente, solo lectura).

EF-12.6: dinámica exclusiva COMB_GNL CHILE para acreedor 4000000387
(todos los registros con nomenclatura), retirado de las demás dinámicas.

EF-13.7: aplica la lógica validada con GASODUCTO GASANDES a todos los
proveedores del universo CLP_USD (COM + NO_COM). Match único
doccompensacion ↔ PROVISIÓN CONTABLE; MONTO USD con signo CLP.
Sin match / match múltiple → hoja PENDIENTES_USD (no se descartan).

EF-14.2: registros sin nomenclatura válida (concepto vacío o fuera del
catálogo CONCEPTOS) se retiran de las dinámicas y se copian a NO_ECM
sin reclasificar ni alterar montos. MATRIZ no se modifica.

EF-14.4/14.8: DATOS_* de pivots jerárquicos (incl. NO_COMB_CLP) son
detalle documental (una fila = una línea MATRIZ). PENDIENTES_USD y
NO_ECM siguen estáticas.

EF-14.7: antes de SaveAs, PivotCache.Refresh() en dinámicas jerárquicas
para reparar el índice de ShowDetail sin alterar DATOS_* ni totales.

EF-14.8: NO_COMB_CLP pasa a PivotTable documental (DATOS_NO_COMB_CLP),
mismo comportamiento que COMB_CLP / USD / GNL.

EF-14.9: CONFIRMING_USD — Pivot documental de acreedores confirming
con estado SIN_MATCH vs matriz de pago ME. No altera matching ni
demás hojas (PENDIENTES_USD se conserva completa).

NO modifica pipeline, MATRIZ, catálogos, concept_classifier ni reglas BUSINESS.
"""

from __future__ import annotations

import logging
import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from ...config.config import CONCEPTOS_PATH, INPUT_DIR, OUTPUT_DIR
from ...services.output_artifact_service import OutputArtifactService
from ..margen.module import InformeMargenModule

logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_DIR = OUTPUT_DIR

# EF-13.7 — matriz pago ME (solo lectura; misma fuente validada en EF-13.6).
PAGO_ME_PATH = INPUT_DIR / "EJM_ARCHIVO_PAGO_MONEDA_EXTRANJERA.xlsx"

# Entregable estable para Power BI (copia validada del timestamp).
ACTUAL_FILENAME = "DINAMICAS_FINALES_ACTUAL.xlsx"

# Excel COM constants (mismo patrón que pivots_tesoreria).
XL_DATABASE = 1
XL_ROW_FIELD = 1
XL_COLUMN_FIELD = 2
XL_DATA_FIELD = 4
XL_SUM = -4157
XL_OPEN_XML_WORKBOOK = 51
XL_SORT_MANUAL = 0

SHEET_MAP: dict[str, str] = {
    "COMB_CLP": InformeMargenModule._TD_SHEET_COM,
    "NO_COMB_CLP": InformeMargenModule._TD_SHEET_NO_COM,
    "COMB_USD": InformeMargenModule._TD_SHEET_CLP_USD_COM,
    "NO_COMB_USD": InformeMargenModule._TD_SHEET_CLP_USD_NO_COM,
}

# EF-12.6 — acreedor exclusivo COMB_GNL CHILE
EXCLUSIVE_ACREEDOR = "4000000387"
EXCLUSIVE_PROVEEDOR = "GNL CHILE SA"
EXCLUSIVE_SHEET = "COMB_GNL CHILE"

# Compatibilidad con scripts previos (piloto Gasandes).
GASANDES_ACREEDOR = "2000168164"
GASANDES_PROVEEDOR = "GASODUCTO GASANDES S.A."

PENDIENTES_SHEET = "PENDIENTES_USD"
NO_ECM_SHEET = "NO_ECM"

# EF-14.9 — Confirming SIN_MATCH (subconjunto de PENDIENTES_USD).
CONFIRMING_SHEET = "CONFIRMING_USD"
CONFIRMING_ACREEDORES: frozenset[str] = frozenset(
    {
        "2000326095",
        "2000706717",
        "2000728821",
        "2000680638",
        "2000736433",
        "2000129154",
        "2000253935",
    }
)

# Universo de proveedores pagados en USD (listado del proyecto).
USD_ACREEDORES_COM = InformeMargenModule._CLP_USD_ACREEDORES_COM
USD_ACREEDORES_NO_COM = InformeMargenModule._CLP_USD_ACREEDORES_NO_COM
USD_ACREEDORES = frozenset(USD_ACREEDORES_COM | USD_ACREEDORES_NO_COM)

# Columnas de detalle para NO_ECM (valores originales de MATRIZ).
NO_ECM_COLUMN_CANDIDATES: tuple[tuple[str, ...], ...] = (
    ("Sociedad", "sociedad"),
    ("Acreedor", "acreedor"),
    ("Name", "name"),
    ("no_documento",),
    ("doccompensacion",),
    ("referencia",),
    ("Texto cab.documento", "texto_cabdocumento"),
    ("moneda_del_documento", "Moneda del documento"),
    ("importe_en_moneda_doc", "Importe en moneda doc."),
    ("Fecha compensación", "fecha_compensacion"),
    ("mes_compensacion",),
    ("concepto_detectado",),
    ("concepto_estado",),
    ("tipo",),
    ("currency_group",),
)

# Última auditoría EF-13.7 / EF-14.2 / EF-14.9 (en memoria; no se escribe en MATRIZ).
_LAST_EF137_AUDIT: dict[str, Any] | None = None
_LAST_EF142_AUDIT: dict[str, Any] | None = None
_LAST_EF149_AUDIT: dict[str, Any] | None = None
# Alias legacy.
_LAST_EF132_AUDIT: dict[str, Any] | None = None

FINAL_SHEET_ORDER: tuple[str, ...] = (
    "COMB_CLP",
    "NO_COMB_CLP",
    "COMB_USD",
    "NO_COMB_USD",
    "CLP_USD",
    EXCLUSIVE_SHEET,
    CONFIRMING_SHEET,
    PENDIENTES_SHEET,
    NO_ECM_SHEET,
)

HIERARCHICAL_SHEETS: tuple[str, ...] = (
    "COMB_CLP",
    "NO_COMB_CLP",
    "COMB_USD",
    "NO_COMB_USD",
    "CLP_USD",
    EXCLUSIVE_SHEET,
    CONFIRMING_SHEET,
)

# CONFIRMING_USD es pivot de montos CLP (nombre histórico).
CLP_PIVOT_SHEETS: frozenset[str] = frozenset(
    {
        "COMB_CLP",
        "NO_COMB_CLP",
        EXCLUSIVE_SHEET,
        CONFIRMING_SHEET,
    }
)
USD_PIVOT_SHEETS: frozenset[str] = frozenset(
    {
        "COMB_USD",
        "NO_COMB_USD",
        "CLP_USD",
    }
)

# Excel es-CL (miles='.' decimal=','). El invariante US '#,##0.00'
# se persiste corrupto como '#.##000'. Usar NumberFormatLocal.
CLP_NUMBER_FORMAT_LOCAL = "#.##0"
USD_NUMBER_FORMAT_LOCAL = "#.##0,00"

# Presentación (no altera valores ni NumberFormatLocal).
PIVOT_HEADER_COLOR_HEX = "#002060"
PIVOT_HEADER_FONT_COLOR = 16777215  # blanco, VBA RGB(255,255,255)
XL_ALIGN_CENTER = -4108
XL_ALIGN_LEFT = -4131
XL_ALIGN_RIGHT = -4152
XL_PIVOT_LINE_GRAND_TOTAL = 2
ROW_LABEL_COL_WIDTH = 24.0
ROW_LABEL_COL_WIDTH_MAX = 32.0
AMOUNT_COL_WIDTH_MIN = 18.0
AMOUNT_COL_WIDTH_MAX = 20.0
STATIC_TEXT_COL_WIDTH_MIN = 12.0
STATIC_TEXT_COL_WIDTH_MAX = 28.0
_GRAND_TOTAL_LABELS = frozenset({"total general", "grand total"})
_PIVOT_TOTAL_ROW_LABELS = frozenset({"Monto Bruto", "Total general", "Grand Total"})

# 9 dinámicas/estáticas + 7 DATOS_* (nombres exactos, incl. espacio en GNL).
EXPECTED_DINAMICAS_SHEETS: frozenset[str] = frozenset(
    {
        *FINAL_SHEET_ORDER,
        *[f"DATOS_{name}"[:31] for name in HIERARCHICAL_SHEETS],
    }
)

MONTH_LABELS: tuple[str, ...] = InformeMargenModule._TD_MONTH_LABELS
CONCEPT_COL = InformeMargenModule._TD_CONCEPT_COLUMN
PROVEEDOR_COL = InformeMargenModule._TD_PROVEEDOR_COLUMN
IMPORTE_DINAMICA_COL = "Importe usado en la dinámica"

# EF-14.4 — columnas documentales mínimas en DATOS_* (detalle MATRIZ).
DOCUMENT_COLUMN_MAP: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("Sociedad", ("Sociedad", "sociedad")),
    ("Acreedor", ("Acreedor", "acreedor")),
    ("Name", ("Name", "name")),
    ("N° documento", ("Nº documento", "N° documento", "no_documento")),
    ("Clase de documento", ("Clase de documento", "clase_de_documento")),
    ("Referencia", ("Referencia", "referencia")),
    ("Fecha de documento", ("Fecha de documento", "fecha_de_documento")),
    ("Fecha de contabilización", ("Fecha contabiliz.", "fecha_contabiliz")),
    ("Fecha de compensación", ("Fecha compensación", "fecha_compensacion")),
    (
        "Documento de compensación",
        ("Doc.compensación", "doccompensacion"),
    ),
    ("Texto cab.documento", ("Texto cab.documento", "texto_cabdocumento")),
    (
        "Importe en moneda doc.",
        ("Importe en moneda doc.", "importe_en_moneda_doc"),
    ),
    (
        "Moneda del documento",
        ("Moneda del documento", "moneda_del_documento"),
    ),
)

_VALID_NOMENCLATURE_CACHE: frozenset[str] | None = None
_LAST_DOCUMENT_DATOS: dict[str, pd.DataFrame] | None = None
_LAST_EF144_AUDIT: dict[str, Any] | None = None


def load_matriz(path: Path) -> pd.DataFrame:
    """Leer hoja MATRIZ_FBL1N del artefacto (solo lectura)."""
    return pd.read_excel(path, sheet_name="MATRIZ_FBL1N", engine="calamine")


def load_valid_nomenclature_concepts() -> frozenset[str]:
    """
    Nomenclaturas definidas (solo lectura).

    Fuente: data/resources/CONCEPTOS.xlsx columna Conceptos.
    Complemento: conceptos estándar del classifier (sin modificar catálogos).
    """
    global _VALID_NOMENCLATURE_CACHE
    if _VALID_NOMENCLATURE_CACHE is not None:
        return _VALID_NOMENCLATURE_CACHE

    concepts: set[str] = set()
    if CONCEPTOS_PATH.exists():
        frame = pd.read_excel(CONCEPTOS_PATH, engine="calamine")
        columns = {str(c).strip().lower(): c for c in frame.columns}
        col = columns.get("conceptos") or columns.get("concepto")
        if col is not None:
            for value in frame[col].tolist():
                if value is None or (isinstance(value, float) and pd.isna(value)):
                    continue
                text = str(value).strip()
                if text:
                    concepts.add(text)
    try:
        from ...services import concept_classifier

        for value in concept_classifier.all_standard_concepts():
            text = str(value).strip()
            if text:
                concepts.add(text)
    except Exception:
        pass

    _VALID_NOMENCLATURE_CACHE = frozenset(concepts)
    return _VALID_NOMENCLATURE_CACHE


def outside_nomenclature_mask(matrix: pd.DataFrame) -> pd.Series:
    """True si concepto vacío o no pertenece a nomenclatura válida."""
    if matrix.empty or "concepto_detectado" not in matrix.columns:
        return pd.Series(False, index=matrix.index)
    concept = matrix["concepto_detectado"].fillna("").astype(str).str.strip()
    valid = load_valid_nomenclature_concepts()
    return concept.eq("") | ~concept.isin(valid)


def _build_no_ecm_sheet(rows: pd.DataFrame) -> pd.DataFrame:
    """
    Copiar registros fuera de nomenclatura sin reclasificar ni alterar montos.
    """
    if rows is None or rows.empty:
        return pd.DataFrame()

    selected: list[str] = []
    for candidates in NO_ECM_COLUMN_CANDIDATES:
        for column in candidates:
            if column in rows.columns:
                selected.append(column)
                break

    if not selected:
        out = rows.copy().reset_index(drop=True)
    else:
        out = rows.loc[:, selected].copy().reset_index(drop=True)

    concept = (
        out["concepto_detectado"].fillna("").astype(str).str.strip()
        if "concepto_detectado" in out.columns
        else pd.Series("", index=out.index)
    )
    out["Motivo"] = concept.map(
        lambda c: "SIN_CONCEPTO" if not c else "FUERA_NOMENCLATURA"
    )
    return out


def _importe_sum(frame: pd.DataFrame) -> float:
    if frame is None or frame.empty or "importe_en_moneda_doc" not in frame.columns:
        return 0.0
    return float(
        pd.to_numeric(frame["importe_en_moneda_doc"], errors="coerce").fillna(0.0).sum()
    )


def _pick_column(frame: pd.DataFrame, candidates: tuple[str, ...]) -> str | None:
    for name in candidates:
        if name in frame.columns:
            return name
    return None


def _filter_dynamic_source(
    matrix: pd.DataFrame,
    *,
    tipo_filter: str | None,
    clp_usd: bool | None,
    clp_usd_acreedores: frozenset[str] | None,
    require_clp: bool = True,
) -> pd.DataFrame:
    """
    Misma segmentación que generate_pivot_clp (sin agregar).
    """
    if matrix is None or matrix.empty:
        return matrix.iloc[0:0].copy() if matrix is not None else pd.DataFrame()

    margen = InformeMargenModule()
    source, _ = margen.apply_clp_pivot_exclusions(matrix)
    if tipo_filter:
        source = margen._filter_by_matrix_tipo(source, tipo_filter)
    if clp_usd is not None:
        acreedores = clp_usd_acreedores or frozenset()
        mask = margen.clp_usd_segment_mask(source, acreedores)
        source = source.loc[mask].copy() if clp_usd else source.loc[~mask].copy()
    if require_clp and "moneda_del_documento" in source.columns:
        mon = (
            source["moneda_del_documento"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
        )
        source = source.loc[mon.eq("CLP")].copy()
    return source


def _build_document_detail(
    source: pd.DataFrame,
    *,
    amount_used: pd.Series | None = None,
    importe_original: pd.Series | None = None,
) -> pd.DataFrame:
    """
    EF-14.4: una fila por registro documental que alimenta la dinámica.
    """
    empty_cols = [
        *[label for label, _ in DOCUMENT_COLUMN_MAP],
        CONCEPT_COL,
        PROVEEDOR_COL,
        "Mes",
        IMPORTE_DINAMICA_COL,
    ]
    if source is None or source.empty:
        return pd.DataFrame(columns=empty_cols)

    work = source.copy()
    out = pd.DataFrame(index=work.index)
    for label, candidates in DOCUMENT_COLUMN_MAP:
        col = _pick_column(work, candidates)
        if col is None:
            out[label] = ""
        else:
            out[label] = work[col]

    # Importe en moneda doc. original (pre-override USD si se provee).
    if importe_original is not None:
        out["Importe en moneda doc."] = pd.to_numeric(
            importe_original.reindex(work.index), errors="coerce"
        )
    else:
        imp_col = _pick_column(
            work, ("Importe en moneda doc.", "importe_en_moneda_doc")
        )
        out["Importe en moneda doc."] = (
            pd.to_numeric(work[imp_col], errors="coerce")
            if imp_col is not None
            else 0.0
        )

    concept_col = _pick_column(work, ("concepto_detectado", "concepto"))
    out[CONCEPT_COL] = (
        work[concept_col].fillna("").astype(str).str.strip()
        if concept_col is not None
        else ""
    )
    name_col = _pick_column(work, ("Name", "name"))
    out[PROVEEDOR_COL] = (
        work[name_col].fillna("").astype(str).str.strip()
        if name_col is not None
        else ""
    )

    mes_col = _pick_column(work, ("mes_compensacion",))
    if mes_col is not None:
        out["Mes"] = work[mes_col].map(InformeMargenModule._month_label_from_period)
    else:
        out["Mes"] = None

    if amount_used is not None:
        out[IMPORTE_DINAMICA_COL] = pd.to_numeric(
            amount_used.reindex(work.index), errors="coerce"
        ).fillna(0.0)
    else:
        out[IMPORTE_DINAMICA_COL] = pd.to_numeric(
            out["Importe en moneda doc."], errors="coerce"
        ).fillna(0.0)

    # Solo filas con mes de dinámica (misma base que el pivot por mes).
    out = out.loc[out["Mes"].isin(MONTH_LABELS)].copy()
    out = out.loc[~out.index.duplicated(keep="first")].copy()
    return out.reset_index(drop=True)


def build_document_datos(
    *,
    matrix_rest: pd.DataFrame,
    matrix_nom: pd.DataFrame,
    matrix_usd_matched: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    """Detalle documental por dinámica jerárquica (EF-14.4)."""
    imp_rest = None
    if "importe_en_moneda_doc" in matrix_rest.columns:
        imp_rest = pd.to_numeric(matrix_rest["importe_en_moneda_doc"], errors="coerce")
    elif "Importe en moneda doc." in matrix_rest.columns:
        imp_rest = pd.to_numeric(
            matrix_rest["Importe en moneda doc."], errors="coerce"
        )

    src_comb_clp = _filter_dynamic_source(
        matrix_rest,
        tipo_filter="COM",
        clp_usd=False,
        clp_usd_acreedores=USD_ACREEDORES_COM,
        require_clp=True,
    )
    src_no_comb_clp = _filter_dynamic_source(
        matrix_rest,
        tipo_filter="NO_COM",
        clp_usd=False,
        clp_usd_acreedores=USD_ACREEDORES_NO_COM,
        require_clp=True,
    )
    src_comb_usd = _filter_dynamic_source(
        matrix_usd_matched,
        tipo_filter="COM",
        clp_usd=True,
        clp_usd_acreedores=USD_ACREEDORES_COM,
        require_clp=True,
    )
    src_no_comb_usd = _filter_dynamic_source(
        matrix_usd_matched,
        tipo_filter="NO_COM",
        clp_usd=True,
        clp_usd_acreedores=USD_ACREEDORES_NO_COM,
        require_clp=True,
    )
    src_clp_usd = pd.concat([src_comb_usd, src_no_comb_usd], axis=0)
    src_clp_usd = src_clp_usd.loc[~src_clp_usd.index.duplicated(keep="first")]

    gnl_mask = _acreedor_mask(matrix_nom, EXCLUSIVE_ACREEDOR)
    src_gnl = matrix_nom.loc[gnl_mask].copy()

    def _used(frame: pd.DataFrame) -> pd.Series | None:
        col = _pick_column(frame, ("importe_en_moneda_doc", "Importe en moneda doc."))
        if col is None:
            return None
        return pd.to_numeric(frame[col], errors="coerce")

    return {
        "COMB_CLP": _build_document_detail(src_comb_clp),
        "NO_COMB_CLP": _build_document_detail(src_no_comb_clp),
        "COMB_USD": _build_document_detail(
            src_comb_usd,
            amount_used=_used(src_comb_usd),
            importe_original=imp_rest,
        ),
        "NO_COMB_USD": _build_document_detail(
            src_no_comb_usd,
            amount_used=_used(src_no_comb_usd),
            importe_original=imp_rest,
        ),
        "CLP_USD": _build_document_detail(
            src_clp_usd,
            amount_used=_used(src_clp_usd),
            importe_original=imp_rest,
        ),
        EXCLUSIVE_SHEET: _build_document_detail(src_gnl),
    }


def _norm_acreedor(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    try:
        number = float(text.replace(",", ""))
        if number.is_integer():
            return str(int(number))
    except ValueError:
        pass
    return text


def _acreedor_mask(matrix: pd.DataFrame, acreedor: str) -> pd.Series:
    if "Acreedor" in matrix.columns:
        series = matrix["Acreedor"]
    elif "acreedor" in matrix.columns:
        series = matrix["acreedor"]
    else:
        raise KeyError("MATRIZ sin columna Acreedor/acreedor")
    return series.map(_norm_acreedor).eq(acreedor)


def _sheet_total(frame: pd.DataFrame) -> float:
    if frame is None or frame.empty or "Total" not in frame.columns:
        return 0.0
    work = frame.copy()
    if CONCEPT_COL in work.columns:
        concept = work[CONCEPT_COL].fillna("").astype(str).str.strip()
        # Excluir solo márgenes; conservar concepto vacío.
        work = work.loc[~concept.isin({"All", "all", "Total", "total"})]
    return float(pd.to_numeric(work["Total"], errors="coerce").fillna(0.0).sum())


def _build_comb_gnl_chile(matrix: pd.DataFrame) -> pd.DataFrame:
    """
    Dinámica exclusiva: todos los registros del acreedor 4000000387.

    Sin filtro de moneda, tipo ni BUSINESS-18.
    No inventa conceptos: vacío se conserva como "".
    """
    mask = _acreedor_mask(matrix, EXCLUSIVE_ACREEDOR)
    source = matrix.loc[mask].copy()
    if source.empty:
        return InformeMargenModule()._empty_td_frame(include_proveedor=True)

    margen = InformeMargenModule()
    source["concepto_detectado"] = (
        source["concepto_detectado"].fillna("").astype(str)
    )
    # Proveedor confirmado; reforzar Name para el pivot.
    if "Name" in source.columns:
        source["Name"] = source["Name"].fillna("").astype(str).str.strip()
        source.loc[source["Name"].eq(""), "Name"] = EXCLUSIVE_PROVEEDOR
    else:
        source["Name"] = EXCLUSIVE_PROVEEDOR

    pivot_df = margen.pivot.create_pivot(
        df=source,
        index=["concepto_detectado", "Name"],
        columns="mes_compensacion",
        values="importe_en_moneda_doc",
        aggfunc="sum",
        filters=None,
        margins=True,
        fill_value=0,
    )
    if pivot_df.empty:
        return margen._empty_td_frame(include_proveedor=True)

    pivot_df = pivot_df.reset_index()
    if "concepto_detectado" in pivot_df.columns:
        pivot_df = pivot_df.rename(
            columns={"concepto_detectado": CONCEPT_COL}
        )
    if "Name" in pivot_df.columns:
        pivot_df = pivot_df.rename(columns={"Name": PROVEEDOR_COL})

    return margen._format_td_clp_structure(
        pivot_df,
        include_proveedor=True,
    )


def _select_confirming_sin_match(
    matrix: pd.DataFrame,
    audit: dict[str, Any],
) -> pd.DataFrame:
    """
    EF-14.9: registros SIN_MATCH de acreedores confirming (montos CLP originales).

    Usa índices ya clasificados por apply_usd_pago_me; no recalcula matching.
    """
    empty = matrix.iloc[0:0].copy() if matrix is not None else pd.DataFrame()
    if matrix is None or matrix.empty:
        return empty
    idxs = set(audit.get("sin_match_index") or [])
    if not idxs:
        return empty
    source = matrix.loc[matrix.index.isin(idxs)].copy()
    if source.empty:
        return empty
    acre_col = "Acreedor" if "Acreedor" in source.columns else (
        "acreedor" if "acreedor" in source.columns else None
    )
    if acre_col is None:
        return empty
    mask = source[acre_col].map(_norm_acreedor).isin(CONFIRMING_ACREEDORES)
    return source.loc[mask].copy()


def _build_confirming_usd(source: pd.DataFrame) -> pd.DataFrame:
    """Pivot jerárquico CONFIRMING_USD (Concepto × Proveedor × Mes)."""
    if source is None or source.empty:
        return InformeMargenModule()._empty_td_frame(include_proveedor=True)

    margen = InformeMargenModule()
    work = source.copy()
    if "concepto_detectado" in work.columns:
        work["concepto_detectado"] = (
            work["concepto_detectado"].fillna("").astype(str)
        )
    if "Name" in work.columns:
        work["Name"] = work["Name"].fillna("").astype(str).str.strip()
    elif "name" in work.columns:
        work["Name"] = work["name"].fillna("").astype(str).str.strip()
    else:
        work["Name"] = ""

    pivot_df = margen.pivot.create_pivot(
        df=work,
        index=["concepto_detectado", "Name"],
        columns="mes_compensacion",
        values="importe_en_moneda_doc",
        aggfunc="sum",
        filters=None,
        margins=True,
        fill_value=0,
    )
    if pivot_df.empty:
        return margen._empty_td_frame(include_proveedor=True)

    pivot_df = pivot_df.reset_index()
    if "concepto_detectado" in pivot_df.columns:
        pivot_df = pivot_df.rename(columns={"concepto_detectado": CONCEPT_COL})
    if "Name" in pivot_df.columns:
        pivot_df = pivot_df.rename(columns={"Name": PROVEEDOR_COL})

    return margen._format_td_clp_structure(pivot_df, include_proveedor=True)


def _norm_doc(value: object) -> str:
    return _norm_acreedor(value)


def _load_pago_me(path: Path = PAGO_ME_PATH) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Matriz pago ME no encontrada: {path}")
    pago = pd.read_excel(path, sheet_name=0, engine="openpyxl")
    pago.columns = [str(c).strip() for c in pago.columns]
    prov_cols = [c for c in pago.columns if "PROVIS" in c.upper()]
    if not prov_cols:
        raise KeyError("Columna PROVISIÓN CONTABLE ausente en matriz de pago")
    if "ACREEDOR" not in pago.columns or "MONTO" not in pago.columns:
        raise KeyError("Columnas ACREEDOR/MONTO ausentes en matriz de pago")
    out = pago.copy()
    out["_acre"] = out["ACREEDOR"].map(_norm_acreedor)
    out["_prov"] = out[prov_cols[0]].map(_norm_doc)
    out["_monto"] = pd.to_numeric(out["MONTO"], errors="coerce")
    return out


def _usd_feed_mask(matrix: pd.DataFrame) -> pd.Series:
    """
    Registros CLP que alimentan dinámicas USD (universo del proyecto).

    Misma base que generate_pivot_clp(clp_usd=True): exclusiones BUSINESS-01
    + segmento BUSINESS-18 (COM ∪ NO_COM) + moneda CLP.
    """
    if matrix.empty:
        return pd.Series(dtype=bool)

    margen = InformeMargenModule()
    source, _ = margen.apply_clp_pivot_exclusions(matrix)
    mask_com = margen.clp_usd_segment_mask(source, USD_ACREEDORES_COM)
    mask_nocom = margen.clp_usd_segment_mask(source, USD_ACREEDORES_NO_COM)
    segment = pd.Series(False, index=matrix.index)
    segment.loc[source.index] = mask_com | mask_nocom
    mon = (
        matrix["moneda_del_documento"].fillna("").astype(str).str.strip().str.upper()
        if "moneda_del_documento" in matrix.columns
        else pd.Series("", index=matrix.index)
    )
    return segment.fillna(False) & mon.eq("CLP")


def apply_usd_pago_me(
    matrix: pd.DataFrame,
    *,
    pago_path: Path = PAGO_ME_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    """
    EF-13.7: copia de MATRIZ con importe USD en matches únicos.

    Lógica validada con Gasandes, extendida a todos los acreedores USD:
    - fila principal = registro CLP del segmento USD;
    - llave = doccompensacion = PROVISIÓN CONTABLE (por acreedor);
    - match único → MONTO USD con signo del CLP;
    - sin match / múltiple → pendiente (sin inventar monto);
    - no duplica provisiones ni montos.
    No escribe la MATRIZ en disco.
    """
    work = matrix.copy()
    pago = _load_pago_me(pago_path)
    pago_usd = pago.loc[pago["_acre"].isin(USD_ACREEDORES)].copy()

    if "doccompensacion" not in work.columns:
        raise KeyError("MATRIZ sin doccompensacion")
    if "importe_en_moneda_doc" not in work.columns:
        raise KeyError("MATRIZ sin importe_en_moneda_doc")

    mask = _usd_feed_mask(work)
    candidates = work.loc[mask].copy()
    if candidates.empty:
        empty_pend = pd.DataFrame(
            columns=[
                "Sociedad",
                "Acreedor",
                "Proveedor",
                "N° documento CLP",
                "doccompensacion",
                "Referencia",
                "Concepto",
                "Moneda",
                "Monto CLP original",
                "Motivo",
                "cantidad de coincidencias encontradas",
            ]
        )
        audit: dict[str, Any] = {
            "procesados": 0,
            "match_unico": 0,
            "sin_match": 0,
            "multi": 0,
            "total_usd_matches": 0.0,
            "proveedores_procesados": [],
            "usd_por_proveedor": {},
            "provisiones_usadas": [],
            "provisiones_duplicadas": False,
            "pendientes": 0,
            "pendientes_descartados": 0,
            "detalle": pd.DataFrame(),
            "pago_path": str(pago_path),
            "matched_index": [],
            "pending_index": [],
            "sin_match_index": [],
        }
        return work, empty_pend, audit

    acre_col = "Acreedor" if "Acreedor" in candidates.columns else "acreedor"
    soc_col = "Sociedad" if "Sociedad" in candidates.columns else (
        "sociedad" if "sociedad" in candidates.columns else None
    )
    name_col = "Name" if "Name" in candidates.columns else (
        "name" if "name" in candidates.columns else None
    )

    candidates["_acre"] = candidates[acre_col].map(_norm_acreedor)
    candidates["_comp"] = candidates["doccompensacion"].map(_norm_doc)
    candidates["_imp_clp"] = pd.to_numeric(
        candidates["importe_en_moneda_doc"], errors="coerce"
    ).fillna(0.0)
    candidates["_doc"] = (
        candidates["no_documento"].map(_norm_doc)
        if "no_documento" in candidates.columns
        else ""
    )
    candidates["_concepto"] = (
        candidates["concepto_detectado"].fillna("").astype(str).str.strip()
        if "concepto_detectado" in candidates.columns
        else ""
    )
    candidates["_ref"] = (
        candidates["referencia"].fillna("").astype(str).str.strip()
        if "referencia" in candidates.columns
        else ""
    )
    candidates["_mon"] = (
        candidates["moneda_del_documento"].fillna("").astype(str).str.strip().str.upper()
        if "moneda_del_documento" in candidates.columns
        else "CLP"
    )
    candidates["_soc"] = (
        candidates[soc_col].fillna("").astype(str).str.strip()
        if soc_col is not None
        else ""
    )
    candidates["_name"] = (
        candidates[name_col].fillna("").astype(str).str.strip()
        if name_col is not None
        else ""
    )

    # Conteos de PROVISIÓN por acreedor en matriz de pago.
    pago_counts = (
        pago_usd.groupby(["_acre", "_prov"]).size().rename("n_pago").reset_index()
    )
    pago_count_map = {
        (r["_acre"], r["_prov"]): int(r["n_pago"]) for _, r in pago_counts.iterrows()
    }
    monto_map = {
        (r["_acre"], r["_prov"]): float(r["_monto"])
        for _, r in pago_usd.iterrows()
        if pd.notna(r["_monto"]) and r["_prov"]
    }

    # Conteos MATRIZ por (acreedor, doccompensacion) para evitar duplicar montos.
    matriz_counts = candidates.groupby(["_acre", "_comp"]).size().to_dict()

    detail_rows: list[dict[str, Any]] = []
    matched_idx: list[Any] = []
    pending_idx: list[Any] = []
    sin_match_idx: list[Any] = []
    used_provisions: list[tuple[str, str]] = []
    usd_por_proveedor: dict[str, float] = {}

    for idx, row in candidates.iterrows():
        acre = row["_acre"]
        comp = row["_comp"]
        key = (acre, comp)
        n_pago = int(pago_count_map.get(key, 0)) if comp else 0
        n_matriz = int(matriz_counts.get(key, 0)) if comp else 0
        clp_amt = float(row["_imp_clp"])
        proveedor = row["_name"] or acre

        if (
            comp
            and n_pago == 1
            and n_matriz == 1
            and key in monto_map
            and key not in used_provisions
        ):
            usd_abs = abs(float(monto_map[key]))
            signed = -usd_abs if clp_amt < 0 else usd_abs
            estado = "MATCH_UNICO"
            motivo = ""
            work.at[idx, "importe_en_moneda_doc"] = signed
            matched_idx.append(idx)
            used_provisions.append(key)
            monto_usd = signed
            usd_por_proveedor[proveedor] = usd_por_proveedor.get(proveedor, 0.0) + signed
            cantidad = 1
        elif n_pago > 1 or n_matriz > 1:
            estado = "MATCH_MULTIPLE"
            motivo = "MATCH_MULTIPLE"
            monto_usd = None
            cantidad = max(n_pago, n_matriz)
            pending_idx.append(idx)
        else:
            estado = "SIN_MATCH"
            motivo = "SIN_MATCH"
            monto_usd = None
            cantidad = n_pago
            pending_idx.append(idx)
            sin_match_idx.append(idx)

        detail_rows.append(
            {
                "sociedad": row["_soc"],
                "acreedor": acre,
                "proveedor": proveedor,
                "documento_clp": row["_doc"],
                "doccompensacion": comp,
                "concepto": row["_concepto"],
                "referencia": row["_ref"],
                "moneda": row["_mon"],
                "monto_clp_anterior": clp_amt,
                "monto_usd_nuevo": monto_usd,
                "estado": estado,
                "motivo": motivo,
                "cantidad_coincidencias": cantidad,
            }
        )

    detail = pd.DataFrame(detail_rows)
    n_match = int((detail["estado"] == "MATCH_UNICO").sum()) if not detail.empty else 0
    n_sin = int((detail["estado"] == "SIN_MATCH").sum()) if not detail.empty else 0
    n_multi = int((detail["estado"] == "MATCH_MULTIPLE").sum()) if not detail.empty else 0
    total_usd = (
        float(
            detail.loc[detail["estado"] == "MATCH_UNICO", "monto_usd_nuevo"]
            .astype(float)
            .sum()
        )
        if n_match
        else 0.0
    )

    pend_src = detail.loc[detail["estado"].isin({"SIN_MATCH", "MATCH_MULTIPLE"})].copy()
    if pend_src.empty:
        pendientes = pd.DataFrame(
            columns=[
                "Sociedad",
                "Acreedor",
                "Proveedor",
                "N° documento CLP",
                "doccompensacion",
                "Referencia",
                "Concepto",
                "Moneda",
                "Monto CLP original",
                "Motivo",
                "cantidad de coincidencias encontradas",
            ]
        )
    else:
        pendientes = pend_src.rename(
            columns={
                "sociedad": "Sociedad",
                "acreedor": "Acreedor",
                "proveedor": "Proveedor",
                "documento_clp": "N° documento CLP",
                "doccompensacion": "doccompensacion",
                "referencia": "Referencia",
                "concepto": "Concepto",
                "moneda": "Moneda",
                "monto_clp_anterior": "Monto CLP original",
                "motivo": "Motivo",
                "cantidad_coincidencias": "cantidad de coincidencias encontradas",
            }
        )[
            [
                "Sociedad",
                "Acreedor",
                "Proveedor",
                "N° documento CLP",
                "doccompensacion",
                "Referencia",
                "Concepto",
                "Moneda",
                "Monto CLP original",
                "Motivo",
                "cantidad de coincidencias encontradas",
            ]
        ].reset_index(drop=True)

    proveedores = sorted({r["proveedor"] for r in detail_rows if r["proveedor"]})
    audit = {
        "procesados": int(len(detail)),
        "match_unico": n_match,
        "sin_match": n_sin,
        "multi": n_multi,
        "total_usd_matches": total_usd,
        "proveedores_procesados": proveedores,
        "usd_por_proveedor": dict(sorted(usd_por_proveedor.items())),
        "provisiones_usadas": [f"{a}:{p}" for a, p in used_provisions],
        "provisiones_duplicadas": len(used_provisions) != len(set(used_provisions)),
        "pendientes": int(len(pendientes)),
        "pendientes_descartados": 0,
        "detalle": detail,
        "pago_path": str(pago_path),
        "matched_index": matched_idx,
        "pending_index": pending_idx,
        "sin_match_index": sin_match_idx,
    }
    return work, pendientes, audit


def apply_gasandes_usd_prueba(
    matrix: pd.DataFrame,
    *,
    pago_path: Path = PAGO_ME_PATH,
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Compatibilidad EF-13.2: delega en apply_usd_pago_me y filtra Gasandes."""
    work, _pendientes, audit = apply_usd_pago_me(matrix, pago_path=pago_path)
    detail = audit.get("detalle")
    if detail is not None and not detail.empty and "acreedor" in detail.columns:
        g = detail.loc[detail["acreedor"].eq(GASANDES_ACREEDOR)].copy()
        n_match = int((g["estado"] == "MATCH_UNICO").sum()) if not g.empty else 0
        audit = {
            **audit,
            "acreedor": GASANDES_ACREEDOR,
            "proveedor": GASANDES_PROVEEDOR,
            "procesados": int(len(g)),
            "match_unico": n_match,
            "sin_match": int((g["estado"] == "SIN_MATCH").sum()) if not g.empty else 0,
            "multi": int((g["estado"] == "MATCH_MULTIPLE").sum()) if not g.empty else 0,
            "total_usd_matches": float(
                g.loc[g["estado"] == "MATCH_UNICO", "monto_usd_nuevo"]
                .astype(float)
                .sum()
            )
            if n_match
            else 0.0,
            "detalle": g,
        }
    return work, audit


def build_dinamicas(matrix: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """
    Construir hojas finales.

    EF-12.6: excluye acreedor 4000000387 de las cinco dinámicas base
    y lo concentra en COMB_GNL CHILE.

    EF-13.7: COMB_USD / NO_COMB_USD / CLP_USD usan montos USD de pago ME
    en matches únicos del universo CLP_USD. Pendientes → PENDIENTES_USD.

    EF-14.2: sin nomenclatura → NO_ECM; dinámicas solo con nomenclatura.
    COMB_CLP / NO_COMB_CLP / COMB_GNL recalculadas sin esos registros.

    EF-14.4: prepara DATOS_* documentales para pivots jerárquicos.

    EF-14.9: CONFIRMING_USD = SIN_MATCH de acreedores confirming
    (vista adicional; no modifica PENDIENTES_USD ni matching).
    """
    global _LAST_EF132_AUDIT, _LAST_EF137_AUDIT, _LAST_EF142_AUDIT
    global _LAST_EF149_AUDIT, _LAST_DOCUMENT_DATOS, _LAST_EF144_AUDIT

    # EF-14.2: partición universo (sin modificar MATRIZ en disco).
    no_ecm_mask = outside_nomenclature_mask(matrix)
    no_ecm_source = matrix.loc[no_ecm_mask].copy()
    matrix_nom = matrix.loc[~no_ecm_mask].copy()

    exclusive_mask = _acreedor_mask(matrix_nom, EXCLUSIVE_ACREEDOR)
    matrix_rest = matrix_nom.loc[~exclusive_mask].copy()

    margen = InformeMargenModule()
    td_con_proveedor = margen.generate_td_clp_dynamics(
        matrix_rest,
        include_proveedor=True,
    )
    td_sin_proveedor = margen.generate_td_clp_dynamics(
        matrix_rest,
        include_proveedor=False,
    )

    # EF-13.7: override USD en memoria; pendientes fuera de pivots USD.
    matrix_usd, pendientes, audit = apply_usd_pago_me(matrix_rest)
    _LAST_EF137_AUDIT = audit
    _LAST_EF132_AUDIT = audit

    pending_index = set(audit.get("pending_index") or [])
    if pending_index:
        matrix_usd_matched = matrix_usd.loc[
            ~matrix_usd.index.isin(pending_index)
        ].copy()
    else:
        matrix_usd_matched = matrix_usd

    # EF-14.9: confirming SIN_MATCH (CLP original; subset de pendientes).
    src_confirming = _select_confirming_sin_match(matrix_rest, audit)
    confirming = _build_confirming_usd(src_confirming)

    # EF-14.4: detalle documental (misma segmentación / montos USD).
    document_datos = build_document_datos(
        matrix_rest=matrix_rest,
        matrix_nom=matrix_nom,
        matrix_usd_matched=matrix_usd_matched,
    )
    document_datos[CONFIRMING_SHEET] = _build_document_detail(src_confirming)
    _LAST_DOCUMENT_DATOS = document_datos

    comb_usd = margen.generate_pivot_clp(
        matrix_usd_matched,
        tipo_filter="COM",
        standardize=True,
        clp_usd=True,
        clp_usd_acreedores=USD_ACREEDORES_COM,
        include_proveedor=True,
    )
    no_comb_usd = margen.generate_pivot_clp(
        matrix_usd_matched,
        tipo_filter="NO_COM",
        standardize=True,
        clp_usd=True,
        clp_usd_acreedores=USD_ACREEDORES_NO_COM,
        include_proveedor=True,
    )

    if comb_usd.empty and no_comb_usd.empty:
        clp_usd = pd.DataFrame()
    elif comb_usd.empty:
        clp_usd = no_comb_usd.copy()
    elif no_comb_usd.empty:
        clp_usd = comb_usd.copy()
    else:
        clp_usd = pd.concat([comb_usd, no_comb_usd], ignore_index=True)

    no_ecm_sheet = _build_no_ecm_sheet(no_ecm_source)
    comb_clp = td_con_proveedor[SHEET_MAP["COMB_CLP"]].copy()
    no_comb_clp = td_sin_proveedor[SHEET_MAP["NO_COMB_CLP"]].copy()
    comb_gnl = _build_comb_gnl_chile(matrix_nom)

    audit["total_usd_COMB_USD"] = _sheet_total(comb_usd)
    audit["total_usd_NO_COMB_USD"] = _sheet_total(no_comb_usd)
    audit["total_usd_CLP_USD"] = _sheet_total(clp_usd)
    audit["total_COMB_CLP"] = _sheet_total(comb_clp)
    audit["total_NO_COMB_CLP"] = _sheet_total(no_comb_clp)
    audit["total_COMB_GNL"] = _sheet_total(comb_gnl)
    audit["total_CONFIRMING_USD"] = _sheet_total(confirming)

    # Conciliación CONFIRMING vs PENDIENTES_USD (SIN_MATCH + acreedores).
    pend_acre = (
        pendientes["Acreedor"].map(_norm_acreedor)
        if not pendientes.empty and "Acreedor" in pendientes.columns
        else pd.Series(dtype=str)
    )
    pend_motivo = (
        pendientes["Motivo"].fillna("").astype(str).str.strip()
        if not pendientes.empty and "Motivo" in pendientes.columns
        else pd.Series(dtype=str)
    )
    pend_conf = (
        pendientes.loc[
            pend_acre.isin(CONFIRMING_ACREEDORES) & pend_motivo.eq("SIN_MATCH")
        ].copy()
        if not pendientes.empty
        else pendientes
    )
    confirming_imp = _importe_sum(src_confirming)
    pend_conf_imp = (
        float(
            pd.to_numeric(pend_conf["Monto CLP original"], errors="coerce")
            .fillna(0.0)
            .sum()
        )
        if not pend_conf.empty and "Monto CLP original" in pend_conf.columns
        else 0.0
    )
    proveedores = sorted(
        {
            str(v).strip()
            for v in (
                src_confirming["Name"]
                if "Name" in src_confirming.columns
                else src_confirming.get("name", pd.Series(dtype=str))
            )
            .fillna("")
            .astype(str)
            .tolist()
            if str(v).strip()
        }
    ) if not src_confirming.empty else []
    acreedores_presentes = sorted(
        {
            _norm_acreedor(v)
            for v in (
                src_confirming["Acreedor"]
                if "Acreedor" in src_confirming.columns
                else src_confirming.get("acreedor", pd.Series(dtype=str))
            ).tolist()
            if _norm_acreedor(v)
        }
    ) if not src_confirming.empty else []

    ef149: dict[str, Any] = {
        "registros": int(len(src_confirming)),
        "proveedores": proveedores,
        "acreedores": acreedores_presentes,
        "todos_sin_match": True,
        "importe_clp": confirming_imp,
        "total_pivot": audit["total_CONFIRMING_USD"],
        "pendientes_confirming_sin_match": int(len(pend_conf)),
        "pendientes_confirming_importe": pend_conf_imp,
        "conciliacion_registros_ok": int(len(src_confirming)) == int(len(pend_conf)),
        "conciliacion_importe_ok": abs(confirming_imp - pend_conf_imp) < 0.01,
        "acreedores_esperados": sorted(CONFIRMING_ACREEDORES),
        "solo_acreedores_confirming": set(acreedores_presentes).issubset(
            CONFIRMING_ACREEDORES
        ),
    }
    _LAST_EF149_AUDIT = ef149

    universo_n = int(len(matrix))
    dinamicas_n = int(len(matrix_nom))
    no_ecm_n = int(len(no_ecm_source))
    universo_imp = _importe_sum(matrix)
    dinamicas_imp = _importe_sum(matrix_nom)
    no_ecm_imp = _importe_sum(no_ecm_source)
    ef142: dict[str, Any] = {
        "registros_trasladados": no_ecm_n,
        "monto_trasladado": no_ecm_imp,
        "universo_registros": universo_n,
        "dinamicas_registros": dinamicas_n,
        "no_ecm_registros": no_ecm_n,
        "universo_importe": universo_imp,
        "dinamicas_importe": dinamicas_imp,
        "no_ecm_importe": no_ecm_imp,
        "conciliacion_registros_ok": universo_n == dinamicas_n + no_ecm_n,
        "conciliacion_importe_ok": abs(universo_imp - (dinamicas_imp + no_ecm_imp))
        < 0.01,
        "duplicados": bool(matrix.index.duplicated().any())
        or universo_n != dinamicas_n + no_ecm_n,
        "sin_concepto": int(
            (
                no_ecm_source["concepto_detectado"].fillna("").astype(str).str.strip()
                == ""
            ).sum()
        )
        if not no_ecm_source.empty and "concepto_detectado" in no_ecm_source.columns
        else 0,
        "fuera_nomenclatura": int(
            (
                no_ecm_source["concepto_detectado"].fillna("").astype(str).str.strip()
                != ""
            ).sum()
        )
        if not no_ecm_source.empty and "concepto_detectado" in no_ecm_source.columns
        else 0,
        "total_COMB_CLP": audit["total_COMB_CLP"],
        "total_NO_COMB_CLP": audit["total_NO_COMB_CLP"],
        "total_COMB_USD": audit["total_usd_COMB_USD"],
        "total_NO_COMB_USD": audit["total_usd_NO_COMB_USD"],
        "total_CLP_USD": audit["total_usd_CLP_USD"],
        "total_COMB_GNL": audit["total_COMB_GNL"],
        "total_CONFIRMING_USD": audit["total_CONFIRMING_USD"],
        "total_NO_ECM": no_ecm_imp,
    }
    _LAST_EF142_AUDIT = ef142

    ef144: dict[str, Any] = {
        "filas_por_datos": {
            name: int(len(frame)) for name, frame in document_datos.items()
        },
        "importe_por_datos": {
            name: float(
                pd.to_numeric(frame[IMPORTE_DINAMICA_COL], errors="coerce")
                .fillna(0.0)
                .sum()
            )
            if not frame.empty and IMPORTE_DINAMICA_COL in frame.columns
            else 0.0
            for name, frame in document_datos.items()
        },
        "duplicados_por_datos": {
            name: int(frame.duplicated().sum()) if not frame.empty else 0
            for name, frame in document_datos.items()
        },
    }
    sample = next((f for f in document_datos.values() if not f.empty), pd.DataFrame())
    ef144["columnas_documentales"] = (
        list(sample.columns)
        if not sample.empty
        else [
            *[label for label, _ in DOCUMENT_COLUMN_MAP],
            CONCEPT_COL,
            PROVEEDOR_COL,
            "Mes",
            IMPORTE_DINAMICA_COL,
        ]
    )
    ef144["total_COMB_CLP"] = audit["total_COMB_CLP"]
    ef144["total_COMB_USD"] = audit["total_usd_COMB_USD"]
    ef144["total_NO_COMB_USD"] = audit["total_usd_NO_COMB_USD"]
    ef144["total_CLP_USD"] = audit["total_usd_CLP_USD"]
    ef144["total_COMB_GNL"] = audit["total_COMB_GNL"]
    ef144["total_CONFIRMING_USD"] = audit["total_CONFIRMING_USD"]
    _LAST_EF144_AUDIT = ef144

    out: dict[str, pd.DataFrame] = {
        "COMB_CLP": comb_clp,
        "NO_COMB_CLP": no_comb_clp,
        "COMB_USD": comb_usd,
        "NO_COMB_USD": no_comb_usd,
        "CLP_USD": clp_usd,
        EXCLUSIVE_SHEET: comb_gnl,
        CONFIRMING_SHEET: confirming,
        PENDIENTES_SHEET: pendientes,
        NO_ECM_SHEET: no_ecm_sheet,
    }

    return {name: out[name] for name in FINAL_SHEET_ORDER}


def _flat_to_long(frame: pd.DataFrame) -> pd.DataFrame:
    """Convertir dinámica ancha a formato largo para PivotTable Excel."""
    if frame is None or frame.empty:
        return pd.DataFrame(
            columns=[CONCEPT_COL, PROVEEDOR_COL, "Mes", "Importe"]
        )

    work = frame.copy()
    if CONCEPT_COL not in work.columns:
        raise KeyError(f"Falta columna {CONCEPT_COL}")
    if PROVEEDOR_COL not in work.columns:
        work[PROVEEDOR_COL] = ""

    concept = work[CONCEPT_COL].fillna("").astype(str)
    # Excluir solo márgenes; conservar concepto vacío (EF-12.6).
    work = work.loc[
        ~concept.str.strip().isin({"All", "all", "Total", "total"})
    ].copy()
    if work.empty:
        return pd.DataFrame(
            columns=[CONCEPT_COL, PROVEEDOR_COL, "Mes", "Importe"]
        )

    work[CONCEPT_COL] = work[CONCEPT_COL].fillna("").astype(str)
    work[PROVEEDOR_COL] = work[PROVEEDOR_COL].fillna("").astype(str).str.strip()

    month_cols = [m for m in MONTH_LABELS if m in work.columns]
    long = work.melt(
        id_vars=[CONCEPT_COL, PROVEEDOR_COL],
        value_vars=month_cols,
        var_name="Mes",
        value_name="Importe",
    )
    long["Importe"] = pd.to_numeric(long["Importe"], errors="coerce").fillna(0.0)
    long["_ord_concepto"] = pd.Categorical(
        long[CONCEPT_COL],
        categories=list(dict.fromkeys(work[CONCEPT_COL].tolist())),
        ordered=True,
    )
    long["_ord_mes"] = pd.Categorical(
        long["Mes"],
        categories=list(MONTH_LABELS),
        ordered=True,
    )
    long = long.sort_values(
        ["_ord_concepto", PROVEEDOR_COL, "_ord_mes"],
        kind="mergesort",
    ).drop(columns=["_ord_concepto", "_ord_mes"])
    return long.reset_index(drop=True)


def _datos_source_sheet(sheet_name: str) -> str:
    name = str(sheet_name or "")
    if name.startswith("DATOS_"):
        return name[6:]
    return name


def _pivot_total_format(sheet_name: str) -> str | None:
    if sheet_name in USD_PIVOT_SHEETS:
        return USD_NUMBER_FORMAT_LOCAL
    if sheet_name in CLP_PIVOT_SHEETS:
        return CLP_NUMBER_FORMAT_LOCAL
    return None


def _column_money_format(sheet_name: str, column_name: str) -> str | None:
    """Formato visual por columna. L (importe doc.) es CLP; Q depende de la hoja."""
    name = str(column_name or "").strip()
    source = _datos_source_sheet(sheet_name)
    if name == "Monto CLP original":
        return CLP_NUMBER_FORMAT_LOCAL
    if name == "Importe en moneda doc.":
        return CLP_NUMBER_FORMAT_LOCAL
    if name in {IMPORTE_DINAMICA_COL, "Importe"}:
        if source in USD_PIVOT_SHEETS:
            return USD_NUMBER_FORMAT_LOCAL
        if source in CLP_PIVOT_SHEETS:
            return CLP_NUMBER_FORMAT_LOCAL
    return None


def _apply_excel_number_format(com_object, format_string: str) -> None:
    """Aplicar formato es-CL. No altera el valor almacenado."""
    try:
        com_object.NumberFormatLocal = format_string
        return
    except Exception:
        pass
    try:
        com_object.NumberFormat = format_string
    except Exception:
        pass


def _apply_pivot_total_format(pivot, sheet_name: str) -> None:
    fmt = _pivot_total_format(sheet_name)
    if not fmt:
        return
    for getter in (
        lambda: pivot.PivotFields("Total"),
        lambda: pivot.DataFields(1),
    ):
        try:
            _apply_excel_number_format(getter(), fmt)
            return
        except Exception:
            continue


def _apply_sheet_column_money_formats(ws) -> None:
    sheet_name = str(ws.Name)
    if sheet_name == NO_ECM_SHEET:
        return
    used = int(ws.UsedRange.Columns.Count)
    last_row = int(ws.UsedRange.Row + ws.UsedRange.Rows.Count - 1)
    if last_row < 2:
        return
    for col_idx in range(1, used + 1):
        header = ws.Cells(1, col_idx).Value
        fmt = _column_money_format(sheet_name, str(header) if header is not None else "")
        if not fmt:
            continue
        _apply_excel_number_format(
            ws.Range(ws.Cells(2, col_idx), ws.Cells(last_row, col_idx)),
            fmt,
        )


def _apply_workbook_money_formats(wb) -> None:
    """Reaplicar formatos de montos a pivots y columnas DATOS/PENDIENTES."""
    for ws in list(wb.Worksheets):
        name = str(ws.Name)
        if name == NO_ECM_SHEET:
            continue
        try:
            _apply_sheet_column_money_formats(ws)
        except Exception:
            pass
        if name in HIERARCHICAL_SHEETS:
            try:
                _apply_pivot_total_format(ws.PivotTables(1), name)
            except Exception:
                pass


def _hex_to_excel_color(hex_color: str) -> int:
    """Convertir #RRGGBB al entero RGB que usa Excel COM (VBA RGB)."""
    text = str(hex_color or "").strip().lstrip("#")
    if len(text) != 6:
        raise ValueError(f"Color hexadecimal inválido: {hex_color!r}")
    red = int(text[0:2], 16)
    green = int(text[2:4], 16)
    blue = int(text[4:6], 16)
    return red + (green * 256) + (blue * 65536)


def _is_sen_concept(name: object) -> bool:
    """True si el concepto, trim + mayúsculas, comienza con SEN_."""
    return str(name or "").strip().upper().startswith("SEN_")


def _is_amount_header(name: object) -> bool:
    text = str(name or "").strip().lower()
    if not text:
        return False
    if text in {
        "importe",
        "importe usado en la dinámica",
        "importe en moneda doc.",
        "monto clp original",
        "monto bruto",
        "total",
        "total general",
        "grand total",
    }:
        return True
    return "importe" in text or text.startswith("monto")


def _clamp_width(value: float, minimum: float, maximum: float) -> float:
    return max(minimum, min(maximum, float(value)))


def _style_header_range(rng, fill_color: int) -> None:
    """Fondo azul oscuro, texto blanco y negrita. No toca NumberFormat."""
    rng.Interior.Color = fill_color
    rng.Font.Color = PIVOT_HEADER_FONT_COLOR
    rng.Font.Bold = True
    rng.HorizontalAlignment = XL_ALIGN_CENTER
    rng.VerticalAlignment = XL_ALIGN_CENTER


def _worksheet_exists(wb, name: str) -> bool:
    try:
        wb.Worksheets(name)
        return True
    except Exception:
        return False


def _apply_pivot_captions(pivot, sheet_name: str) -> None:
    """Captions de presentación. No cambia SourceName ni valores."""
    try:
        pivot.CompactLayoutRowHeader = "Concepto"
    except Exception as exc:
        logger.warning(
            "%s: no se pudo asignar CompactLayoutRowHeader 'Concepto' (%s)",
            sheet_name,
            exc,
        )
    try:
        pivot.CompactLayoutColumnHeader = "Mes de Pago"
    except Exception as exc:
        logger.warning(
            "%s: no se pudo asignar CompactLayoutColumnHeader 'Mes de Pago' (%s)",
            sheet_name,
            exc,
        )
    try:
        pivot.PivotFields("Mes").Caption = "Mes de Pago"
    except Exception as exc:
        logger.warning(
            "%s: no se pudo asignar Caption del campo Mes (%s)",
            sheet_name,
            exc,
        )
    caption_set = False
    for getter in (
        lambda: pivot.DataFields(1),
        lambda: pivot.PivotFields("Total"),
    ):
        try:
            getter().Caption = sheet_name
            caption_set = True
            break
        except Exception:
            continue
    if not caption_set:
        logger.warning(
            "%s: no se pudo renombrar el data field 'Total' al nombre de hoja",
            sheet_name,
        )


def _collapse_sen_row_items(pivot, sheet_name: str) -> None:
    """Contraer solo PivotItems SEN_ (ShowDetail=False). No usa Visible=False."""
    try:
        field = pivot.PivotFields(CONCEPT_COL)
        items = field.PivotItems()
        count = int(items.Count)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo enumerar PivotItems de concepto (%s)",
            sheet_name,
            exc,
        )
        return
    try:
        pivot_name = str(pivot.Name)
    except Exception:
        pivot_name = ""
    for index in range(1, count + 1):
        try:
            item = items.Item(index)
            name = str(item.Name)
        except Exception as exc:
            logger.warning(
                "%s: fallo al recuperar PivotItems.Item(%s) "
                "(pivot=%s, operación=PivotItems.Item). (%s)",
                sheet_name,
                index,
                pivot_name or "N/D",
                exc,
            )
            continue
        if not _is_sen_concept(name):
            continue
        try:
            item.ShowDetail = False
        except Exception as exc:
            logger.warning(
                "%s: fallo al aplicar ShowDetail=False al concepto SEN_ %r "
                "(pivot=%s, índice=%s, operación=ShowDetail=False). (%s)",
                sheet_name,
                name,
                pivot_name or "N/D",
                index,
                exc,
            )


def _rename_grand_total_label(pivot, sheet_name: str) -> None:
    """
    Renombrar 'Total general' / 'Grand Total' a 'Monto Bruto'.

    Excel COM no expone una propiedad estable GrandTotalName. Se escribe
    sobre celdas de texto del TableRange. Si COM rechaza la escritura,
    se registra warning y se continúa (fallback: conservar etiqueta original).
    """
    try:
        rng = pivot.TableRange1
        start_row = int(rng.Row)
        start_col = int(rng.Column)
        n_rows = int(rng.Rows.Count)
        n_cols = int(rng.Columns.Count)
        ws = rng.Worksheet
    except Exception as exc:
        logger.warning(
            "%s: no se pudo leer TableRange1 para 'Monto Bruto' (%s)",
            sheet_name,
            exc,
        )
        return
    renamed = 0
    for row in range(start_row, start_row + n_rows):
        for col in range(start_col, start_col + n_cols):
            cell = ws.Cells(row, col)
            try:
                raw = cell.Value
            except Exception:
                continue
            if raw is None or isinstance(raw, (int, float)):
                continue
            text = str(raw).strip()
            if text.casefold() not in _GRAND_TOTAL_LABELS:
                continue
            try:
                cell.Value = "Monto Bruto"
                renamed += 1
            except Exception as exc:
                logger.warning(
                    "%s: Excel COM rechazó escribir 'Monto Bruto' en %s "
                    "(se conserva %r). Fallback: etiqueta original. (%s)",
                    sheet_name,
                    cell.Address,
                    text,
                    exc,
                )
    if renamed == 0:
        logger.warning(
            "%s: no se encontró celda 'Total general'/'Grand Total' para "
            "renombrar a 'Monto Bruto' (fallback: se conserva el texto de Excel)",
            sheet_name,
        )


def _pivot_display_name(pivot) -> str:
    try:
        return str(pivot.Name)
    except Exception:
        return "N/D"


def _grand_total_row_range(ws, pivot):
    """
    Rango de la fila de total general. No usa dirección fija ni UsedRange.

    1) PivotRowAxis.PivotLines con LineType = xlPivotLineGrandTotal.
    2) Respaldo: última fila de TableRange1 solo si RowGrand y la primera
       celda es exactamente Monto Bruto / Total general / Grand Total.
    """
    table = pivot.TableRange1
    first_col = int(table.Column)
    last_col = first_col + int(table.Columns.Count) - 1

    try:
        lines = pivot.PivotRowAxis.PivotLines
        count = int(lines.Count)
        for index in range(1, count + 1):
            line = lines.Item(index)
            if int(line.LineType) != XL_PIVOT_LINE_GRAND_TOTAL:
                continue
            row_num = int(line.PivotLineCells.Item(1).Range.Row)
            return ws.Range(
                ws.Cells(row_num, first_col),
                ws.Cells(row_num, last_col),
            )
    except Exception:
        pass

    try:
        row_grand = bool(pivot.RowGrand)
    except Exception:
        row_grand = False
    if not row_grand:
        return None
    last_row = int(table.Row) + int(table.Rows.Count) - 1
    try:
        raw = ws.Cells(last_row, first_col).Value
    except Exception:
        return None
    if raw is None or isinstance(raw, (int, float)):
        return None
    if str(raw).strip() not in _PIVOT_TOTAL_ROW_LABELS:
        return None
    return ws.Range(
        ws.Cells(last_row, first_col),
        ws.Cells(last_row, last_col),
    )


def _apply_pivot_grand_total_row_style(ws, pivot, sheet_name: str, fill: int) -> None:
    """Estilo de la fila de total general. Solo Interior/Font/alineación."""
    pivot_name = _pivot_display_name(pivot)
    try:
        rng = _grand_total_row_range(ws, pivot)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo identificar la fila de total general "
            "(pivot=%s, operación=PivotRowAxis.PivotLines/TableRange1). (%s)",
            sheet_name,
            pivot_name,
            exc,
        )
        return
    if rng is None:
        logger.warning(
            "%s: no se pudo identificar la fila de total general "
            "(pivot=%s, operación=PivotRowAxis.PivotLines/TableRange1).",
            sheet_name,
            pivot_name,
        )
        return
    try:
        _style_header_range(rng, fill)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo aplicar estilo a la fila de total general "
            "(pivot=%s, operación=_style_header_range). (%s)",
            sheet_name,
            pivot_name,
            exc,
        )


def _apply_pivot_header_style(ws, pivot, sheet_name: str) -> None:
    """Estilo de encabezados del pivot. No usa TableStyle2 ni toca el cuerpo."""
    fill = _hex_to_excel_color(PIVOT_HEADER_COLOR_HEX)
    try:
        table = pivot.TableRange1
        first_row = int(table.Row)
        first_col = int(table.Column)
        last_col = first_col + int(table.Columns.Count) - 1
        header = ws.Range(
            ws.Cells(first_row, first_col),
            ws.Cells(first_row, last_col),
        )
        _style_header_range(header, fill)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo aplicar estilo a la fila de encabezados (%s)",
            sheet_name,
            exc,
        )
    try:
        _style_header_range(pivot.ColumnRange, fill)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo aplicar estilo a ColumnRange (%s)",
            sheet_name,
            exc,
        )
    try:
        # Encabezado compacto de fila: misma fila que ColumnRange, primera
        # columna de TableRange1 (equivale a A4 sin fijar la dirección).
        compact = ws.Cells(
            int(pivot.ColumnRange.Row),
            int(pivot.TableRange1.Column),
        )
        _style_header_range(compact, fill)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo aplicar estilo al encabezado compacto de fila "
            "(pivot=%s, operación=ColumnRange.Row/TableRange1.Column). (%s)",
            sheet_name,
            _pivot_display_name(pivot),
            exc,
        )
    _apply_pivot_grand_total_row_style(ws, pivot, sheet_name, fill)


def _apply_pivot_column_widths(ws, pivot, sheet_name: str) -> None:
    """Anchos controlados. No AutoFit ilimitado. No convierte valores a texto."""
    try:
        table = pivot.TableRange1
        first_col = int(table.Column)
        n_cols = int(table.Columns.Count)
    except Exception as exc:
        logger.warning(
            "%s: no se pudo ajustar anchos del PivotTable (%s)",
            sheet_name,
            exc,
        )
        return
    try:
        ws.Columns(first_col).ColumnWidth = _clamp_width(
            ROW_LABEL_COL_WIDTH,
            ROW_LABEL_COL_WIDTH,
            ROW_LABEL_COL_WIDTH_MAX,
        )
        amount_width = _clamp_width(
            AMOUNT_COL_WIDTH_MIN,
            AMOUNT_COL_WIDTH_MIN,
            AMOUNT_COL_WIDTH_MAX,
        )
        for offset in range(1, n_cols):
            ws.Columns(first_col + offset).ColumnWidth = amount_width
    except Exception as exc:
        logger.warning(
            "%s: fallo al asignar ColumnWidth (%s)",
            sheet_name,
            exc,
        )


def _apply_static_sheet_com_presentation(ws) -> None:
    """Encabezado y anchos para PENDIENTES_USD / NO_ECM. No PivotTable."""
    sheet_name = str(ws.Name)
    if sheet_name.startswith("DATOS_"):
        return
    try:
        used = ws.UsedRange
        n_cols = int(used.Columns.Count)
        n_rows = int(used.Rows.Count)
        first_row = int(used.Row)
        first_col = int(used.Column)
    except Exception as exc:
        logger.warning("%s: no se pudo leer UsedRange (%s)", sheet_name, exc)
        return
    if n_cols < 1:
        return
    fill = _hex_to_excel_color(PIVOT_HEADER_COLOR_HEX)
    try:
        header = ws.Range(
            ws.Cells(first_row, first_col),
            ws.Cells(first_row, first_col + n_cols - 1),
        )
        _style_header_range(header, fill)
        header.WrapText = True
    except Exception as exc:
        logger.warning("%s: no se pudo estilizar encabezado estático (%s)", sheet_name, exc)
    last_row = first_row + n_rows - 1
    for offset in range(n_cols):
        col_idx = first_col + offset
        try:
            header_value = ws.Cells(first_row, col_idx).Value
            column = ws.Columns(col_idx)
            if _is_amount_header(header_value):
                column.ColumnWidth = _clamp_width(
                    AMOUNT_COL_WIDTH_MIN,
                    AMOUNT_COL_WIDTH_MIN,
                    AMOUNT_COL_WIDTH_MAX,
                )
                if last_row > first_row:
                    data = ws.Range(
                        ws.Cells(first_row + 1, col_idx),
                        ws.Cells(last_row, col_idx),
                    )
                    data.HorizontalAlignment = XL_ALIGN_RIGHT
            else:
                raw_len = len(str(header_value or "").strip()) + 2
                column.ColumnWidth = _clamp_width(
                    max(STATIC_TEXT_COL_WIDTH_MIN, raw_len),
                    STATIC_TEXT_COL_WIDTH_MIN,
                    STATIC_TEXT_COL_WIDTH_MAX,
                )
                if last_row > first_row:
                    data = ws.Range(
                        ws.Cells(first_row + 1, col_idx),
                        ws.Cells(last_row, col_idx),
                    )
                    data.HorizontalAlignment = XL_ALIGN_LEFT
        except Exception as exc:
            logger.warning(
                "%s: no se pudo ajustar columna %s (%s)",
                sheet_name,
                col_idx,
                exc,
            )


def _apply_openpyxl_static_presentation(ws) -> None:
    """Estilo equivalente para NO_ECM (se escribe con openpyxl, no COM)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if ws.max_column < 1:
        return
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = fill
        cell.font = font
        cell.alignment = header_align
        header = cell.value
        letter = get_column_letter(col_idx)
        if _is_amount_header(header):
            width = _clamp_width(
                AMOUNT_COL_WIDTH_MIN,
                AMOUNT_COL_WIDTH_MIN,
                AMOUNT_COL_WIDTH_MAX,
            )
            data_align = Alignment(horizontal="right", vertical="center")
        else:
            raw_len = len(str(header or "").strip()) + 2
            width = _clamp_width(
                max(STATIC_TEXT_COL_WIDTH_MIN, raw_len),
                STATIC_TEXT_COL_WIDTH_MIN,
                STATIC_TEXT_COL_WIDTH_MAX,
            )
            data_align = Alignment(horizontal="left", vertical="center")
        ws.column_dimensions[letter].width = width
        if ws.max_row >= 2:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, col_idx).alignment = data_align


def _apply_presentation_after_refresh(wb) -> None:
    """
    Presentación posterior al último Refresh.

    Orden: captions → colapso SEN_ (solo NO_COMB_CLP) → Total general
    → estilo encabezados → anchos. DATOS_* no se tocan.
    """
    for sheet_name in HIERARCHICAL_SHEETS:
        try:
            ws = wb.Worksheets(sheet_name)
        except Exception as exc:
            logger.warning(
                "%s: no se pudo abrir la hoja jerárquica esperada "
                "(pivot=N/D, operación=Worksheets). (%s)",
                sheet_name,
                exc,
            )
            continue
        try:
            pivot = ws.PivotTables(1)
        except Exception as exc:
            logger.warning(
                "%s: hoja jerárquica esperada sin PivotTable "
                "(pivot=N/D, operación=PivotTables(1)). (%s)",
                sheet_name,
                exc,
            )
            continue
        try:
            _apply_pivot_captions(pivot, sheet_name)
            if sheet_name == "NO_COMB_CLP":
                _collapse_sen_row_items(pivot, sheet_name)
            _rename_grand_total_label(pivot, sheet_name)
            _apply_pivot_header_style(ws, pivot, sheet_name)
            _apply_pivot_column_widths(ws, pivot, sheet_name)
        except Exception as exc:
            logger.warning(
                "%s: presentación de PivotTable incompleta (%s)",
                sheet_name,
                exc,
            )
    for static_name in (PENDIENTES_SHEET, NO_ECM_SHEET):
        if not _worksheet_exists(wb, static_name):
            continue
        try:
            _apply_static_sheet_com_presentation(wb.Worksheets(static_name))
        except Exception as exc:
            logger.warning(
                "%s: fallo en presentación estática (%s)",
                static_name,
                exc,
            )


def _write_sheet_values(ws, df: pd.DataFrame) -> tuple[int, int]:
    headers = list(df.columns)
    for col_idx, name in enumerate(headers, start=1):
        ws.Cells(1, col_idx).Value = name
    values = df.where(pd.notna(df), None).values.tolist()
    n_rows = len(values)
    n_cols = len(headers)
    if n_rows:
        ws.Range(ws.Cells(2, 1), ws.Cells(n_rows + 1, n_cols)).Value = values
    sheet_name = str(ws.Name)
    for amount_name in ("Importe", IMPORTE_DINAMICA_COL, "Importe en moneda doc."):
        if amount_name in headers and n_rows:
            idx = headers.index(amount_name) + 1
            fmt = _column_money_format(sheet_name, amount_name)
            if fmt:
                _apply_excel_number_format(
                    ws.Range(ws.Cells(2, idx), ws.Cells(n_rows + 1, idx)),
                    fmt,
                )
    return n_rows + 1, n_cols


def _drop_export_all_row(frame: pd.DataFrame) -> pd.DataFrame:
    """
    EF-14.0: quitar solo la fila de margen pandas Concepto='All' al exportar.

    No altera el pivot ni margins=True; filtro de presentación exclusivo
    de la hoja estática NO_COMB_CLP.
    """
    if frame is None or frame.empty or CONCEPT_COL not in frame.columns:
        return frame
    concept = frame[CONCEPT_COL].fillna("").astype(str).str.strip()
    return frame.loc[concept.ne("All")].copy()


def _label_empty_concept_row(frame: pd.DataFrame) -> pd.DataFrame:
    """
    EF-14.1: etiquetar Concepto vacío como 'SIN CONCEPTO' al exportar.

    Solo presentación en NO_COMB_CLP. No elimina filas ni altera montos.
    No confundir con Total General ni con la fila All.
    """
    if frame is None or frame.empty or CONCEPT_COL not in frame.columns:
        return frame
    out = frame.copy()
    concept = out[CONCEPT_COL]
    empty = concept.isna() | (concept.astype(str).str.strip() == "")
    out.loc[empty, CONCEPT_COL] = "SIN CONCEPTO"
    return out


def _prepare_no_comb_clp_export(frame: pd.DataFrame) -> pd.DataFrame:
    """Presentación estática NO_COMB_CLP: sin All + etiqueta SIN CONCEPTO."""
    return _label_empty_concept_row(_drop_export_all_row(frame))


def _write_static_sheet(ws, df: pd.DataFrame) -> None:
    headers = list(df.columns)
    for col_idx, name in enumerate(headers, start=1):
        ws.Cells(1, col_idx).Value = name
    values = df.where(pd.notna(df), None).values.tolist()
    if values:
        n_rows = len(values)
        n_cols = len(headers)
        ws.Range(ws.Cells(2, 1), ws.Cells(n_rows + 1, n_cols)).Value = values
        if "Monto CLP original" in headers:
            idx = headers.index("Monto CLP original") + 1
            _apply_excel_number_format(
                ws.Range(ws.Cells(2, idx), ws.Cells(n_rows + 1, idx)),
                CLP_NUMBER_FORMAT_LOCAL,
            )


def _set_row_fields(pivot, names: list[str]) -> None:
    for name in reversed(names):
        field = pivot.PivotFields(name)
        field.Orientation = XL_ROW_FIELD
        field.Position = 1
    for idx, name in enumerate(names, start=1):
        pivot.PivotFields(name).Position = idx


def _order_pivot_items(field, ordered_names: list[str]) -> None:
    """Forzar orden de ítems del campo (meses / conceptos)."""
    try:
        field.AutoSort(XL_SORT_MANUAL, field.SourceName)
    except Exception:
        pass
    position = 1
    for name in ordered_names:
        try:
            item = field.PivotItems(name)
            item.Position = position
            position += 1
        except Exception:
            continue


def _prepare_document_for_excel(frame: pd.DataFrame) -> pd.DataFrame:
    """Normalizar tipos para escritura COM (fechas → texto)."""
    if frame is None or frame.empty:
        return frame
    out = frame.copy()
    for col in out.columns:
        series = out[col]
        if pd.api.types.is_datetime64_any_dtype(series):
            out[col] = series.dt.strftime("%Y-%m-%d")
        elif series.dtype == object:
            # Timestamps sueltos / NaT
            out[col] = series.map(
                lambda v: (
                    v.strftime("%Y-%m-%d")
                    if hasattr(v, "strftime") and pd.notna(v)
                    else (None if v is None or (isinstance(v, float) and pd.isna(v)) else v)
                )
            )
    return out


def _create_hierarchical_pivot(
    wb,
    cache,
    sheet_name: str,
    table_name: str,
    *,
    amount_field: str = IMPORTE_DINAMICA_COL,
) -> object:
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = sheet_name
    pivot = cache.CreatePivotTable(
        TableDestination=ws.Range("A3"),
        TableName=table_name,
    )
    _set_row_fields(pivot, [CONCEPT_COL, PROVEEDOR_COL])
    mes_field = pivot.PivotFields("Mes")
    mes_field.Orientation = XL_COLUMN_FIELD
    mes_field.Position = 1
    _order_pivot_items(mes_field, list(MONTH_LABELS))
    pivot.AddDataField(
        pivot.PivotFields(amount_field),
        "Total",
        XL_SUM,
    )
    _apply_pivot_total_format(pivot, sheet_name)
    try:
        pivot.EnableDrilldown = True
    except Exception:
        pass
    try:
        pivot.RowAxisLayout(0)  # xlCompactRow
    except Exception:
        pass
    try:
        pivot.PivotFields(CONCEPT_COL).LayoutCompactRow = True
        pivot.PivotFields(PROVEEDOR_COL).LayoutCompactRow = True
    except Exception:
        pass
    try:
        concept_field = pivot.PivotFields(CONCEPT_COL)
        for i in range(1, 13):
            concept_field.Subtotals[i] = False
    except Exception:
        pass
    return pivot


def export_dinamicas(
    sheets: dict[str, pd.DataFrame],
    output_path: Path,
    document_datos: dict[str, pd.DataFrame] | None = None,
) -> Path:
    """
    Exportar dinámicas:
    - pivots jerárquicos (incl. NO_COMB_CLP / CONFIRMING_USD): DATOS_* documentales
    - PENDIENTES_USD / NO_ECM: estáticas
    """
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "Se requiere pywin32 y Microsoft Excel para crear "
            "PivotTables nativas jerárquicas."
        ) from exc

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        output_path.unlink()

    datos_map = document_datos if document_datos is not None else (_LAST_DOCUMENT_DATOS or {})

    excel = None
    wb = None
    no_ecm_df = sheets.get(NO_ECM_SHEET, pd.DataFrame())
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False

        wb = excel.Workbooks.Add()
        while wb.Worksheets.Count > 1:
            wb.Worksheets(wb.Worksheets.Count).Delete()

        ws_pend = wb.Worksheets(1)
        ws_pend.Name = PENDIENTES_SHEET
        _write_static_sheet(
            ws_pend,
            sheets.get(PENDIENTES_SHEET, pd.DataFrame()),
        )

        for sheet_name in HIERARCHICAL_SHEETS:
            # EF-14.4: origen documental; fallback agregado solo si falta detalle.
            detail = datos_map.get(sheet_name)
            if detail is None or detail.empty:
                long_df = _flat_to_long(sheets[sheet_name])
                amount_field = "Importe"
            else:
                long_df = _prepare_document_for_excel(detail)
                amount_field = IMPORTE_DINAMICA_COL

            ws_datos = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
            datos_name = f"DATOS_{sheet_name}"[:31]
            ws_datos.Name = datos_name
            n_rows, n_cols = _write_sheet_values(ws_datos, long_df)
            if n_rows < 2:
                ws_pivot = wb.Worksheets.Add(
                    After=wb.Worksheets(wb.Worksheets.Count)
                )
                ws_pivot.Name = sheet_name
                ws_pivot.Range("A1").Value = CONCEPT_COL
                ws_pivot.Range("B1").Value = PROVEEDOR_COL
                continue

            source = f"{datos_name}!R1C1:R{n_rows}C{n_cols}"
            cache = wb.PivotCaches().Create(
                SourceType=XL_DATABASE,
                SourceData=source,
            )
            _create_hierarchical_pivot(
                wb,
                cache,
                sheet_name=sheet_name,
                table_name=f"PT_{sheet_name.replace(' ', '_')}"[:31],
                amount_field=amount_field,
            )
            try:
                concepts = [
                    c
                    for c in sheets[sheet_name][CONCEPT_COL]
                    .fillna("")
                    .astype(str)
                    .tolist()
                    if str(c).strip() not in {"All", "all", "Total", "total"}
                ]
                concepts = list(dict.fromkeys(concepts))
                _order_pivot_items(
                    wb.Worksheets(sheet_name).PivotTables(1).PivotFields(
                        CONCEPT_COL
                    ),
                    concepts,
                )
            except Exception:
                pass

        for name in reversed(FINAL_SHEET_ORDER):
            if name == NO_ECM_SHEET:
                continue  # se agrega después con openpyxl
            wb.Worksheets(name).Move(Before=wb.Worksheets(1))

        # EF-14.7: Refresh PivotCache desde DATOS_* (visibles) antes de guardar.
        # Repara ShowDetail vacío sin cambiar datos/totales/estructura.
        # DATOS_* permanecen xlSheetVisible para consumo Power BI.
        for ws in list(wb.Worksheets):
            if str(ws.Name).startswith("DATOS_"):
                ws.Visible = -1  # xlSheetVisible
        for sheet_name in HIERARCHICAL_SHEETS:
            try:
                pt = wb.Worksheets(sheet_name).PivotTables(1)
            except Exception:
                continue
            try:
                pt.SaveData = True
            except Exception:
                pass
            cache = pt.PivotCache()
            cache.Refresh()
            _ = cache.RecordCount
            _apply_pivot_total_format(pt, sheet_name)
        for ws in list(wb.Worksheets):
            if str(ws.Name).startswith("DATOS_"):
                ws.Visible = -1  # xlSheetVisible
        _apply_workbook_money_formats(wb)

        wb.SaveAs(str(output_path.resolve()), FileFormat=XL_OPEN_XML_WORKBOOK)
        wb.Close(SaveChanges=False)
        wb = None
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass

    _append_no_ecm_sheet(output_path, no_ecm_df)
    # EF-14.7: reabrir y Refresh+Save para persistir caches de drill-through
    # (SaveAs inicial puede dejar definitions 2–5 apuntando a records1).
    _persist_refreshed_pivot_caches(output_path)
    return output_path


def _persist_refreshed_pivot_caches(output_path: Path) -> None:
    """
    EF-14.7: reabre el xlsx, refresca cada PivotCache jerárquico y guarda.

    No altera DATOS_*, campos ni totales; solo reconstruye el cache embebido
    para que ShowDetail resuelva filas documentales.
    """
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise ImportError(
            "Se requiere pywin32 y Microsoft Excel para refrescar PivotCache."
        ) from exc

    excel = None
    wb = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        wb = excel.Workbooks.Open(str(Path(output_path).resolve()), UpdateLinks=0)

        for ws in list(wb.Worksheets):
            if str(ws.Name).startswith("DATOS_"):
                ws.Visible = -1  # xlSheetVisible

        for sheet_name in HIERARCHICAL_SHEETS:
            try:
                pt = wb.Worksheets(sheet_name).PivotTables(1)
            except Exception:
                continue
            try:
                pt.SaveData = True
            except Exception:
                pass
            cache = pt.PivotCache()
            cache.Refresh()
            _ = cache.RecordCount
            _apply_pivot_total_format(pt, sheet_name)

        for ws in list(wb.Worksheets):
            if str(ws.Name).startswith("DATOS_"):
                ws.Visible = -1  # xlSheetVisible
        _apply_workbook_money_formats(wb)
        try:
            _apply_presentation_after_refresh(wb)
        except Exception as exc:
            logger.warning(
                "Presentación visual no aplicada (generación continúa): %s",
                exc,
            )

        wb.Save()
        wb.Close(SaveChanges=False)
        wb = None
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def _append_no_ecm_sheet(output_path: Path, no_ecm_df: pd.DataFrame) -> None:
    """Agregar hoja NO_ECM al libro ya generado (detalle de revisión)."""
    from openpyxl import load_workbook

    frame = no_ecm_df.copy() if no_ecm_df is not None else pd.DataFrame()
    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        frame.to_excel(writer, sheet_name=NO_ECM_SHEET, index=False)

    wb = load_workbook(output_path)
    if NO_ECM_SHEET in wb.sheetnames:
        _apply_openpyxl_static_presentation(wb[NO_ECM_SHEET])
    desired = [n for n in FINAL_SHEET_ORDER if n in wb.sheetnames]
    extras = [n for n in wb.sheetnames if n not in desired]
    for idx, name in enumerate(desired + extras, start=1):
        current = wb.sheetnames.index(name) + 1
        if current != idx:
            wb.move_sheet(name, offset=idx - current)
    wb.save(output_path)
    wb.close()


def summarize(sheets: dict[str, pd.DataFrame]) -> dict[str, int]:
    return {name: int(len(frame)) for name, frame in sheets.items()}


def default_output_path(output_dir: Path = DEFAULT_OUTPUT_DIR) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(output_dir) / f"DINAMICAS_FINALES_{stamp}.xlsx"


def _validate_dinamicas_workbook(path: Path) -> None:
    """
    Validar workbook generado: existe, abre y trae las 16 hojas esperadas.

    No altera contenido. Lanza si no es apto para publicar ACTUAL.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Workbook inexistente: {path}")
    if path.stat().st_size <= 0:
        raise ValueError(f"Workbook vacío: {path}")

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        names = set(wb.sheetnames)
    finally:
        wb.close()

    missing = sorted(EXPECTED_DINAMICAS_SHEETS - names)
    if missing:
        raise ValueError(
            "Workbook incompleto para ACTUAL; faltan hojas: " + ", ".join(missing)
        )
    if len(names) < len(EXPECTED_DINAMICAS_SHEETS):
        raise ValueError(
            f"Workbook con {len(names)} hojas; se esperan al menos "
            f"{len(EXPECTED_DINAMICAS_SHEETS)}"
        )


def publish_dinamicas_actual(
    source_path: Path,
    *,
    output_dir: Path | None = None,
) -> Path:
    """
    Publicar DINAMICAS_FINALES_ACTUAL.xlsx por copia segura/atómica.

    Conserva el histórico timestamp. Solo reemplaza ACTUAL si la fuente
    existe, es válida y contiene las 16 hojas esperadas.
    """
    source_path = Path(source_path).resolve()
    out_dir = Path(output_dir) if output_dir is not None else Path(DEFAULT_OUTPUT_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    actual_path = out_dir / ACTUAL_FILENAME

    _validate_dinamicas_workbook(source_path)

    tmp_path = out_dir / f".{ACTUAL_FILENAME}.{os.getpid()}.tmp.xlsx"
    try:
        shutil.copy2(source_path, tmp_path)
        _validate_dinamicas_workbook(tmp_path)
        os.replace(tmp_path, actual_path)
    except Exception:
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass
        raise

    return actual_path


def run(
    matriz_path: Path | None = None,
    output_path: Path | None = None,
) -> tuple[Path, dict[str, pd.DataFrame], dict[str, int], Path]:
    """
    Generar DINAMICAS_FINALES_*.xlsx y publicar ACTUAL si la generación es válida.

    Returns:
        (ruta_salida, hojas, conteos, ruta_matriz_fuente)
    """
    source = (
        Path(matriz_path) if matriz_path else OutputArtifactService().latest_or_raise()
    )
    matrix = load_matriz(source)
    sheets = build_dinamicas(matrix)
    out = export_dinamicas(
        sheets,
        Path(output_path) if output_path else default_output_path(),
        document_datos=_LAST_DOCUMENT_DATOS,
    )
    publish_dinamicas_actual(out, output_dir=Path(out).parent)
    return out, sheets, summarize(sheets), source


def get_last_ef132_audit() -> dict[str, Any] | None:
    return _LAST_EF132_AUDIT


def get_last_ef137_audit() -> dict[str, Any] | None:
    return _LAST_EF137_AUDIT


def get_last_ef142_audit() -> dict[str, Any] | None:
    return _LAST_EF142_AUDIT


def get_last_ef144_audit() -> dict[str, Any] | None:
    return _LAST_EF144_AUDIT


def get_last_ef149_audit() -> dict[str, Any] | None:
    return _LAST_EF149_AUDIT


def get_last_document_datos() -> dict[str, pd.DataFrame] | None:
    return _LAST_DOCUMENT_DATOS


__all__ = (
    "ACTUAL_FILENAME",
    "EXCLUSIVE_ACREEDOR",
    "EXCLUSIVE_PROVEEDOR",
    "EXCLUSIVE_SHEET",
    "CONFIRMING_SHEET",
    "CONFIRMING_ACREEDORES",
    "EXPECTED_DINAMICAS_SHEETS",
    "GASANDES_ACREEDOR",
    "GASANDES_PROVEEDOR",
    "PENDIENTES_SHEET",
    "NO_ECM_SHEET",
    "PAGO_ME_PATH",
    "USD_ACREEDORES",
    "FINAL_SHEET_ORDER",
    "HIERARCHICAL_SHEETS",
    "SHEET_MAP",
    "IMPORTE_DINAMICA_COL",
    "apply_gasandes_usd_prueba",
    "apply_usd_pago_me",
    "build_dinamicas",
    "build_document_datos",
    "export_dinamicas",
    "get_last_ef132_audit",
    "get_last_ef137_audit",
    "get_last_ef142_audit",
    "get_last_ef144_audit",
    "get_last_ef149_audit",
    "get_last_document_datos",
    "load_valid_nomenclature_concepts",
    "outside_nomenclature_mask",
    "publish_dinamicas_actual",
    "run",
    "summarize",
)
