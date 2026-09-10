from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ...config.logger import LoggerManager
from ...services.output_artifact_service import OutputArtifactService


logger = LoggerManager.get_logger(__name__)


@dataclass(frozen=True)
class ConceptMetric:
    concept: str
    count: int
    percentage: float


@dataclass
class ResumenMetrics:
    total_rows: int = 0
    classified: int = 0
    unclassified: int = 0
    coverage_pct: float = 0.0
    concepts: list[ConceptMetric] = field(default_factory=list)

    @property
    def top5(self) -> list[ConceptMetric]:
        return self.concepts[:5]


def count_matriz_rows(path: Path) -> int:
    """Filas de datos en MATRIZ_FBL1N vía metadata (sin cargar valores)."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["MATRIZ_FBL1N"]
        max_row = worksheet.max_row or 0
        return max(0, max_row - 1)
    finally:
        workbook.close()


def load_resumen_metrics(
    output_service: OutputArtifactService,
    artifact: Path,
    total_rows: int | None = None,
) -> ResumenMetrics:
    """
    Métricas oficiales desde RESUMEN_CONCEPTOS + conteo de filas.

    No lee MATRIZ_FBL1N completa ni NO_CLASIFICADOS.
    """

    if total_rows is None:
        total_rows = count_matriz_rows(artifact)

    resumen = output_service.read_sheet(artifact, "RESUMEN_CONCEPTOS")
    classified = 0
    unclassified = 0
    concepts: list[ConceptMetric] = []

    if {"Concepto", "Cantidad"}.issubset(resumen.columns):
        frame = resumen.copy()
        frame["Concepto"] = frame["Concepto"].astype(str).str.strip()
        frame["Cantidad"] = (
            pd.to_numeric(frame["Cantidad"], errors="coerce").fillna(0).astype(int)
        )

        sin_mask = frame["Concepto"] == "SIN CLASIFICAR"
        if sin_mask.any():
            unclassified = int(frame.loc[sin_mask, "Cantidad"].iloc[0])

        body = frame.loc[~sin_mask].copy()
        if not body.empty:
            classified = int(body["Cantidad"].sum())
            body = body.sort_values(
                by=["Cantidad", "Concepto"],
                ascending=[False, True],
                kind="mergesort",
            )
            denom = total_rows if total_rows > 0 else max(classified, 1)
            for row in body.itertuples(index=False):
                count = int(row.Cantidad)
                concepts.append(
                    ConceptMetric(
                        concept=str(row.Concepto),
                        count=count,
                        percentage=round((count / denom) * 100, 1),
                    )
                )

    coverage = (
        round((classified / total_rows) * 100, 1) if total_rows > 0 else 0.0
    )

    return ResumenMetrics(
        total_rows=total_rows,
        classified=classified,
        unclassified=unclassified,
        coverage_pct=coverage,
        concepts=concepts,
    )


__all__ = (
    "ConceptMetric",
    "ResumenMetrics",
    "count_matriz_rows",
    "load_resumen_metrics",
)
