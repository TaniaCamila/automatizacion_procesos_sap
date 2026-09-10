from __future__ import annotations

import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ...config.logger import LoggerManager
from ...services.output_artifact_service import OutputArtifactService


logger = LoggerManager.get_logger(__name__)

ND = "N/D"
TRANSACTION_LIMIT = 50

# Columnas candidatas solo para la muestra de transacciones.
_TRANSACTION_COLUMNS = (
    "concepto_detectado",
    "concepto_estado",
    "sociedad_nombre",
    "Sociedad",
    "sociedad",
    "Moneda del documento",
    "moneda_del_documento",
    "Importe en moneda doc.",
    "importe_en_moneda_doc",
    "Nº documento",
    "Name",
    "Fecha compensación",
    "fecha_compensacion",
    "mes_compensacion",
)


def _fmt_int(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def _fmt_amount(value: object) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ND
    return f"{number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _safe_str(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ND
    text = str(value).strip()
    return text if text else ND


def _resolve_col(df: pd.DataFrame, *candidates: str) -> str | None:
    for name in candidates:
        if name in df.columns:
            return name
    return None


def _count_matriz_rows(path: Path) -> int:
    """
    Contar filas de datos en MATRIZ_FBL1N sin cargar el contenido.

    Usa openpyxl en modo read_only (max_row - 1 por el encabezado).
    """

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["MATRIZ_FBL1N"]
        max_row = worksheet.max_row or 0
        return max(0, max_row - 1)
    finally:
        workbook.close()


def map_margin(
    output_service: OutputArtifactService,
    path: Path | None = None,
) -> dict:
    """
    Mapear el último Excel MATRIZ al contrato MarginMockData.

    Optimizado: no carga la hoja MATRIZ completa (~69k filas).
    Solo lee 50 filas para transactions + RESUMEN_CONCEPTOS + metadata.
    """

    total_started = time.perf_counter()

    artifact = path or output_service.latest_or_raise()
    meta = output_service.file_metadata(artifact)
    last_update = str(meta["modified_display"])

    # --- Total de registros vía metadata Excel (sin cargar valores) ---
    meta_started = time.perf_counter()
    total_rows = _count_matriz_rows(artifact)
    logger.info(
        "Contando filas MATRIZ (metadata)... %.2f s (%d registros)",
        time.perf_counter() - meta_started,
        total_rows,
    )

    # --- Muestra MATRIZ: solo 50 filas y columnas necesarias ---
    matriz_started = time.perf_counter()
    header = pd.read_excel(artifact, sheet_name="MATRIZ_FBL1N", nrows=0)
    available = list(header.columns)
    usecols = [c for c in _TRANSACTION_COLUMNS if c in available]

    matrix = output_service.read_sheet(
        artifact,
        "MATRIZ_FBL1N",
        usecols=usecols or None,
        nrows=TRANSACTION_LIMIT,
    )
    logger.info(
        "Leyendo MATRIZ... %.2f s (%d filas cargadas)",
        time.perf_counter() - matriz_started,
        len(matrix),
    )

    concept_col = _resolve_col(matrix, "concepto_detectado")
    estado_col = _resolve_col(matrix, "concepto_estado")
    sociedad_col = _resolve_col(
        matrix,
        "sociedad_nombre",
        "Sociedad",
        "sociedad",
    )
    moneda_col = _resolve_col(
        matrix,
        "moneda_del_documento",
        "Moneda del documento",
    )
    importe_col = _resolve_col(
        matrix,
        "importe_en_moneda_doc",
        "Importe en moneda doc.",
    )
    doc_col = _resolve_col(matrix, "Nº documento")
    name_col = _resolve_col(matrix, "Name")
    fecha_col = _resolve_col(
        matrix,
        "Fecha compensación",
        "fecha_compensacion",
        "mes_compensacion",
    )

    # --- RESUMEN_CONCEPTOS: fuente oficial de agregados / gráficos / filtros ---
    resumen_started = time.perf_counter()
    top_concepts: list[dict] = []
    distribution: list[dict] = []
    filter_fields: list[dict] = []
    classified = 0
    unclassified = 0

    try:
        resumen = output_service.read_sheet(artifact, "RESUMEN_CONCEPTOS")
        logger.info(
            "Leyendo RESUMEN... %.2f s",
            time.perf_counter() - resumen_started,
        )

        if {"Concepto", "Cantidad"}.issubset(resumen.columns):
            resumen = resumen.copy()
            resumen["Concepto"] = resumen["Concepto"].astype(str).str.strip()
            resumen["Cantidad"] = pd.to_numeric(
                resumen["Cantidad"],
                errors="coerce",
            ).fillna(0).astype(int)

            sin_mask = resumen["Concepto"] == "SIN CLASIFICAR"
            if sin_mask.any():
                unclassified = int(resumen.loc[sin_mask, "Cantidad"].iloc[0])

            body = resumen.loc[~sin_mask].copy()
            if not body.empty:
                classified = int(body["Cantidad"].sum())
                body_sorted = body.sort_values(
                    by=["Cantidad", "Concepto"],
                    ascending=[False, True],
                    kind="mergesort",
                )

                # Filtro de conceptos: catálogo real completo (RESUMEN).
                concept_options = body_sorted["Concepto"].tolist()
                filter_fields.append(
                    {
                        "id": "concept",
                        "label": "Concepto",
                        "type": "select",
                        "placeholder": "Todos los conceptos",
                        "options": [
                            {"label": opt, "value": opt}
                            for opt in concept_options
                        ],
                    }
                )

                # Distribución y topConcepts: mismos datos, ordenados por cantidad.
                denom = total_rows if total_rows > 0 else max(classified, 1)
                for i, row in enumerate(body_sorted.itertuples(index=False)):
                    concept = str(row.Concepto)
                    count = int(row.Cantidad)
                    pct = round((count / denom) * 100, 1)
                    distribution.append(
                        {
                            "id": f"concept-dist-{i}",
                            "label": concept,
                            "value": float(pct),
                            "displayValue": f"{pct}%",
                        }
                    )

                for i, row in enumerate(
                    body_sorted.head(5).itertuples(index=False)
                ):
                    top_concepts.append(
                        {
                            "id": f"concept-{i}",
                            "label": _safe_str(row.Concepto),
                            "value": _fmt_int(int(row.Cantidad)),
                            "detail": "Registros clasificados",
                        }
                    )
    except Exception as exc:  # noqa: BLE001
        logger.warning(
            "No se pudo leer RESUMEN_CONCEPTOS (%.2f s): %s",
            time.perf_counter() - resumen_started,
            exc,
        )

    build_started = time.perf_counter()

    tx_rows: list[dict[str, str]] = []
    for _, row in matrix.iterrows():
        tx_rows.append(
            {
                "company": _safe_str(row[sociedad_col]) if sociedad_col else ND,
                "document": _safe_str(row[doc_col]) if doc_col else ND,
                "client": _safe_str(row[name_col]) if name_col else ND,
                "concept": _safe_str(row[concept_col]) if concept_col else ND,
                "currency": _safe_str(row[moneda_col]) if moneda_col else ND,
                "amount": (
                    _fmt_amount(row[importe_col]) if importe_col else ND
                ),
                "margin": ND,
                "status": _safe_str(row[estado_col]) if estado_col else ND,
                "date": _safe_str(row[fecha_col]) if fecha_col else ND,
            }
        )

    exports_rows = [
        {
            "file": str(meta["name"]),
            "date": last_update.split(" · ")[0],
            "user": ND,
            "status": "Disponible",
            "size": f"{float(meta['size_bytes']) / (1024 * 1024):.2f} MB",
        }
    ]

    alerts: list[dict] = []
    if unclassified > 0:
        alerts.append(
            {
                "id": "unclassified",
                "title": "Registros sin clasificar",
                "description": (
                    f"{_fmt_int(unclassified)} registros sin concepto detectado."
                ),
                "type": "alert",
                "date": last_update.split(" · ")[0],
            }
        )

    payload = {
        "header": {
            "lastExecution": last_update,
            "status": "Disponible",
        },
        "kpis": [
            {
                "id": "income",
                "label": "Ingresos",
                "value": ND,
                "hint": "No disponible en el artefacto",
                "icon": "income",
            },
            {
                "id": "margin",
                "label": "Margen total",
                "value": ND,
                "hint": "No disponible en el artefacto",
                "icon": "margin",
            },
            {
                "id": "average",
                "label": "Margen promedio",
                "value": ND,
                "hint": "No disponible en el artefacto",
                "icon": "average",
            },
            {
                "id": "records",
                "label": "Registros procesados",
                "value": _fmt_int(total_rows),
                "hint": f"Clasificados: {_fmt_int(classified)}",
                "icon": "records",
            },
        ],
        "filters": filter_fields,
        "transactions": {
            "columns": [
                {"key": "company", "header": "Sociedad"},
                {"key": "document", "header": "Documento"},
                {"key": "client", "header": "Cliente"},
                {"key": "concept", "header": "Concepto"},
                {"key": "currency", "header": "Moneda"},
                {"key": "amount", "header": "Monto", "align": "right"},
                {"key": "margin", "header": "Margen", "align": "right"},
                {"key": "status", "header": "Estado"},
                {"key": "date", "header": "Fecha"},
            ],
            "rows": tx_rows,
        },
        "executive": {
            "summary": [
                {
                    "id": "coverage",
                    "label": "Cobertura de clasificación",
                    "value": (
                        f"{_fmt_int(classified)} / {_fmt_int(total_rows)}"
                        if total_rows
                        else ND
                    ),
                    "detail": f"Sin clasificar: {_fmt_int(unclassified)}",
                },
                {
                    "id": "artifact",
                    "label": "Artefacto",
                    "value": str(meta["name"]),
                    "detail": last_update,
                },
            ],
            "topConcepts": top_concepts,
            "alerts": alerts,
            "distribution": distribution,
            "lastUpdate": last_update,
        },
        "activity": {
            "columns": [
                {"key": "date", "header": "Fecha"},
                {"key": "user", "header": "Usuario"},
                {"key": "event", "header": "Evento"},
                {"key": "result", "header": "Resultado"},
                {"key": "duration", "header": "Duración", "align": "right"},
            ],
            "rows": [],
        },
        "exports": {
            "columns": [
                {"key": "file", "header": "Archivo"},
                {"key": "date", "header": "Fecha"},
                {"key": "user", "header": "Usuario"},
                {"key": "status", "header": "Estado"},
                {"key": "size", "header": "Tamaño", "align": "right"},
            ],
            "rows": exports_rows,
        },
    }

    logger.info(
        "Construyendo payload... %.2f s",
        time.perf_counter() - build_started,
    )
    logger.info(
        "Total endpoint... %.2f s",
        time.perf_counter() - total_started,
    )

    return payload
