from __future__ import annotations

import copy
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "_docx_work"
OUT = ROOT / "documentos_completados"
DIAGRAMS = WORK / "diagrams"
REQ_TEMPLATE = Path(r"C:\Users\cl183288210\Downloads\Documento 01 - Documento de Requerimientos Simplificado.docx")
DESIGN_TEMPLATE = Path(r"C:\Users\cl183288210\Downloads\Documento 02 - Documento de Diseño.docx")
REQ_OUTPUT = OUT / "Documento 01 - Requerimientos Operations Analytics.docx"
DESIGN_OUTPUT = OUT / "Documento 02 - Diseño Operations Analytics.docx"

DATE_LONG = "1 de octubre de 2026"
DATE_SHORT = "01-10-2026"
PROJECT = "Operations Analytics Automatización de Procesos SAP"
BLUE = "#2F5D95"
LIGHT_BLUE = "#EAF2FA"
ORANGE = "#F5A623"
GRAY = "#F3F4F6"
TEXT = "#1F2937"


def set_font(run, size=10.5, bold=False, italic=False, color="000000"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_text(paragraph, text, *, size=10.5, bold=False, italic=False, align=None):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    if align is not None:
        paragraph.alignment = align
    return paragraph


def insert_paragraph_after(paragraph, text=""):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_p.addnext(new_paragraph._p)
    new_paragraph._p.getparent().remove(new_p)
    if text:
        set_paragraph_text(new_paragraph, text)
    return new_paragraph


def paragraph_by_text(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise KeyError(text)


def following_paragraph(doc, heading_text):
    paragraphs = doc.paragraphs
    for idx, p in enumerate(paragraphs):
        if p.text.strip() == heading_text:
            return paragraphs[idx + 1]
    raise KeyError(heading_text)


def replace_text_nodes(doc, replacements):
    for paragraph in doc.element.iter(qn("w:p")):
        nodes = list(paragraph.iter(qn("w:t")))
        if not nodes:
            continue
        combined = "".join(node.text or "" for node in nodes)
        revised = combined
        for old, new in replacements.items():
            revised = revised.replace(old, new)
        if revised != combined:
            nodes[0].text = revised
            for node in nodes[1:]:
                node.text = ""
    for node in doc.element.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
        elif node.text:
            for old, new in replacements.items():
                if old in node.text:
                    node.text = node.text.replace(old, new)


def set_cell_text(cell, text, *, size=9.3, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = 1


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_common_tables(doc):
    for table in doc.tables:
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        set_font(run, size=9.3, bold=run.bold, italic=run.italic)
        if table.rows:
            for cell in table.rows[0].cells:
                shade_cell(cell, "E7E6E6")
                for run in cell.paragraphs[0].runs:
                    set_font(run, size=9.3, bold=True)


def fill_cover_and_metadata(doc, *, design=False):
    title = "Documento de Diseño de Software" if design else "Especificación de Requisitos de Software"
    replace_text_nodes(
        doc,
        {
            "Propuesta de Proyecto       y Especificación de Requisitos de Software": title,
            "Proyecto: [Insertar Nombre de Proyecto]": f"Proyecto: {PROJECT}",
            "[01]": "[1.0]",
            "[Seleccionar fecha]": DATE_LONG,
        },
    )

    table = doc.tables[0]
    values = [DATE_SHORT, "1.0", "Tania Herrera y Camila Armijo", "Documento completado según el estado actual del proyecto"]
    for ci, value in enumerate(values):
        set_cell_text(table.cell(1, ci), value, size=9.0)
    for ci in range(4):
        set_cell_text(table.cell(2, ci), "", size=9.0)

    team = doc.tables[1]
    set_cell_text(team.cell(1, 0), "Tania Herrera", bold=True)
    set_cell_text(team.cell(1, 1), "Líder técnica, desarrolladora principal, automatización, integración y despliegue")
    set_cell_text(team.cell(2, 0), "Camila Armijo", bold=True)
    set_cell_text(team.cell(2, 1), "Analista funcional, documentación, validación y apoyo de QA")
    for ri in range(3, len(team.rows)):
        set_cell_text(team.cell(ri, 0), "")
        set_cell_text(team.cell(ri, 1), "")

    for p in doc.paragraphs:
        if p.text.startswith("Documento validado por las partes en fecha"):
            set_paragraph_text(p, "Documento pendiente de validación por las partes.", italic=True)


def fill_case_table(table, data):
    set_cell_text(table.cell(0, 0), data["id"], bold=True)
    set_cell_text(table.cell(0, 1), data["name"], bold=True)
    set_cell_text(table.cell(1, 1), f"1.0 - {DATE_SHORT}")
    set_cell_text(table.cell(2, 1), data["actors"])
    set_cell_text(table.cell(3, 1), data["objectives"])
    set_cell_text(table.cell(4, 1), data["related"])
    set_cell_text(table.cell(5, 1), data["description"])
    set_cell_text(table.cell(6, 1), data["precondition"])
    set_cell_text(table.cell(7, 0), "Secuencia normal\nFlujo principal sin errores.", bold=True)
    set_cell_text(table.cell(7, 1), "Paso", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(7, 2), "Acción", bold=True)
    for offset, action in enumerate(data["steps"], start=8):
        set_cell_text(table.cell(offset, 1), str(offset - 7), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(offset, 2), action)
    set_cell_text(table.cell(15, 1), data["postcondition"])
    set_cell_text(table.cell(16, 0), "Excepciones\nDesviaciones controladas del flujo normal.", bold=True)
    set_cell_text(table.cell(16, 1), "Paso", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(16, 2), "Acción", bold=True)
    for offset, action in enumerate(data["exceptions"], start=17):
        set_cell_text(table.cell(offset, 1), str(offset - 16), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(offset, 2), action)
    set_cell_text(table.cell(20, 0), "Rendimiento\nLímites técnicos configurados.", bold=True)
    set_cell_text(table.cell(20, 1), "Paso", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(20, 2), "Cota de tiempo", bold=True)
    for offset, (step, limit) in enumerate(data["performance"], start=21):
        set_cell_text(table.cell(offset, 1), step, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(offset, 2), limit)
    set_cell_text(table.cell(23, 1), data["frequency"])
    set_cell_text(table.cell(24, 1), data["comments"])


def build_requirements():
    doc = Document(REQ_TEMPLATE)
    fill_cover_and_metadata(doc)

    content = {
        "Requisitos comunes de las interfaces": "La solución integra una interfaz web de consulta con un pipeline de automatización que opera sobre SAP FBL1N, Excel, SharePoint o OneDrive, Outlook y Power BI. Las interfaces deben conservar la trazabilidad del origen y evitar cambios directos sobre datos productivos desde la capa web.",
        "Interfaces de usuario": "Operations Analytics ofrece una interfaz React con navegación por Inicio, Dashboard, Tesorería, Margen, SEN, IMG, Configuración e Historial. La interfaz consume endpoints de lectura de FastAPI, presenta indicadores y estados, y restringe las vistas según permisos configurados.",
        "Interfaces de hardware": "La ejecución empresarial requiere un equipo Windows 11 de 64 bits con sesión interactiva desbloqueada. Debe disponer de capacidad para ejecutar Python, SAP GUI, Microsoft Excel y Outlook, y de almacenamiento suficiente para fuentes, resultados, respaldos, estados y logs.",
        "Interfaces de software": "El backend utiliza Python 3.13, pandas, openpyxl, pywin32, FastAPI, Uvicorn y Pydantic. La interfaz utiliza React, TypeScript y Vite. La automatización se integra con SAP GUI Scripting, Excel COM, Outlook COM, Windows Task Scheduler, SharePoint o OneDrive y Power BI.",
        "Interfaces de comunicación": "La interfaz web se comunica con FastAPI mediante HTTP o HTTPS y solicitudes JSON. La publicación utiliza una carpeta corporativa sincronizada con SharePoint o OneDrive; las notificaciones se envían mediante Outlook COM. SAP GUI Scripting y Excel COM operan localmente en la sesión autorizada.",
        "Requisitos funcionales": "Los requisitos funcionales se agrupan en tres casos de uso que cubren el ciclo diario, la consulta operativa y el mantenimiento controlado de parámetros.",
        "Requisitos de rendimiento": "El sistema debe comprobar la estabilidad de una fuente mediante tres verificaciones y abortar la espera después de 300 segundos. La extracción SAP tiene un límite configurable de 45 minutos; el procesamiento principal y la construcción del reporte tienen límites de 3 horas y 1 hora. Las solicitudes de la interfaz web utilizan un timeout de 120 segundos.",
        "Seguridad": "El repositorio no debe contener credenciales, archivos productivos ni rutas personales. El acceso a SAP, SharePoint o OneDrive, Outlook y los reportes debe limitarse a usuarios autorizados. La configuración sensible debe residir en archivos locales excluidos del control de versiones.",
        "Fiabilidad": "Cada ejecución debe validar estructura, columnas, fechas, estabilidad del archivo y conteos de filas antes de publicar. El orquestador debe usar bloqueo de concurrencia, huellas SHA-256, estados persistentes, logs y respaldos. Ante fallos de publicación o correo debe reintentar solo la etapa pendiente cuando corresponda.",
        "Disponibilidad": "La tarea se programa de lunes a viernes a las 09:00 en el equipo autorizado. La ejecución completa depende de una sesión interactiva con SAP GUI, Excel y Outlook disponibles. La consulta web puede funcionar en modo de lectura mientras existan artefactos válidos publicados.",
        "Mantenibilidad": "La solución debe conservar una arquitectura modular por capas, configuración centralizada, catálogos externos y pruebas automatizadas. Los cambios deben quedar trazados en Git y en los logs de ejecución, sin duplicar reglas de negocio entre Excel, API e interfaz.",
        "Portabilidad": "El pipeline productivo está diseñado para Windows 11 por sus dependencias de SAP GUI y Microsoft Office COM. El backend Python y la interfaz React pueden ejecutarse en entornos compatibles, pero la automatización completa requiere Python 3.13 de 64 bits y las aplicaciones corporativas instaladas.",
        "Otros Requisitos": "La matriz base debe conservar una fila por transacción, nombres y semántica de las columnas SAP, y registros sin coincidencia en catálogos. No se deben convertir importes entre monedas. Power BI consume los resultados publicados y permanece sujeto a ajustes visuales y funcionales.",
    }
    for heading, text in content.items():
        if heading == "Mantenibilidad":
            insert_paragraph_after(paragraph_by_text(doc, heading), text)
        else:
            set_paragraph_text(following_paragraph(doc, heading), text)

    cases = [
        {
            "id": "RF-01",
            "name": "Ejecutar actualización automática FBL1N",
            "actors": "Primarios: Windows Task Scheduler y analista autorizado. Secundarios: SAP GUI, pipeline Python, Excel, SharePoint o OneDrive y Outlook.",
            "objectives": "Automatizar la obtención, validación, consolidación, generación y publicación de información FBL1N con trazabilidad.",
            "related": "RF-02, RF-03 y requisitos de rendimiento, seguridad, fiabilidad y disponibilidad.",
            "description": "Ejecuta el ciclo diario, detecta cambios en las fuentes, construye la matriz y los reportes, publica los resultados y notifica cuando corresponde.",
            "precondition": "El equipo autorizado tiene sesión interactiva; SAP GUI Scripting, Python, Excel, Outlook y las rutas configuradas están disponibles.",
            "steps": [
                "El programador inicia scripts/run_diario.py a las 09:00.",
                "El wrapper ejecuta la extracción controlada de FBL1N mediante SAP GUI Scripting.",
                "El sistema espera la liberación y estabilidad del archivo exportado.",
                "El orquestador valida estructura y calcula la identidad SHA-256 y semántica de las fuentes.",
                "Si existen cambios, el pipeline genera MATRIZ_FBL1N y los reportes dinámicos.",
                "El sistema valida los artefactos y publica la versión actual en SharePoint o OneDrive.",
                "El sistema registra estado y logs, y Outlook envía la notificación el día configurado.",
            ],
            "postcondition": "Los artefactos válidos quedan publicados y la ejecución queda registrada con estado, huella, rutas, fechas y resultado de correo.",
            "exceptions": [
                "Si SAP GUI o la sesión no están disponibles, se aborta sin reemplazar el artefacto vigente.",
                "Si la fuente es inestable, ilegible o no cumple la estructura, se registra el error y no se publica.",
                "Si falla publicación o correo, el sistema conserva el estado para reintentar solo la etapa pendiente.",
            ],
            "performance": [("2", "Extracción SAP: máximo configurable de 45 minutos."), ("4-6", "Estabilidad: 300 s; proceso principal: 3 h; reporte: 1 h.")],
            "frequency": "Una ejecución programada de lunes a viernes a las 09:00, más ejecuciones manuales controladas.",
            "comments": "El modo detect-only y la simulación permiten validar cambios sin ejecutar publicación ni correo.",
        },
        {
            "id": "RF-02",
            "name": "Consultar resultados y trazabilidad",
            "actors": "Primario: analista de Back Office u Operaciones Comerciales. Secundarios: interfaz React, API FastAPI y repositorio de artefactos.",
            "objectives": "Entregar una vista consistente de indicadores, reportes, estado del sistema e historial sin recalcular FBL1N.",
            "related": "RF-01 y requisitos de seguridad, disponibilidad y rendimiento.",
            "description": "Permite consultar módulos analíticos, indicadores, servicios, configuración visible e historial a partir de los artefactos generados por el pipeline.",
            "precondition": "Existen artefactos válidos y el usuario cuenta con acceso a la interfaz y a los reportes autorizados.",
            "steps": [
                "El usuario abre Operations Analytics.",
                "La aplicación aplica el contexto de autenticación y permisos configurados.",
                "La interfaz solicita datos a los endpoints GET de FastAPI.",
                "La API lee los artefactos, estados JSON y rutas configuradas sin ejecutar el pipeline.",
                "La interfaz presenta indicadores, módulos, estado de servicios e historial.",
                "El usuario revisa el reporte o abre el archivo publicado en SharePoint o OneDrive.",
                "Cuando corresponde, el usuario consulta la visualización disponible en Power BI.",
            ],
            "postcondition": "El usuario obtiene información de lectura actualizada según el último artefacto válido, sin alterar las fuentes ni los resultados.",
            "exceptions": [
                "Si no existe un artefacto válido, la interfaz muestra un estado vacío o no disponible.",
                "Si la API supera 120 segundos o falla, la interfaz informa el error sin presentar datos incompletos como vigentes.",
                "Si el usuario carece de permisos, el módulo o recurso no queda disponible.",
            ],
            "performance": [("3", "Timeout de solicitud web: 120 segundos."), ("4-5", "Lectura de artefactos sin recalcular el pipeline.")],
            "frequency": "Bajo demanda durante la jornada operativa.",
            "comments": "Los endpoints son de lectura; la interfaz no modifica directamente los archivos SAP ni los catálogos productivos.",
        },
        {
            "id": "RF-03",
            "name": "Administrar configuración y catálogos",
            "actors": "Primarios: líder técnica y analista funcional autorizado. Secundarios: archivos de configuración, catálogos Excel, pruebas automatizadas y Git.",
            "objectives": "Mantener sociedades, monedas, conceptos, rutas y parámetros sin introducir secretos ni cambios no validados.",
            "related": "RF-01, RF-02 y requisitos de seguridad, fiabilidad y mantenibilidad.",
            "description": "Actualiza parámetros y catálogos de negocio, valida su estructura y prueba el efecto antes de habilitar una ejecución controlada.",
            "precondition": "El responsable tiene acceso autorizado a los archivos de configuración y conoce la estructura exigida para cada catálogo.",
            "steps": [
                "El responsable identifica el parámetro o catálogo que debe cambiar.",
                "Actualiza sociedades, monedas, conceptos o variables locales con valores autorizados.",
                "Comprueba que no se incorporen credenciales, datos productivos ni rutas personales al repositorio.",
                "Ejecuta validaciones, pruebas automatizadas o detección en modo controlado.",
                "Corrige errores de estructura, rutas o reglas encontrados durante la validación.",
                "Ejecuta el pipeline controlado y revisa logs y artefactos resultantes.",
                "Registra el cambio en Git y actualiza la documentación cuando corresponda.",
            ],
            "postcondition": "La configuración validada queda disponible para futuras ejecuciones y el cambio mantiene trazabilidad técnica y funcional.",
            "exceptions": [
                "Si faltan columnas o existen valores inválidos, el sistema rechaza el catálogo y conserva la configuración anterior.",
                "Si una ruta o aplicación externa no está disponible, la validación falla antes de publicar resultados.",
                "Si existe una ejecución concurrente, el bloqueo impide iniciar un segundo proceso.",
            ],
            "performance": [("4", "Las pruebas y la detección deben finalizar dentro de sus timeouts configurados."), ("6", "La ejecución usa los límites definidos en RF-01.")],
            "frequency": "Según cambios de negocio, incorporación de sociedades o ajustes operativos.",
            "comments": "Las reglas modificables por negocio se mantienen en catálogos Excel; la lógica estructural permanece en módulos versionados.",
        },
    ]
    for table, data in zip(doc.tables[2:5], cases):
        fill_case_table(table, data)

    version = doc.tables[5]
    set_cell_text(version.cell(1, 0), DATE_SHORT, size=9.0)
    set_cell_text(version.cell(1, 1), "Versión 1.0 completada con requisitos e interfaces del proyecto", size=9.0)
    set_cell_text(version.cell(1, 2), "README, arquitectura, código fuente, pruebas y configuración", size=9.0)
    for ri in range(2, len(version.rows)):
        for ci in range(3):
            set_cell_text(version.cell(ri, ci), "", size=9.0)

    format_common_tables(doc)
    set_update_fields(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(REQ_OUTPUT)


CANVAS = (2200, 1300)
REGULAR_FONT = Path(r"C:\Windows\Fonts\calibri.ttf")
BOLD_FONT = Path(r"C:\Windows\Fonts\calibrib.ttf")


def font(size, bold=False):
    return ImageFont.truetype(str(BOLD_FONT if bold else REGULAR_FONT), size)


def canvas():
    image = Image.new("RGB", CANVAS, "white")
    return image, ImageDraw.Draw(image)


def centered(draw, bounds, text, *, size=34, bold=False, fill=TEXT, spacing=6):
    x1, y1, x2, y2 = bounds
    f = font(size, bold)
    bbox = draw.multiline_textbbox((0, 0), text, font=f, align="center", spacing=spacing)
    w, h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=f, fill=fill, align="center", spacing=spacing)


def pbox(draw, bounds, text, *, face=GRAY, edge=BLUE, size=30, bold=False, radius=18):
    draw.rounded_rectangle(bounds, radius=radius, fill=face, outline=edge, width=4)
    centered(draw, bounds, text, size=size, bold=bold)


def parrow(draw, start, end, *, color=TEXT, width=4):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    p1 = (end[0] - length * math.cos(angle - 0.55), end[1] - length * math.sin(angle - 0.55))
    p2 = (end[0] - length * math.cos(angle + 0.55), end[1] - length * math.sin(angle + 0.55))
    draw.polygon([end, p1, p2], fill=color)


def save_image(image, name):
    path = DIAGRAMS / name
    image.save(path, dpi=(220, 220))
    return path


def make_use_case_diagram():
    image, draw = canvas()
    draw.rectangle((500, 100, 1700, 1180), outline=BLUE, width=5)
    centered(draw, (650, 110, 1550, 190), "Operations Analytics", size=42, bold=True, fill=BLUE)
    for cx, title in [(200, "Programador\nde tareas"), (1990, "Analista\nautorizado")]:
        centered(draw, (cx - 130, 165, cx + 130, 285), title, size=30)
        draw.ellipse((cx - 28, 300, cx + 28, 356), outline=TEXT, width=4)
        draw.line((cx, 356, cx, 510), fill=TEXT, width=4); draw.line((cx - 70, 400, cx + 70, 400), fill=TEXT, width=4)
        draw.line((cx, 510, cx - 60, 605), fill=TEXT, width=4); draw.line((cx, 510, cx + 60, 605), fill=TEXT, width=4)
    cases = [(1100, 360, "Actualizar FBL1N\ny publicar"), (1100, 650, "Consultar resultados\ny trazabilidad"), (1100, 940, "Administrar parámetros\ny catálogos")]
    for cx, cy, title in cases:
        draw.ellipse((cx - 390, cy - 90, cx + 390, cy + 90), fill=LIGHT_BLUE, outline=BLUE, width=4)
        centered(draw, (cx - 360, cy - 75, cx + 360, cy + 75), title, size=34)
    draw.line((270, 400, 710, 360), fill=TEXT, width=4)
    for cy in (360, 650, 940): draw.line((1920, 400, 1490, cy), fill=TEXT, width=4)
    return save_image(image, "casos_uso.png")


def entity(draw, bounds, title, fields):
    x1, y1, x2, y2 = bounds
    draw.rectangle(bounds, fill="white", outline=BLUE, width=4)
    draw.rectangle((x1, y1, x2, y1 + 85), fill=LIGHT_BLUE, outline=BLUE, width=4)
    centered(draw, (x1, y1, x2, y1 + 85), title, size=29, bold=True, fill=BLUE)
    f = font(25)
    for i, value in enumerate(fields):
        draw.text((x1 + 22, y1 + 110 + i * 45), value, font=f, fill=TEXT)


def make_er_diagram():
    image, draw = canvas()
    entity(draw, (60, 100, 660, 590), "TRANSACCION_FBL1N", ["documento", "sociedad_codigo", "moneda_codigo", "concepto", "importe", "fecha_compensacion"])
    entity(draw, (800, 100, 1210, 340), "SOCIEDAD", ["codigo (PK)", "nombre"])
    entity(draw, (800, 420, 1210, 660), "MONEDA", ["codigo (PK)", "descripcion"])
    entity(draw, (800, 740, 1210, 980), "CONCEPTO", ["patron (PK)", "clasificacion"])
    entity(draw, (1450, 100, 2050, 485), "EJECUCION", ["sha256 (PK)", "inicio", "estado", "filas", "duracion"])
    entity(draw, (1450, 700, 2050, 1040), "ARTEFACTO", ["ruta (PK)", "tipo", "fecha", "publicado"])
    for start, end, label, pos in [((660, 230),(800,230),"N:1",(700,185)),((660,380),(800,520),"N:1",(690,430)),((660,510),(800,840),"N:1",(680,670)),((1450,300),(660,330),"procesa",(1030,270)),((1750,485),(1750,700),"1:N",(1780,575))]:
        parrow(draw, start, end, color=BLUE); draw.text(pos, label, font=font(23), fill=TEXT)
    centered(draw, (200, 1120, 2000, 1260), "Modelo lógico basado en archivos Excel y estados JSON; no existe una base de datos relacional.", size=28)
    return save_image(image, "modelo_er.png")


def class_box(draw, bounds, title, attrs, methods):
    x1, y1, x2, y2 = bounds
    draw.rectangle(bounds, fill="white", outline=BLUE, width=4)
    draw.rectangle((x1, y1, x2, y1 + 75), fill=LIGHT_BLUE, outline=BLUE, width=4)
    centered(draw, (x1, y1, x2, y1 + 75), title, size=26, bold=True, fill=BLUE)
    f = font(22); y = y1 + 95
    for value in attrs: draw.text((x1 + 16, y), value, font=f, fill=TEXT); y += 36
    draw.line((x1, y + 5, x2, y + 5), fill=BLUE, width=2); y += 22
    for value in methods: draw.text((x1 + 16, y), value, font=f, fill=TEXT); y += 36


def make_class_diagram():
    image, draw = canvas()
    specs = [
        ((40,80,430,370),"Application",["- config","- logger"],["+ run()"]),
        ((510,80,1000,400),"InformeMargenModule",["- services","- processor"],["+ run()","+ load_catalogs()"]),
        ((1080,80,1510,400),"FBL1NProcessor",["- validator"],["+ process(df)","+ normalize_columns()"]),
        ((1600,80,2150,400),"FastAPI Routes",["- dependencies"],["+ get_dashboard()","+ get_history()"]),
        ((170,650,650,980),"ExcelLoader",["- engine"],["+ load_fbl1n(path)","+ load_catalog(path)"]),
        ((830,650,1270,980),"MatrizService",["- classifiers"],["+ build(df)","+ enrich(df)"]),
        ((1460,650,2070,980),"OutputArtifactService",["- output_dir"],["+ latest()","+ read_state()"]),
    ]
    for spec in specs: class_box(draw, *spec)
    for start,end in [((430,220),(510,220)),((1000,220),(1080,220)),((760,400),(410,650)),((820,400),(1050,650)),((1870,400),(1770,650)),((1270,820),(1460,820))]: parrow(draw,start,end,color=BLUE)
    centered(draw, (200, 1100, 2000, 1260), "Las rutas HTTP consumen artefactos generados por el pipeline y no recalculan FBL1N.", size=28)
    return save_image(image, "clases.png")


def make_sequence_diagram():
    image, draw = canvas()
    actors = [(170,"Scheduler"),(520,"run_diario"),(870,"SAP GUI"),(1220,"Orquestador"),(1600,"Excel /\nSharePoint"),(1980,"Outlook")]
    for x,label in actors:
        pbox(draw,(x-130,70,x+130,180),label,face=LIGHT_BLUE,size=25,bold=True)
        draw.line((x,180,x,1180),fill="#9CA3AF",width=3)
    messages=[(170,520,260,"iniciar tarea"),(520,870,370,"extraer FBL1N"),(870,520,480,"archivo exportado"),(520,1220,590,"ejecutar actualización"),(1220,1600,700,"generar y validar"),(1600,1220,810,"artefactos listos"),(1220,1600,920,"publicar versión actual"),(1220,1980,1030,"notificar"),(1220,520,1140,"resultado y estado")]
    for x1,x2,y,label in messages:
        parrow(draw,(x1,y),(x2,y),color=BLUE)
        centered(draw,(min(x1,x2),y-55,max(x1,x2),y-5),label,size=22)
    return save_image(image, "secuencia.png")


def make_component_diagram():
    image, draw = canvas()
    top=[(40,"React + TypeScript\nInterfaz web"),(570,"FastAPI\nAPI de lectura"),(1100,"Servicios de lectura\ny mapeadores"),(1630,"Artefactos Excel\ny estados JSON")]
    for x,label in top: pbox(draw,(x,100,x+480,300),label,face=LIGHT_BLUE,size=28,bold=True)
    pbox(draw,(570,500,1050,700),"Orquestador Python\ny validadores",size=28,bold=True)
    bottom=[(40,"SAP GUI\nFBL1N"),(570,"Excel COM\nReportes dinámicos"),(1100,"SharePoint / OneDrive\nPublicación"),(1630,"Outlook y Power BI\nNotificación y análisis")]
    for x,label in bottom: pbox(draw,(x,900,x+480,1100),label,face="#FFF7E6",edge=ORANGE,size=27,bold=True)
    for start,end in [((520,200),(570,200)),((1050,200),(1100,200)),((1580,200),(1630,200)),((810,500),(810,300)),((280,900),(650,700)),((810,700),(810,900)),((1050,600),(1340,900)),((1050,580),(1870,900)),((1340,900),(1740,300))]: parrow(draw,start,end,color=BLUE)
    centered(draw,(150,1160,2050,1270),"El pipeline productivo requiere Windows y sesión interactiva; la API solo lee resultados ya generados.",size=28)
    return save_image(image, "componentes.png")


def add_diagram(paragraph, image_path, caption, description):
    set_paragraph_text(paragraph, description)
    run = paragraph.add_run()
    run.add_break()
    run.add_picture(str(image_path), width=Inches(6.15))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = insert_paragraph_after(paragraph, caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_font(run, size=9.0, italic=True, color="404040")
    cap.paragraph_format.keep_with_next = False


def build_design():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    diagrams = {
        "use": make_use_case_diagram(), "er": make_er_diagram(), "classes": make_class_diagram(),
        "sequence": make_sequence_diagram(), "components": make_component_diagram(),
    }
    doc = Document(DESIGN_TEMPLATE)
    fill_cover_and_metadata(doc, design=True)
    set_paragraph_text(following_paragraph(doc, "Diseño"), "El diseño implementa una arquitectura modular en capas. El pipeline productivo obtiene y valida datos SAP FBL1N, genera artefactos Excel y publica resultados; la API FastAPI y la interfaz React exponen consultas de lectura sobre esos artefactos.")

    paragraph_by_text(doc, "Diagrama de Casos de Uso").paragraph_format.page_break_before = True
    add_diagram(following_paragraph(doc, "Diagrama de Casos de Uso"), diagrams["use"], "Figura 1 Casos de uso principales", "Los actores principales son el programador de tareas y el analista autorizado. El sistema separa el ciclo automático, la consulta operativa y el mantenimiento de parámetros.")

    paragraph_by_text(doc, "Vista Lógica").paragraph_format.page_break_before = True
    intro = insert_paragraph_after(paragraph_by_text(doc, "Vista Lógica"), "La vista lógica representa los datos y clases que sostienen el procesamiento. La persistencia se basa en archivos Excel y estados JSON, no en una base de datos relacional.")
    intro.paragraph_format.keep_with_next = True
    add_diagram(following_paragraph(doc, "Modelo E-R"), diagrams["er"], "Figura 2 Modelo lógico de datos", "El modelo relaciona transacciones FBL1N con catálogos parametrizables y vincula cada ejecución con sus artefactos publicados.")

    paragraph_by_text(doc, "Diagrama de Clases").paragraph_format.page_break_before = True
    add_diagram(following_paragraph(doc, "Diagrama de Clases"), diagrams["classes"], "Figura 3 Clases y servicios principales", "La aplicación delega en módulos, procesadores y servicios especializados. Las rutas FastAPI leen artefactos mediante servicios de salida y no ejecutan el pipeline.")

    paragraph_by_text(doc, "Diagrama de Secuencia").paragraph_format.page_break_before = True
    add_diagram(following_paragraph(doc, "Diagrama de Secuencia"), diagrams["sequence"], "Figura 4 Secuencia de actualización diaria", "La secuencia muestra la extracción SAP, la ejecución del pipeline, la generación y publicación de artefactos y la notificación controlada.")

    paragraph_by_text(doc, "Vista Despliegue").paragraph_format.page_break_before = True
    intro2 = insert_paragraph_after(paragraph_by_text(doc, "Vista Despliegue"), "La solución se despliega en un equipo Windows autorizado para las integraciones COM y SAP GUI. Los resultados se comparten mediante SharePoint o OneDrive y pueden alimentar Power BI.")
    intro2.paragraph_format.keep_with_next = True
    add_diagram(following_paragraph(doc, "Diagrama de Componentes"), diagrams["components"], "Figura 5 Componentes e integraciones", "Los componentes de lectura web se mantienen separados del pipeline productivo. Esta separación evita que una consulta de interfaz reprocesa o modifique la fuente FBL1N.")

    format_common_tables(doc)
    set_update_fields(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(DESIGN_OUTPUT)


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    build_requirements()
    build_design()
    print(REQ_OUTPUT)
    print(DESIGN_OUTPUT)


if __name__ == "__main__":
    main()
