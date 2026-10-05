from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document


PLACEHOLDERS = [
    "[Insertar Nombre de Proyecto]",
    "[Seleccionar fecha]",
    "nombre del requerimiento funcional",
    "número de versión y fecha",
    "comentarios adicionales",
    "Replicar el punto íntegramente",
    "Adaptar el documento anterior",
]


def main(path: Path) -> None:
    doc = Document(path)
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        custom_xml = [name for name in archive.namelist() if name.startswith("customXml/")]
    found = [item for item in PLACEHOLDERS if item in document_xml]
    print(f"file={path}")
    print(f"zip_error={bad_member}")
    print(f"sections={len(doc.sections)} tables={len(doc.tables)} paragraphs={len(doc.paragraphs)}")
    print(f"media={len(media)} customXml={len(custom_xml)} updateFields={'updateFields' in settings_xml}")
    print(f"placeholders={found}")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
