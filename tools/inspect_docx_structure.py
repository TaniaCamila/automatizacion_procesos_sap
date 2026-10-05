from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def paragraph_record(paragraph, index: int) -> dict:
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "alignment": str(paragraph.alignment),
    }


def cell_record(cell) -> dict:
    return {
        "text": "\n".join(p.text for p in cell.paragraphs),
        "paragraphs": [paragraph_record(p, i) for i, p in enumerate(cell.paragraphs)],
    }


def main() -> None:
    path = Path(sys.argv[1])
    concise = "--concise" in sys.argv[2:]
    doc = Document(path)
    if concise:
        print(f"PATH: {path}")
        print("PARAGRAPHS")
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"P{i:02d} [{p.style.name if p.style else ''}] {p.text}")
        print("TABLES")
        for ti, table in enumerate(doc.tables):
            print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else ''}")
            for ri, row in enumerate(table.rows):
                values = [cell.text.replace("\n", " / ") for cell in row.cells]
                print(f"  R{ri:02d}: " + " || ".join(values))
        print("PACKAGE TEXT")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with zipfile.ZipFile(path) as archive:
            document_root = etree.fromstring(archive.read("word/document.xml"))
            print("CONTENT CONTROLS")
            for si, sdt in enumerate(document_root.xpath(".//w:sdt", namespaces=ns)):
                tag = sdt.xpath("string(./w:sdtPr/w:tag/@w:val)", namespaces=ns)
                alias = sdt.xpath("string(./w:sdtPr/w:alias/@w:val)", namespaces=ns)
                text_value = "".join(sdt.xpath(".//w:sdtContent//w:t/text()", namespaces=ns)).strip()
                print(f"  SDT {si}: tag={tag!r} alias={alias!r} text={text_value!r}")
            for part in sorted(archive.namelist()):
                if not (part.startswith("word/") and part.endswith(".xml")):
                    continue
                root = etree.fromstring(archive.read(part))
                lines = []
                for p in root.xpath(".//w:p", namespaces=ns):
                    text_value = "".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
                    if text_value:
                        lines.append(text_value)
                if lines:
                    print(f"PART {part}")
                    for line in lines:
                        print(f"  {line}")
        return
    result = {
        "path": str(path),
        "paragraphs": [paragraph_record(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [],
        "sections": [],
        "package_parts": [],
    }
    for ti, table in enumerate(doc.tables):
        result["tables"].append({
            "index": ti,
            "style": table.style.name if table.style else None,
            "rows": [
                [cell_record(cell) for cell in row.cells]
                for row in table.rows
            ],
        })
    for si, section in enumerate(doc.sections):
        result["sections"].append({
            "index": si,
            "header": [paragraph_record(p, i) for i, p in enumerate(section.header.paragraphs)],
            "first_page_header": [paragraph_record(p, i) for i, p in enumerate(section.first_page_header.paragraphs)],
            "footer": [paragraph_record(p, i) for i, p in enumerate(section.footer.paragraphs)],
            "first_page_footer": [paragraph_record(p, i) for i, p in enumerate(section.first_page_footer.paragraphs)],
        })
    with zipfile.ZipFile(path) as archive:
        result["package_parts"] = sorted(archive.namelist())
        document_xml = archive.read("word/document.xml")
        result["content_control_count"] = document_xml.count(b"<w:sdt")
        result["drawing_count"] = document_xml.count(b"<w:drawing")
        result["field_count"] = document_xml.count(b"<w:fldChar")
        result["bookmark_count"] = document_xml.count(b"<w:bookmarkStart")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
