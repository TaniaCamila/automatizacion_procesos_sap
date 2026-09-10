#!/usr/bin/env python
"""Reanudar publicación/states tras MATRIZ + DINÁMICAS ya generadas.

Dry-run por defecto (cero escrituras). Escrituras solo con --apply.

No reconstruye MATRIZ ni DINÁMICAS. No abre Power BI / Outlook Send
(MAIL_AUTO_SEND debe ser false).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import load_workbook

from src.config.config import (
    FBL1N_COMPENSATION_CUTOFF,
    FBL1N_FUTURE_ENABLED,
    FBL1N_FUTURE_PATH,
    FBL1N_PATH,
    MAIL_AUTO_SEND,
    OUTPUT_DIR,
    PUBLICATION_DIR,
    resolve_state_dir,
)
from src.loaders.excel_loader import ExcelLoader
from src.modules.actualizacion_automatica.mail import (
    MAIL_PENDING,
    load_mail_state,
    run_mail_stage,
)
from src.modules.actualizacion_automatica.module import (
    PipelineAborted,
    PipelineLock,
    _validate_matrix,
    load_state,
    save_state,
    sha256_file,
    wait_until_stable,
)
from src.modules.dinamicas_finales.module import (
    ACTUAL_FILENAME,
    EXPECTED_DINAMICAS_SHEETS,
    FINAL_SHEET_ORDER,
    _validate_dinamicas_workbook,
    build_dinamicas,
    get_last_ef142_audit,
    get_last_ef149_audit,
    load_matriz,
)
from src.modules.fbl1n_fuentes.detection import (
    DECISION_CHANGED,
    build_combined_state_payload,
    plan_source_detection,
)
from src.modules.publicacion_final.module import (
    DINAMICAS_DEST_NAME,
    PAGO_ME_DEST_NAME,
    PUBLISH_SUCCEEDED,
    load_publish_state,
    run_publish,
)

EXPECTED_MATRIX_ROWS = 92056
CANONICAL_SHEET_COUNT = 16
CUTOFF_DEFAULT = date(2026, 8, 31)

# Partes OOXML de negocio: se comparan byte-a-byte (SHA256) entre ACTUAL y timestamp.
_BUSINESS_EXACT_PARTS: frozenset[str] = frozenset(
    {
        "xl/workbook.xml",
        "xl/sharedStrings.xml",
        "xl/styles.xml",
    }
)
_BUSINESS_PREFIXES: tuple[str, ...] = (
    "xl/worksheets/",
    "xl/pivotCache/",
    "xl/pivotTables/",
    "xl/tables/",
)
# Metadatos/empaquetado OneDrive-SharePoint: nunca deciden equivalencia.
_IGNORED_PREFIXES: tuple[str, ...] = (
    "docProps/",
    "customXml/",
)


class ResumeAborted(Exception):
    """Fallo controlado de la reanudación."""


@dataclass
class ResumePlan:
    matriz: Path
    dinamicas: Path
    actual: Path
    cutoff: date
    decision_status: str
    alias: str
    historical_fp: dict[str, Any]
    decision: Any
    last_state: dict[str, Any]
    fbl_rows: int
    matrix_rows: int
    months: list[str]
    month_labels: list[str]
    latest_compensation_date: str
    dinamicas_sha256: str
    actual_sha256: str
    actual_matches_dinamicas: bool
    state_cas: dict[str, str] = field(default_factory=dict)
    published_cas: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def _sha_optional(path: Path) -> str:
    if not path.is_file():
        return ""
    return sha256_file(path)


def _assert_ooxml(path: Path, label: str) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise ResumeAborted(f"{label} inexistente o vacío: {path}")
    with path.open("rb") as handle:
        magic = handle.read(4)
    if magic[:2] != b"PK":
        raise ResumeAborted(f"{label} no es OOXML/ZIP: {path}")
    try:
        with zipfile.ZipFile(path, "r") as zf:
            bad = zf.testzip()
        if bad is not None:
            raise ResumeAborted(f"{label} ZIP corrupto ({bad}): {path}")
    except zipfile.BadZipFile as exc:
        raise ResumeAborted(f"{label} no es un ZIP válido: {path}") from exc


def _norm_zip_name(name: str) -> str:
    return str(name or "").replace("\\", "/").lstrip("/")


def _is_business_ooxml_part(name: str) -> bool:
    n = _norm_zip_name(name)
    if not n or n.endswith("/"):
        return False
    if any(n.startswith(prefix) for prefix in _IGNORED_PREFIXES):
        return False
    if n in _BUSINESS_EXACT_PARTS:
        return True
    return any(n.startswith(prefix) for prefix in _BUSINESS_PREFIXES)


def _zip_business_digests(path: Path, label: str) -> dict[str, str]:
    """SHA256 de cada parte de negocio; exige testzip() is None."""

    digests: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path, "r") as zf:
            bad = zf.testzip()
            if bad is not None:
                raise ResumeAborted(f"{label} ZIP corrupto ({bad}): {path}")
            for info in zf.infolist():
                name = _norm_zip_name(info.filename)
                if not _is_business_ooxml_part(name):
                    continue
                payload = zf.read(info.filename)
                digests[name] = hashlib.sha256(payload).hexdigest()
    except zipfile.BadZipFile as exc:
        raise ResumeAborted(f"{label} no es un ZIP válido: {path}") from exc
    return digests


def _sheet_order_from_ooxml(path: Path, label: str) -> list[str]:
    """Orden de hojas desde xl/workbook.xml (sin Excel COM)."""

    try:
        with zipfile.ZipFile(path, "r") as zf:
            try:
                raw = zf.read("xl/workbook.xml")
            except KeyError as exc:
                raise ResumeAborted(
                    f"{label} sin xl/workbook.xml: {path}"
                ) from exc
    except zipfile.BadZipFile as exc:
        raise ResumeAborted(f"{label} no es un ZIP válido: {path}") from exc

    root = ET.fromstring(raw)
    sheets: list[str] = []
    for el in root.iter():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "sheet" and "name" in el.attrib:
            sheets.append(str(el.attrib["name"]))
    if not sheets:
        raise ResumeAborted(f"{label} sin hojas en workbook.xml: {path}")
    return sheets


def assert_business_ooxml_equivalent(
    left: Path,
    right: Path,
    *,
    left_label: str = "ACTUAL",
    right_label: str = "DINÁMICAS",
    logger: logging.Logger | None = None,
) -> None:
    """Equivalencia estricta de contenido de negocio OOXML (ignora docProps/customXml)."""

    _assert_ooxml(left, left_label)
    _assert_ooxml(right, right_label)

    left_sha = sha256_file(left)
    right_sha = sha256_file(right)
    if logger is not None:
        logger.info(
            "SHA físicos (solo auditoría; no deciden equivalencia): "
            "%s=%s %s=%s",
            left_label,
            left_sha,
            right_label,
            right_sha,
        )

    left_sheets = _sheet_order_from_ooxml(left, left_label)
    right_sheets = _sheet_order_from_ooxml(right, right_label)
    if left_sheets != right_sheets:
        raise ResumeAborted(
            f"Hojas distintas entre {left_label} y {right_label}: "
            f"{left_sheets!r} vs {right_sheets!r}"
        )
    if len(left_sheets) != CANONICAL_SHEET_COUNT:
        raise ResumeAborted(
            f"{left_label}/{right_label} tienen {len(left_sheets)} hojas; "
            f"se exigen {CANONICAL_SHEET_COUNT}."
        )

    left_parts = _zip_business_digests(left, left_label)
    right_parts = _zip_business_digests(right, right_label)
    left_keys = set(left_parts)
    right_keys = set(right_parts)
    if left_keys != right_keys:
        missing = sorted(right_keys - left_keys)
        extra = sorted(left_keys - right_keys)
        raise ResumeAborted(
            f"Partes de negocio distintas {left_label} vs {right_label}: "
            f"missing_in_{left_label}={missing} extra_in_{left_label}={extra}"
        )
    for name in sorted(left_keys):
        if left_parts[name] != right_parts[name]:
            raise ResumeAborted(
                f"Parte de negocio difiere ({name}): "
                f"{left_label}={left_parts[name][:16]}... "
                f"{right_label}={right_parts[name][:16]}..."
            )
    if logger is not None:
        logger.info(
            "Equivalencia OOXML de negocio OK (%d partes; hojas=%d)",
            len(left_keys),
            len(left_sheets),
        )


def _self_pid_tree() -> set[int]:
    """PID actual + ancestros (p.ej. py.exe → python.exe)."""

    pids = {int(os.getpid())}
    try:
        import psutil  # type: ignore

        current = psutil.Process(os.getpid())
        for parent in current.parents():
            try:
                pids.add(int(parent.pid))
            except (TypeError, ValueError, psutil.Error):  # type: ignore[attr-defined]
                continue
    except Exception:
        pass
    return pids


# Solo estos entrypoints cuentan como Python bloqueante (no IDE/Cursor/LSP).
_PIPELINE_PYTHON_MARKERS: tuple[str, ...] = (
    "run_actualizacion_automatica",
    "construir_dinamicas_finales",
    "publicar_reporte_mensual",
    "reanudar_desde_dinamicas",
    "watch_fbl1n",
    "src\\main.py",
    "src/main.py",
    "scripts\\run_ejecucion_diaria",
    "scripts/run_ejecucion_diaria",
)


def _cmdline_lower(pid: int) -> str:
    try:
        import psutil  # type: ignore

        parts = psutil.Process(pid).cmdline()
        return " ".join(str(p) for p in parts).lower()
    except Exception:
        return ""


def _is_pipeline_python(pid: int, name_base: str) -> bool:
    if name_base not in {"python", "pythonw"}:
        return False
    cmd = _cmdline_lower(pid)
    if not cmd:
        # Sin cmdline (acceso denegado): no tratar como pipeline.
        return False
    return any(marker in cmd for marker in _PIPELINE_PYTHON_MARKERS)


def _count_blocking_processes() -> tuple[dict[str, int], int]:
    """
    Contadores de procesos que bloquean apply/dry-run.

    Returns:
        blockers: EXCEL/SAP/VBS/PBI + python de pipeline FBL1N
        other_python: otros python/pythonw (IDE, etc.) — no bloquean
    """

    hard_names = ("EXCEL", "saplogon", "cscript", "wscript", "PBIDesktop")
    blockers = {name: 0 for name in hard_names}
    blockers["python_pipeline"] = 0
    other_python = 0
    ignore_pids = _self_pid_tree()

    try:
        import psutil  # type: ignore
    except ImportError:
        psutil = None

    if psutil is not None:
        for proc in psutil.process_iter(["pid", "name"]):
            try:
                pname = str(proc.info.get("name") or "")
                pid = int(proc.info.get("pid") or 0)
            except (TypeError, ValueError):
                continue
            if pid in ignore_pids:
                continue
            base = pname.lower().removesuffix(".exe")
            if base in {"python", "pythonw"}:
                if _is_pipeline_python(pid, base):
                    blockers["python_pipeline"] += 1
                else:
                    other_python += 1
                continue
            for wanted in hard_names:
                if base == wanted.lower():
                    blockers[wanted] += 1
        return blockers, other_python

    # Fallback Windows: tasklist sin cmdline → no atribuir python ajeno como pipeline.
    import subprocess

    try:
        raw = subprocess.check_output(
            ["tasklist", "/FO", "CSV", "/NH"],
            text=True,
            errors="ignore",
        )
    except (OSError, subprocess.CalledProcessError):
        return blockers, other_python
    for line in raw.splitlines():
        parts = [p.strip().strip('"') for p in line.split(",")]
        if len(parts) < 2:
            continue
        name = parts[0].lower().removesuffix(".exe")
        try:
            pid = int(parts[1])
        except ValueError:
            continue
        if pid in ignore_pids:
            continue
        if name in {"python", "pythonw"}:
            other_python += 1
            continue
        for wanted in hard_names:
            if name == wanted.lower():
                blockers[wanted] += 1
    return blockers, other_python


def precheck_runtime(*, apply: bool, logger: logging.Logger | None = None) -> None:
    lock = resolve_state_dir() / "pipeline.lock"
    if lock.is_file():
        raise ResumeAborted(f"pipeline.lock presente: {lock}")
    blockers, other_python = _count_blocking_processes()
    hard = {
        k: v
        for k, v in blockers.items()
        if v > 0 and k in {"EXCEL", "saplogon", "cscript", "wscript", "PBIDesktop"}
    }
    pipeline_py = int(blockers.get("python_pipeline") or 0)

    if hard:
        raise ResumeAborted(f"Procesos bloqueantes: {hard}")
    if apply and pipeline_py > 0:
        raise ResumeAborted(
            f"Procesos pipeline Python activos (bloqueantes): {pipeline_py}"
        )
    if not apply and pipeline_py > 0:
        raise ResumeAborted(
            f"Procesos pipeline Python activos (bloquearían --apply): {pipeline_py}"
        )
    if other_python and logger is not None:
        logger.info(
            "Python ajenos detectados=%d (IDE/LSP/u otros); no bloquean la reanudación",
            other_python,
        )


def _require_config() -> date:
    if MAIL_AUTO_SEND:
        raise ResumeAborted("MAIL_AUTO_SEND=true; abortado (debe ser false).")
    if not FBL1N_FUTURE_ENABLED:
        raise ResumeAborted("FBL1N_FUTURE_ENABLED debe ser true para esta reanudación.")
    if FBL1N_COMPENSATION_CUTOFF is None:
        raise ResumeAborted("FBL1N_COMPENSATION_CUTOFF ausente.")
    if FBL1N_COMPENSATION_CUTOFF != CUTOFF_DEFAULT:
        raise ResumeAborted(
            f"FBL1N_COMPENSATION_CUTOFF={FBL1N_COMPENSATION_CUTOFF.isoformat()} "
            f"≠ {CUTOFF_DEFAULT.isoformat()}."
        )
    return FBL1N_COMPENSATION_CUTOFF


def _validate_matriz_rows_and_cutoff(
    matriz: Path,
    fbl1n: Path,
    cutoff: date,
    logger: logging.Logger,
) -> tuple[int, int, list[str], list[str], str]:
    _assert_ooxml(matriz, "MATRIZ")
    fbl_rows, matrix_rows, months, labels = _validate_matrix(matriz, fbl1n, logger)
    if matrix_rows != EXPECTED_MATRIX_ROWS:
        raise ResumeAborted(
            f"MATRIZ filas={matrix_rows}; se exigen {EXPECTED_MATRIX_ROWS}."
        )
    frame = pd.read_excel(
        matriz,
        sheet_name="MATRIZ_FBL1N",
        usecols=["fecha_compensacion"],
        engine="calamine",
    )
    dates = pd.to_datetime(frame["fecha_compensacion"], errors="coerce")
    if dates.isna().all():
        raise ResumeAborted("MATRIZ sin fecha_compensacion parseable.")
    max_ts = dates.max()
    max_iso = max_ts.date().isoformat()
    if max_ts.date() != cutoff:
        raise ResumeAborted(
            f"MATRIZ fecha máxima {max_iso}; se exige {cutoff.isoformat()}."
        )
    return fbl_rows, matrix_rows, months, labels, max_iso


def _validate_dinamicas_workbook_exact(path: Path) -> list[str]:
    _assert_ooxml(path, "DINÁMICAS")
    _validate_dinamicas_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()
    name_set = set(names)
    if name_set != set(EXPECTED_DINAMICAS_SHEETS):
        missing = sorted(set(EXPECTED_DINAMICAS_SHEETS) - name_set)
        extra = sorted(name_set - set(EXPECTED_DINAMICAS_SHEETS))
        raise ResumeAborted(
            f"DINÁMICAS hojas no canónicas missing={missing} extra={extra}"
        )
    if len(names) != CANONICAL_SHEET_COUNT:
        raise ResumeAborted(
            f"DINÁMICAS tiene {len(names)} hojas; se exigen {CANONICAL_SHEET_COUNT}."
        )
    # Orden canónico de las 9 finales (DATOS_* pueden ir al final en cualquier orden).
    finals = [n for n in names if n in FINAL_SHEET_ORDER]
    if finals != list(FINAL_SHEET_ORDER):
        raise ResumeAborted(
            f"Orden de hojas finales distinto al canónico: {finals}"
        )
    return names


def _validate_audits_from_matriz(matriz: Path, logger: logging.Logger) -> None:
    logger.info("Recalculando auditorías EF en memoria (sin exportar Excel)")
    matrix = load_matriz(matriz)
    build_dinamicas(matrix)
    ef142 = get_last_ef142_audit()
    if ef142 is None:
        raise ResumeAborted("No se obtuvo auditoría EF-14.2.")
    if not ef142.get("conciliacion_registros_ok"):
        raise ResumeAborted("conciliacion_registros_ok=False (EF-14.2).")
    if not ef142.get("conciliacion_importe_ok"):
        raise ResumeAborted("conciliacion_importe_ok=False (EF-14.2).")
    if ef142.get("duplicados"):
        raise ResumeAborted("duplicados=True en auditoría EF-14.2.")
    ef149 = get_last_ef149_audit()
    if ef149 is not None:
        if not ef149.get("conciliacion_registros_ok"):
            raise ResumeAborted("conciliacion_registros_ok=False (EF-14.9).")
        if not ef149.get("conciliacion_importe_ok"):
            raise ResumeAborted("conciliacion_importe_ok=False (EF-14.9).")
    logger.info(
        "Auditorías OK duplicados=%s ef142_reg_ok=%s ef142_imp_ok=%s",
        ef142.get("duplicados"),
        ef142.get("conciliacion_registros_ok"),
        ef142.get("conciliacion_importe_ok"),
    )


def build_plan(
    *,
    matriz: Path,
    dinamicas: Path,
    logger: logging.Logger,
    skip_runtime_precheck: bool = False,
) -> ResumePlan:
    cutoff = _require_config()
    if not skip_runtime_precheck:
        precheck_runtime(apply=False, logger=logger)

    matriz = matriz.resolve()
    dinamicas = dinamicas.resolve()
    actual = (Path(OUTPUT_DIR) / ACTUAL_FILENAME).resolve()

    last = load_state()
    if last is None:
        raise ResumeAborted("last_fbl1n_state.json ausente; no se reanuda sin línea base.")

    fbl_rows, matrix_rows, months, labels, max_iso = _validate_matriz_rows_and_cutoff(
        matriz, Path(FBL1N_PATH), cutoff, logger
    )
    _validate_dinamicas_workbook_exact(dinamicas)
    _validate_audits_from_matriz(matriz, logger)

    din_sha = sha256_file(dinamicas)
    act_sha = _sha_optional(actual)
    if not act_sha:
        raise ResumeAborted(f"DINAMICAS_FINALES_ACTUAL ausente: {actual}")
    # OneDrive puede reempaquetar docProps/customXml → SHA físico distinto.
    assert_business_ooxml_equivalent(
        actual,
        dinamicas,
        left_label="ACTUAL",
        right_label="DINÁMICAS",
        logger=logger,
    )
    actual_match = True

    logger.info("Estabilizando FBL1N histórico y futuro")
    hist_fp = wait_until_stable(Path(FBL1N_PATH), logger, checks=2, interval=1.0, timeout=30.0)
    fut_fp = wait_until_stable(Path(FBL1N_FUTURE_PATH), logger, checks=2, interval=1.0, timeout=30.0)
    loader = ExcelLoader()
    hist_df = loader.load_fbl1n(Path(FBL1N_PATH))
    fut_df = loader.load_fbl1n(Path(FBL1N_FUTURE_PATH))

    decision = plan_source_detection(
        future_enabled=True,
        historical_fp=hist_fp,
        last=last,
        historical_df=hist_df,
        future_df=fut_df,
        future_file_sha256=str(fut_fp.get("sha256") or ""),
        compensation_cutoff=cutoff,
        logger=logger,
    )
    if decision.status != DECISION_CHANGED:
        raise ResumeAborted(
            f"Se exige Decision=Changed; se obtuvo {decision.status}: {decision.message}"
        )
    alias = str(decision.source_fbl1n_sha256 or "").strip()
    if not alias:
        raise ResumeAborted("Alias combinado vacío.")
    latest = str(decision.latest_compensation_date or max_iso)
    if latest != cutoff.isoformat():
        # Con cutoff, la fecha combinada debe cerrar en el corte.
        if latest > cutoff.isoformat():
            raise ResumeAborted(
                f"latest_compensation_date={latest} supera cutoff {cutoff.isoformat()}"
            )

    state_dir = resolve_state_dir()
    state_cas = {
        "last_fbl1n_state.json": _sha_optional(state_dir / "last_fbl1n_state.json"),
        "last_publish_state.json": _sha_optional(state_dir / "last_publish_state.json"),
        "last_mail_state.json": _sha_optional(state_dir / "last_mail_state.json"),
    }
    pub_dir = Path(PUBLICATION_DIR)
    published_cas = {
        DINAMICAS_DEST_NAME: _sha_optional(pub_dir / DINAMICAS_DEST_NAME),
        PAGO_ME_DEST_NAME: _sha_optional(pub_dir / PAGO_ME_DEST_NAME),
    }

    return ResumePlan(
        matriz=matriz,
        dinamicas=dinamicas,
        actual=actual,
        cutoff=cutoff,
        decision_status=decision.status,
        alias=alias,
        historical_fp=hist_fp,
        decision=decision,
        last_state=last,
        fbl_rows=fbl_rows,
        matrix_rows=matrix_rows,
        months=months,
        month_labels=labels,
        latest_compensation_date=latest,
        dinamicas_sha256=din_sha,
        actual_sha256=act_sha,
        actual_matches_dinamicas=actual_match,
        state_cas=state_cas,
        published_cas=published_cas,
    )


def _verify_cas(
    plan: ResumePlan,
    *,
    state_keys: tuple[str, ...] | None = None,
    check_published: bool = True,
) -> None:
    """CAS: aborta si un tercero cambió targets aún no tocados por nosotros."""

    state_dir = resolve_state_dir()
    keys = state_keys if state_keys is not None else tuple(plan.state_cas)
    for name in keys:
        expected = plan.state_cas.get(name, "")
        current = _sha_optional(state_dir / name)
        if current != expected:
            raise ResumeAborted(
                f"CAS fallido: {name} cambió durante la reanudación "
                f"(antes={expected[:12] or 'ABSENT'} ahora={current[:12] or 'ABSENT'})."
            )
    if not check_published:
        return
    pub_dir = Path(PUBLICATION_DIR)
    for name, expected in plan.published_cas.items():
        current = _sha_optional(pub_dir / name)
        if current != expected:
            raise ResumeAborted(
                f"CAS fallido: publicado {name} cambió "
                f"(antes={expected[:12] or 'ABSENT'} ahora={current[:12] or 'ABSENT'})."
            )


def _backup_targets(backup_dir: Path, plan: ResumePlan, logger: logging.Logger) -> dict[str, str]:
    backup_dir.mkdir(parents=True, exist_ok=False)
    mapping: dict[str, str] = {}
    state_dir = resolve_state_dir()
    for name in (
        "last_fbl1n_state.json",
        "last_publish_state.json",
        "last_mail_state.json",
    ):
        src = state_dir / name
        if src.is_file():
            dest = backup_dir / "state" / name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest)
            mapping[str(src)] = str(dest)
            logger.info("Backup state %s -> %s", src, dest)
    pub_dir = Path(PUBLICATION_DIR)
    for name in (DINAMICAS_DEST_NAME, PAGO_ME_DEST_NAME):
        src = pub_dir / name
        if src.is_file():
            dest = backup_dir / "published" / name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dest)
            mapping[str(src)] = str(dest)
            logger.info("Backup published %s -> %s", src, dest)
    (backup_dir / "BACKUP_MAP.json").write_text(
        json.dumps(mapping, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return mapping


def _rollback(backup_map: dict[str, str], logger: logging.Logger) -> None:
    for original, backup in backup_map.items():
        src = Path(backup)
        dst = Path(original)
        if not src.is_file():
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        tmp = dst.with_suffix(dst.suffix + ".rollback.tmp")
        shutil.copy2(src, tmp)
        os.replace(tmp, dst)
        logger.info("Rollback restaurado %s", dst)


def _write_manifest(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def apply_plan(
    plan: ResumePlan,
    *,
    logger: logging.Logger,
    work_dir: Path,
    skip_runtime_precheck: bool = False,
) -> dict[str, Any]:
    if not skip_runtime_precheck:
        precheck_runtime(apply=True, logger=logger)
    _require_config()
    _verify_cas(plan)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = work_dir / f"backup_{stamp}"
    manifest_before = {
        "at": datetime.now().isoformat(timespec="seconds"),
        "phase": "before",
        "alias": plan.alias,
        "state_cas": plan.state_cas,
        "published_cas": plan.published_cas,
        "matriz": str(plan.matriz),
        "dinamicas": str(plan.dinamicas),
        "actual": str(plan.actual),
        "decision": plan.decision_status,
        "matrix_rows": plan.matrix_rows,
        "latest_compensation_date": plan.latest_compensation_date,
    }
    _write_manifest(work_dir / "MANIFEST_BEFORE.json", manifest_before)

    backup_map = _backup_targets(backup_dir, plan, logger)
    lock = PipelineLock()
    started_writes = False
    try:
        lock.acquire()
        # Antes de tocar nada: states + publicados intactos.
        _verify_cas(plan)
        started_writes = True

        # Publicación productiva (lee ACTUAL de OUTPUT_DIR).
        published = run_publish(plan.alias, logger)
        if published.status != PUBLISH_SUCCEEDED:
            raise ResumeAborted(
                f"Publicación no Succeeded: {published.status} {published.error}"
            )

        # run_publish ya reescribió last_publish_state + archivos publicados.
        # Solo CAS de lo que aún no tocamos.
        _verify_cas(
            plan,
            state_keys=("last_fbl1n_state.json", "last_mail_state.json"),
            check_published=False,
        )

        state_payload = build_combined_state_payload(
            historical_fp=plan.historical_fp,
            decision=plan.decision,
            started_at=datetime.now().isoformat(timespec="seconds"),
            last_success_at=datetime.now().isoformat(timespec="seconds"),
            fbl1n_rows=plan.fbl_rows,
            matrix_rows=plan.matrix_rows,
            months=plan.months,
            matriz=str(plan.matriz),
            dinamicas=str(plan.dinamicas),
            actual=str(plan.actual),
            log=str((work_dir / "reanudar_desde_dinamicas.log").resolve()),
        )
        save_state(state_payload)

        _verify_cas(
            plan,
            state_keys=("last_mail_state.json",),
            check_published=False,
        )
        mail = run_mail_stage(
            plan.alias,
            logger,
            months=plan.months,
            latest_compensation_date=plan.latest_compensation_date,
            future_enabled=True,
            decision=DECISION_CHANGED,
            publish_ok=True,
        )
        if mail.executed:
            raise ResumeAborted("mail.executed=True; no se permite Outlook/Send.")
        if mail.status != MAIL_PENDING:
            raise ResumeAborted(f"mail_status={mail.status}; se exige Pending.")

        pub_state = load_publish_state()
        if not pub_state or str(pub_state.get("publish_status")) != PUBLISH_SUCCEEDED:
            raise ResumeAborted("last_publish_state no quedó Succeeded.")

        after = {
            "at": datetime.now().isoformat(timespec="seconds"),
            "phase": "after",
            "alias": plan.alias,
            "publish_status": PUBLISH_SUCCEEDED,
            "mail_status": MAIL_PENDING,
            "mail_subject": mail.mail_subject,
            "state_sha": {
                "last_fbl1n_state.json": _sha_optional(
                    resolve_state_dir() / "last_fbl1n_state.json"
                ),
                "last_publish_state.json": _sha_optional(
                    resolve_state_dir() / "last_publish_state.json"
                ),
                "last_mail_state.json": _sha_optional(
                    resolve_state_dir() / "last_mail_state.json"
                ),
            },
            "published_sha": {
                DINAMICAS_DEST_NAME: _sha_optional(
                    Path(PUBLICATION_DIR) / DINAMICAS_DEST_NAME
                ),
                PAGO_ME_DEST_NAME: _sha_optional(
                    Path(PUBLICATION_DIR) / PAGO_ME_DEST_NAME
                ),
            },
            "backup_dir": str(backup_dir),
        }
        _write_manifest(work_dir / "MANIFEST_AFTER.json", after)
        return after
    except Exception as exc:
        logger.error("Fallo en apply: %s", exc)
        if started_writes and backup_map:
            try:
                _rollback(backup_map, logger)
            except Exception as rollback_exc:
                logger.error("Rollback incompleto: %s", rollback_exc)
                raise ResumeAborted(
                    f"Fallo apply ({exc}); rollback también falló ({rollback_exc})."
                ) from rollback_exc
            raise ResumeAborted(f"Fallo apply; rollback ejecutado: {exc}") from exc
        raise
    finally:
        lock.release()


def plan_to_public_dict(plan: ResumePlan) -> dict[str, Any]:
    return {
        "matriz": str(plan.matriz),
        "dinamicas": str(plan.dinamicas),
        "actual": str(plan.actual),
        "cutoff": plan.cutoff.isoformat(),
        "decision": plan.decision_status,
        "alias": plan.alias,
        "matrix_rows": plan.matrix_rows,
        "fbl_rows": plan.fbl_rows,
        "months": plan.months,
        "latest_compensation_date": plan.latest_compensation_date,
        "actual_matches_dinamicas": plan.actual_matches_dinamicas,
        "dinamicas_sha256": plan.dinamicas_sha256,
        "actual_sha256": plan.actual_sha256,
        "state_cas": plan.state_cas,
        "published_cas": plan.published_cas,
        "mail_auto_send": bool(MAIL_AUTO_SEND),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Reanudar publicación y states tras MATRIZ+DINÁMICAS ya generadas. "
            "Dry-run por defecto; requiere --apply para escribir."
        )
    )
    parser.add_argument("--matriz", type=Path, required=True)
    parser.add_argument("--dinamicas", type=Path, required=True)
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Ejecutar escrituras (publicación + states + mail Pending).",
    )
    parser.add_argument(
        "--work-dir",
        type=Path,
        default=None,
        help="Carpeta para manifest/backup/log (default: temp/reanudar_desde_dinamicas_<ts>).",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    work_dir: Path | None = None
    if args.apply:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        work_dir = Path(
            args.work_dir or (ROOT / "temp" / f"reanudar_desde_dinamicas_{stamp}")
        )
        work_dir.mkdir(parents=True, exist_ok=True)
        handlers.append(
            logging.FileHandler(
                work_dir / "reanudar_desde_dinamicas.log", encoding="utf-8"
            )
        )
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=handlers,
        force=True,
    )
    logger = logging.getLogger("reanudar_desde_dinamicas")
    try:
        plan = build_plan(matriz=args.matriz, dinamicas=args.dinamicas, logger=logger)
        public = plan_to_public_dict(plan)
        print(
            json.dumps(
                {"mode": "dry-run" if not args.apply else "apply-plan", **public},
                indent=2,
            )
        )
        if not args.apply:
            print("DRY-RUN OK: cero escrituras productivas.")
            return 0
        assert work_dir is not None
        after = apply_plan(plan, logger=logger, work_dir=work_dir)
        print(json.dumps({"mode": "apply", "result": after}, indent=2))
        print("APPLY OK")
        return 0
    except (ResumeAborted, PipelineAborted) as exc:
        logger.error("%s", exc)
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
